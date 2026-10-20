"""A rounded block carrying a heading, a line of copy and an optional icon.

The plate's colour is a declared pair, so its ink is contrast-checked against its own
fill rather than against the slide's.
"""

from __future__ import annotations


from pptx.util import Inches

from deckwright.errors import LayoutError
from deckwright.icons.draw import place_icon
from deckwright.layouts.components import BodyResult, RevealItem, component
from deckwright.layouts.registry import SlideCtx
from deckwright.theme.media import resolve_media
from deckwright.theme.model import Rect
from deckwright.utils.a11y import describe
from deckwright.utils.shapes import ANCHOR, para, rrect, textbox
from deckwright.utils.text import estimate_caveat, inches_apart, overlong_word

from deckwright.components._shape import (
    anchored,
    figure_text,
    _where,
    known_fields,
    number,
    pair_named,
    shadow,
    visible_edge,
)
from deckwright.components._cardgeom import (
    RADIUS_DEFAULT,
    copy_height,
    corner_reach,
    icon_side,
)
from deckwright.components._shared import HEAD_SPACE_AFTER_PT, mark_colour

_FIELDS = ("pair", "heading", "body", "icon", "alt", "decorative", "radius", "shadow")
_PAIR_DEFAULT = "surface"
_MAX_RADIUS = 0.5
# The icon is a mark beside the copy, not an illustration: two heading line-heights square.


@component("card")
def card(ctx: SlideCtx) -> BodyResult:
    """Paint the placement as a plate and set the card's contents on it."""
    known_fields(ctx, _FIELDS)
    heading = _text(ctx, "heading")
    copy = _text(ctx, "body")
    icon = ctx.body.get("icon")
    if not (heading or copy or icon):
        raise LayoutError(
            f"slide {ctx.spec.index} (component 'card'): a card needs a 'heading', a "
            f"'body' or an 'icon' — an empty plate is the 'panel' component"
        )
    pair = pair_named(ctx, _PAIR_DEFAULT)
    rect = ctx.body_rect
    radius = _radius(ctx)
    plate = rrect(
        ctx.slide,
        rect.left,
        rect.top,
        rect.width,
        rect.height,
        ctx.rgb(pair.bg),
        line=visible_edge(ctx, pair),
        radius=radius,
    )
    shadow(ctx, plate)
    stands_off = str(ctx.body.get("pair", _PAIR_DEFAULT)) != "surface"
    ctx.manifest.record(
        plate,
        bg=pair.bg,
        fill=pair.bg if stands_off else None,
        ground=ctx.behind(rect, ink=pair.bg) if stands_off else None,
    )
    # Chrome drawn over this plate reads its fill, not the slide's surface.
    ctx.painted.append((rect, pair.bg))
    shapes: list[RevealItem] = [(plate.shape_id, "surface")]

    inset = ctx.grid.gutter
    side = inset + corner_reach(rect, radius, depth=inset)
    area = Rect(rect.left + side, rect.top + inset, rect.width - 2 * side, rect.height - 2 * inset)
    if area.width <= 0 or area.height <= 0:
        raise LayoutError(
            f"slide {ctx.spec.index} (component 'card'): the placement is smaller than "
            f"the {inset:.2f}in the theme's gutter insets a card by on every side"
        )
    if icon:
        area = _settled(ctx, area, heading=heading, copy=copy)
        area = _icon(ctx, area, str(icon), shapes, worded=bool(heading or copy))
    elif "alt" in ctx.body or "decorative" in ctx.body:
        raise LayoutError(
            f"{_where(ctx)}: 'alt' and 'decorative' describe the card's icon, and it has none"
        )
    if heading or copy:
        shapes.append(
            (
                _copy(
                    ctx,
                    area,
                    heading=heading,
                    copy=copy,
                    ink=pair.fg,
                    bg=pair.bg,
                    anchor="top" if icon else ctx.anchor,
                ),
                "text",
            )
        )
    return BodyResult(groups=[shapes], height=rect.height)


def _settled(ctx: SlideCtx, area: Rect, *, heading: str, copy: str) -> Rect:
    """What is left of ``area`` once ``anchor:`` has moved an icon and its type down
    together; the type's own frame anchor cannot move the icon drawn above it."""
    depth = icon_side(ctx.theme)
    if heading or copy:
        depth += ctx.grid.gutter + copy_height(
            ctx.theme, width=area.width, heading=heading, copy=copy
        )
    slack = area.height - depth
    if ctx.anchor == "top" or slack <= 0:
        return area
    drop = slack / 2 if ctx.anchor == "middle" else slack
    return Rect(area.left, area.top + drop, area.width, area.height - drop)


def _refuse_overlong(ctx: SlideCtx, text: str, *, width: float, size_pt: float, face: str) -> None:
    found = overlong_word(text, width_in=width, size_pt=size_pt, face=face)
    if found is None:
        return
    word, need = found
    raise LayoutError(
        f"{_where(ctx)}: the word {word!r} needs {need:.2f}in at {size_pt:.1f}pt but the card "
        f"leaves {width:.2f}in, so it would break mid-word — widen the placement, or use a "
        f"shorter word{estimate_caveat(face)}"
    )


def _text(ctx: SlideCtx, key: str) -> str:
    value = ctx.body.get(key)
    return "" if value is None else str(value)


def _radius(ctx: SlideCtx) -> float:
    value = number(ctx, "radius", default=RADIUS_DEFAULT)
    if not 0.0 <= value <= _MAX_RADIUS:
        raise LayoutError(
            f"slide {ctx.spec.index} (component 'card'): 'radius' is a fraction of the "
            f"plate's short side, 0..{_MAX_RADIUS} ({_MAX_RADIUS} is a stadium); "
            f"got {value}"
        )
    return value


def _icon(ctx: SlideCtx, area: Rect, name: str, shapes: list[RevealItem], *, worded: bool) -> Rect:
    """Place the icon at the top of the card's inner area; return what is left below it.

    A bare name is a glyph, drawn as vector and painted from the palette; anything
    carrying a suffix or a path separator is a picture file, placed as it is.
    """
    side = icon_side(ctx.theme)
    if side > area.height:
        raise LayoutError(
            f"slide {ctx.spec.index} (component 'card'): the icon alone wants "
            f"{side:.2f}in of the {area.height:.2f}in inside the plate — grow the "
            f"placement or drop the icon"
        )
    box = anchored(area, side, side, align=ctx.align, anchor="top")
    if _is_file(name):
        media = resolve_media(name, template=ctx.theme.template, roots=ctx.media_roots)
        shape = ctx.slide.shapes.add_picture(
            str(media), Inches(box.left), Inches(box.top), Inches(side), Inches(side)
        )
        alt, decorative = figure_text(ctx)
        # Beside the card's own words an unlabelled icon is a mark; alone it is the content.
        describe(shape, alt=alt, decorative=decorative or (worded and alt is None))
        ctx.manifest.record(shape, rendered="image")
    else:
        fill = mark_colour(ctx, box)
        shape = place_icon(ctx.slide, name, box, fill=fill, theme=ctx.theme)
        alt, decorative = figure_text(ctx)
        if alt is not None or decorative:
            describe(shape, alt=alt, decorative=decorative)
        ctx.manifest.record(shape, fg=fill, bg=ctx.behind(box, ink=fill))
    shapes.append((shape.shape_id, "figure"))
    gap = ctx.grid.gutter
    return Rect(area.left, area.top + side + gap, area.width, area.height - side - gap)


def _is_file(name: str) -> bool:
    return "." in name or "/" in name or "\\" in name


def _copy(
    ctx: SlideCtx, area: Rect, *, heading: str, copy: str, ink: str, bg: str, anchor: str
) -> int:
    """Set the card's type in one frame, refusing copy that will not fit the plate."""
    head_style, body_style = ctx.style("head"), ctx.style("body")
    for text, style in ((heading, head_style), (copy, body_style)):
        _refuse_overlong(
            ctx, text, width=area.width, size_pt=style.size, face=ctx.theme.font_for(style)
        )
    needed = copy_height(ctx.theme, width=area.width, heading=heading, copy=copy)
    if needed > area.height:
        wants, left = inches_apart(needed, area.height)
        raise LayoutError(
            f"{_where(ctx)}: the card's type wants "
            f"{wants} but only {left} is left inside the plate — "
            f"shorten the copy or grow the placement"
            f"{estimate_caveat(ctx.theme.font_for(head_style), ctx.theme.font_for(body_style))}"
        )
    tf = textbox(ctx.slide, area.left, area.top, area.width, area.height, anchor=ANCHOR[anchor])
    lines: list[str] = []
    if heading:
        para(
            tf,
            heading,
            head_style.size,
            ctx.rgb(ink),
            bold=head_style.bold,
            italic=head_style.italic,
            align=ctx.text_align(),
            first=True,
            space_after=HEAD_SPACE_AFTER_PT,
            font=ctx.theme.font_for(head_style),
        )
        lines.append(heading)
    if copy:
        para(
            tf,
            copy,
            body_style.size,
            ctx.rgb(ink),
            italic=body_style.italic,
            align=ctx.text_align(),
            first=not heading,
            space_after=0,
            font=ctx.theme.font_for(body_style),
        )
        lines.append(copy)
    ctx.manifest.record(
        tf._parent,
        lines=lines,
        font_pt=head_style.size if heading else body_style.size,
        line_pt=([head_style.size] if heading else []) + ([body_style.size] if copy else []),
        fg=ink,
        bg=bg,
    )
    return tf._parent.shape_id
