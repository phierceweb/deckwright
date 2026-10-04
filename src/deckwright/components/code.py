"""A monospace listing on a themed plate.

The only other way to put code on a slide is `document:`, which rasterizes through a
browser — this draws real text, so it is selectable, themed, and in the manifest as
lines QA can measure.
"""

from __future__ import annotations

from math import ceil, floor

from pptx.enum.text import MSO_ANCHOR
from pptx.util import Inches, Pt

from deckwright.errors import LayoutError
from deckwright.layouts.components import BodyResult, RevealItem, component
from deckwright.layouts.registry import SlideCtx
from deckwright.utils.shapes import para, rrect
from deckwright.utils.text import LINE_HEIGHT, estimate_caveat, listing_em, listing_rows

from deckwright.components._shape import flag, known_fields, pair_named
from deckwright.components._shared import head as head_line

_FIELDS = ("lines", "text", "heading", "pair", "accent", "wrap", "size")
_PAIR_DEFAULT = "surface"
# The plate's own padding and the space a heading takes, in inches.
_PAD_X = 0.26
_PAD_Y = 0.2
_HEADING_H = 0.42
# The plate's corner radius, and the share of it a roundRect's text rectangle is inset from
# each side by (1 - cos 45°, its preset geometry's ``il``), inside the frame's own margins.
_RADIUS = 0.04
_INSET = 0.29289
# Line advance as a multiple of the type size, sizing the plate and set exactly in each
# paragraph, and the rung a listing is set at.
_LINE_ADVANCE = 1.55
_SIZE_RUNG = "caption"
# Columns per tab stop. A tab is drawn as spaces, which PowerPoint and LibreOffice agree on.
_TAB = 4


@component("code")
def code(ctx: SlideCtx) -> BodyResult:
    """One plate, one line of monospace per entry."""
    known_fields(ctx, _FIELDS)
    lines = _lines(ctx)
    pair = pair_named(ctx, _PAIR_DEFAULT)
    accents = tuple(str(a) for a in (ctx.body.get("accent") or ()))
    heading = str(ctx.body.get("heading", "")) or None
    size = _size(ctx)

    rect = ctx.body_rect
    top = rect.top
    groups: list[list[RevealItem]] = []
    if heading is not None:
        frame = _heading(ctx, heading, top)
        groups.append([(frame.shape_id, "text")])
        top += _HEADING_H

    wrap = flag(ctx, "wrap")
    face = ctx.theme.mono
    pitch = Pt(size * _LINE_ADVANCE)
    advance = pitch.inches
    if _set_width(rect.width, advance + 2 * _PAD_Y) <= 0:
        raise LayoutError(
            f"slide {ctx.spec.index} (component 'code'): the placement is {rect.width:.2f}in "
            f"wide, and a plate's padding and rounded corners leave no room for a line in it — "
            f"widen the placement"
        )
    if wrap:
        rows = _wrapped_rows(
            lines,
            width=rect.width,
            room=rect.bottom - top,
            advance=advance,
            size_pt=size,
            face=face,
        )
    else:
        rows = len(lines)
    plate_h = advance * rows + 2 * _PAD_Y
    if not wrap:
        width_in = _set_width(rect.width, plate_h)
        _refuse_overwide(ctx, lines, width_in=width_in, size_pt=size, face=face)
    if top + plate_h > rect.bottom:
        wrapped = f" wrap to {rows} rows and" if rows != len(lines) else ""
        raise LayoutError(
            f"slide {ctx.spec.index} (component 'code'): {len(lines)} lines{wrapped} need "
            f"{plate_h:.2f}in but only {rect.bottom - top:.2f}in is left — shorten the "
            f"listing, split the slide, or lower 'size'"
            f"{estimate_caveat(face, mono=True) if wrapped else ''}"
        )

    plate = rrect(ctx.slide, rect.left, top, rect.width, plate_h, ctx.rgb(pair.bg), radius=_RADIUS)
    frame = plate.text_frame
    frame.word_wrap = wrap
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = Inches(_PAD_X)
    # LibreOffice sets a row's surplus over single spacing above it, so the first row's
    # surplus comes off the top margin and goes to the bottom one.
    lift = min(Pt(size * (_LINE_ADVANCE - LINE_HEIGHT)), Inches(_PAD_Y))
    frame.margin_top = Inches(_PAD_Y) - lift
    frame.margin_bottom = Inches(_PAD_Y) + lift

    ink = ctx.rgb(ctx.ink_on(pair.bg))
    lit = ctx.rgb(ctx.accent_on(pair.bg, size_pt=size))
    drawn: list[str] = []
    for index, line in enumerate(lines):
        emphasised = any(line.lstrip().startswith(a) for a in accents)
        para(
            frame,
            line or " ",
            size,
            lit if emphasised else ink,
            bold=emphasised,
            first=(index == 0),
            space_after=0,
            font=face,
            links=False,
        ).line_spacing = pitch
        drawn.append(line or " ")
    ctx.manifest.record(
        plate,
        lines=drawn,
        font_pt=size,
        fg=ctx.ink_on(pair.bg),
        bg=pair.bg,
        plate=True,
        literal=True,
    )
    groups.append([(plate.shape_id, "text")])
    ctx.painted.append((_plate_rect(rect, top, plate_h), pair.bg))
    return BodyResult(groups=groups, height=(top - rect.top) + plate_h)


def _lines(ctx: SlideCtx) -> list[str]:
    """``lines:`` as written, or ``text:`` split — one of the two is required."""
    raw, text = ctx.body.get("lines"), ctx.body.get("text")
    if raw is not None and text is not None:
        raise LayoutError(
            f"slide {ctx.spec.index} (component 'code'): give 'lines' or 'text', "
            f"not both — 'text' is a block scalar, 'lines' a list"
        )
    if text is not None:
        return [line.expandtabs(_TAB) for line in str(text).rstrip("\n").split("\n")]
    if not isinstance(raw, list) or not raw:
        raise LayoutError(
            f"slide {ctx.spec.index} (component 'code'): needs 'lines' (a list) or "
            f"'text' (a block scalar) — a listing with nothing in it draws an empty plate"
        )
    return [str(line).expandtabs(_TAB) for line in raw]


def _set_width(width: float, height: float) -> float:
    """The width a plate ``width`` by ``height`` sets its lines in, inside its text rectangle."""
    return width - 2 * (_INSET * _RADIUS * min(width, height) + _PAD_X)


def _wrapped_rows(
    lines: list[str], *, width: float, room: float, advance: float, size_pt: float, face: str
) -> int:
    """Rows ``lines`` wrap to on a plate as deep as those rows, or more if none fits ``room``.

    A shallower plate insets its text less, so each pass, from a plate filling ``room``,
    sets the lines no narrower than the last and needs no more rows.
    """
    rows, height = 0, room
    while True:
        set_in = _set_width(width, height)
        fit = sum(listing_rows(line, width_in=set_in, size_pt=size_pt, face=face) for line in lines)
        depth = advance * fit + 2 * _PAD_Y
        if fit == rows or depth > room:
            return fit
        rows, height = fit, depth


def _refuse_overwide(
    ctx: SlideCtx, lines: list[str], *, width_in: float, size_pt: float, face: str
) -> None:
    """Refuse an unwrapped line wider than the ``width_in`` its plate sets lines in.

    Raises:
        LayoutError: a line is wider than ``width_in``.
    """
    capacity = width_in * 72 / size_pt
    for number, line in enumerate(lines, 1):
        need = listing_em(line, face)
        if need > capacity:
            # Rounded apart, so a line a hair too wide never reads as needing what it has.
            need_in = ceil(need * size_pt / 72 * 100) / 100
            room_in = max(0.0, floor(width_in * 100) / 100)
            raise LayoutError(
                f"slide {ctx.spec.index} (component 'code'): line {number} needs "
                f"{need_in:.2f}in but the plate holds {room_in:.2f}in — shorten "
                f"it, widen the placement, lower 'size', or set 'wrap: true'"
                f"{estimate_caveat(face, mono=True)}"
            )


def _size(ctx: SlideCtx) -> float:
    raw = ctx.body.get("size")
    if raw is None:
        return ctx.style(_SIZE_RUNG).size
    try:
        value = float(raw)
    except (TypeError, ValueError):
        raise LayoutError(
            f"slide {ctx.spec.index} (component 'code'): 'size' is a point size, got {raw!r}"
        ) from None
    if value < ctx.theme.min_pt:
        raise LayoutError(
            f"slide {ctx.spec.index} (component 'code'): 'size' {value}pt is below the "
            f"theme's {ctx.theme.min_pt:.1f}pt minimum"
        )
    return value


def _heading(ctx: SlideCtx, text: str, top: float):
    from deckwright.utils.shapes import textbox

    frame = textbox(ctx.slide, ctx.body_rect.left, top, ctx.body_rect.width, _HEADING_H)
    ink, paper = head_line(ctx, frame, text, first=True)
    ctx.manifest.record(frame._parent, text=text, font_pt=ctx.style("head").size, fg=ink, bg=paper)
    return frame._parent


def _plate_rect(rect, top: float, height: float):
    from deckwright.theme.model import Rect

    return Rect(left=rect.left, top=top, width=rect.width, height=height)
