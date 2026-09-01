"""Which of a slide's words are its chrome, and which are body.

A title placeholder outranks a shape named ``s1.chrome.title``, which outranks a line
read by its position in a stacked ``s1.chrome`` frame; a losing claim keeps its place in
reading order as body.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from statistics import median

from pptx.enum.shapes import PP_PLACEHOLDER

from deckwright.layouts.chrome import CHROME_ORDER
from deckwright.spec._tree import Unreadable, flat

Lines = tuple[str, ...]

_TITLES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)

# How much larger than the slide's median run a lone top line must be set before it
# reads as the title.
_TITLE_RATIO = 1.4

_CHROME_NAME = re.compile(rf"^s\d+\.chrome(?:\.({'|'.join(CHROME_ORDER)}))?$")

_PLACEHOLDER, _NAMED, _STACKED, _BODY = 0, 1, 2, 3


@dataclass(frozen=True)
class _Text:
    """One run of words, the type they are set in, and their claim on a chrome field.

    ``rank`` orders competing claims: a real placeholder outranks a shape name, which
    outranks a line read by its position in a stacked frame.
    """

    lines: Lines
    largest: float
    field: str | None = None
    rank: int = _BODY


def _placeholder_type(shape) -> PP_PLACEHOLDER | None:
    if not shape.is_placeholder:
        return None
    return shape.placeholder_format.type


def _run_sizes(shape) -> list[float]:
    return [
        run.font.size.pt
        for para in shape.text_frame.paragraphs
        for run in para.runs
        if run.font.size is not None
    ]


def _largest(shape) -> float:
    return max(_run_sizes(shape), default=0.0)


def _stacked(shape, taken: frozenset[str]) -> dict[str, str]:
    """The fields a shared chrome frame carries, one paragraph each.

    ``layouts/compose.py`` stacks three lines in :data:`CHROME_ORDER`, so three are read
    by position. Two cannot say which fields they are unless a sibling frame already
    holds the title: a kicker over a subtitle is set exactly like a kicker over a title.
    """
    lines = [flat(para.text) for para in shape.text_frame.paragraphs if flat(para.text)]
    if not lines:
        return {}
    if len(lines) == 1:
        return {"title": lines[0]}
    if len(lines) >= len(CHROME_ORDER):
        kicker, title, *subtitle = lines
        return {"kicker": kicker, "title": title, "subtitle": " ".join(subtitle)}
    if "title" in taken:
        return {"kicker": lines[0], "subtitle": lines[1]}
    sizes = _paragraph_sizes(shape)
    if sizes[0] > sizes[1]:
        return {"title": lines[0], "subtitle": lines[1]}
    return {"kicker": lines[0], "title": lines[1]}


def _paragraph_sizes(shape) -> list[float]:
    """The largest run in each paragraph that carries words."""
    return [
        max((run.font.size.pt for run in para.runs if run.font.size is not None), default=0.0)
        for para in shape.text_frame.paragraphs
        if para.text.strip()
    ]


def _named_field(shape) -> str | None:
    """The chrome field a shape's name claims outright — ``s1.chrome.title``."""
    match = _CHROME_NAME.match(shape.name)
    return match.group(1) or None if match else None


def _claimed(shapes: list) -> frozenset[str]:
    """The chrome fields a placeholder or a named frame holds, before any stacked frame
    is read by position."""
    fields: set[str] = set()
    for shape in shapes:
        if isinstance(shape, Unreadable) or not shape.has_text_frame:
            continue
        kind = _placeholder_type(shape)
        if kind in _TITLES:
            fields.add("title")
        elif kind is PP_PLACEHOLDER.SUBTITLE:
            fields.add("subtitle")
        elif (named := _named_field(shape)) is not None:
            fields.add(named)
    return frozenset(fields)


def _leads(bodies: list[_Text], sizes: list[float]) -> bool:
    """Whether the first body is the slide's title — a foreign deck's headline is often
    an ordinary text box, and a wrong title is worse than none."""
    if not bodies or len(bodies[0].lines) != 1 or not sizes:
        return False
    head = bodies[0].largest
    if head <= 0 or any(body.largest >= head for body in bodies[1:]):
        return False
    return head >= median(sizes) * _TITLE_RATIO


def _texts(shape, lines: Lines, taken: frozenset[str]) -> list[_Text]:
    """What claim this shape's words have on the slide's chrome fields.

    A deck deckwright built has no title placeholder: its chrome is text boxes named
    ``s1.chrome.title``, or one ``s1.chrome`` frame stacking several fields.
    """
    largest = _largest(shape)
    kind = _placeholder_type(shape)
    if kind in _TITLES:
        return [_Text(lines, largest, "title", _PLACEHOLDER)]
    if kind is PP_PLACEHOLDER.SUBTITLE:
        return [_Text(lines, largest, "subtitle", _PLACEHOLDER)]
    if (named := _named_field(shape)) is not None:
        return [_Text((" ".join(lines),), largest, named, _NAMED)]
    if _CHROME_NAME.match(shape.name) is None:
        return [_Text(lines, largest)]
    stacked = _stacked(shape, taken)
    return [_Text((text,), largest, field, _STACKED) for field, text in stacked.items()]


def _resolve(texts: list[_Text]) -> tuple[dict[str, str], list[_Text]]:
    """The slide's chrome and the words left over — a losing claimant keeps its place in
    reading order rather than vanishing."""
    fields: dict[str, str] = {}
    claimed: set[int] = set()
    for rank in (_PLACEHOLDER, _NAMED, _STACKED):
        for position, text in enumerate(texts):
            if text.rank != rank or text.field is None or text.field in fields:
                continue
            fields[text.field] = " ".join(text.lines)
            claimed.add(position)
    return fields, [t for i, t in enumerate(texts) if i not in claimed]
