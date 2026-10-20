"""A deckwright deck's placements, read back whole from the names its shapes carry.

Each shape the build drew is named for the placement that drew it, so the shapes of one
placement can be gathered and redrafted as that component — a card as a card, a chart as
a chart with its data — where a stranger's deck gives only its words.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

from deckwright.spec._bodyback import component_body, text_lines
from deckwright.spec._figures import Picture, picture
from deckwright.spec._glyphback import glyph_name
from deckwright.spec._motionback import built_spids, read_animate, read_reveals, read_transition
from deckwright.spec._origin import Origin, origin_of
from deckwright.spec._tree import (
    MAX_GROUP_DEPTH,
    flat,
    is_picture,
    linked_text,
    members,
    shape_type,
    slide_box,
)
from deckwright.utils.a11y import described, is_file_name

Box = tuple[float, float, float, float]  # x, y, w, h as fractions of the canvas

_AUTO = re.compile(r"p\d+")
_PAIRED = ("card", "panel")
# A box this far past an edge, in canvas fractions, runs off the slide rather than meeting it.
_SLACK = 0.001


@dataclass(frozen=True)
class Recovered:
    """A placement read back: its component's mapping, where it sat, and what it is called."""

    component: str
    body: dict[str, Any]
    box: Box
    key: str  # the placement's id, its pN, or its morph name
    id: str | None = None
    morph: str | None = None
    reveals: str | None = None
    note: str | None = None
    words: tuple[str, ...] = ()
    media: tuple[tuple[str, bytes], ...] = ()  # a picture's file, by the name ``body`` gives it
    painted: tuple[str, str | None] | None = None  # a plate's fill and its words' ink, as hex
    inked: str | None = None  # an icon's fill, as hex
    sizes: tuple[float | None, ...] = ()  # each line over an image, its point size
    token: str = ""  # the origin its shapes were named for, which a note names it by
    still: bool = False  # on an animated slide, none of its shapes is in the build
    bleed: bool = False  # declared to run off the slide, or found running off it


@dataclass(frozen=True)
class Recovery:
    """A slide's placements read back, the shapes they account for, and the slide's motion."""

    placed: tuple[Recovered, ...] = ()
    notes: tuple[str, ...] = ()
    claimed: frozenset[Any] = frozenset()
    animate: str | None = None
    transition: str | None = None


def recover(
    slide,
    *,
    index: int,
    canvas: tuple[int, int],
    boxes: dict[str, Box],
    icons: tuple[Path, ...] = (),
    bleeds: frozenset[str] = frozenset(),
) -> Recovery:
    """The placements ``slide``'s named shapes came from; nothing for a stranger's slide.

    ``boxes`` maps a placement's origin to the rectangle its manifest recorded. Without one,
    a placement's box is the union of its shapes'. An icon is named by its drawing, from the
    glyphs in ``icons`` and the shipped set.
    """
    animate = read_animate(slide)
    built = built_spids(slide) if animate not in (None, "none") else None
    groups: dict[str, list[Any]] = {}
    origins: dict[str, Origin] = {}
    for shape in _built(slide.shapes):
        found = origin_of(str(shape.name))
        if found is not None:
            groups.setdefault(found.token, []).append(shape)
            origins.setdefault(found.token, found)
    if not groups:
        # A slide holding only its chrome still arrives the way it was told to.
        return Recovery(transition=read_transition(slide))
    placed: list[Recovered] = []
    notes: list[str] = []
    claimed: set[Any] = set()
    for token, shapes in groups.items():
        if len({str(s.name) for s in shapes}) < len(shapes):
            continue  # a copy kept its original's name: two placements, left to the words walk
        origin = origins[token]
        name = f"slide{index}-{_safe(origin.placement)}"
        box = boxes.get(token) or _union(shapes, canvas)
        inked = None
        sizes: tuple[float | None, ...] = ()
        if origin.component == "image":
            sizes = _over_sizes(shapes)
            body, note, words, media = _image(shapes, name=name)
        elif origin.component == "icon":
            (body, inked), note, words, media = _icon(shapes, box, canvas, icons), None, (), ()
        else:
            (body, note, words), media = component_body(origin.component, shapes, icons), ()
            if body is None and not words:
                body, _, words, media = _image(shapes, name=name)
                note = (
                    f"drafted as an image: a {origin.component} placement's source is not in "
                    "the file, only its picture"
                )
        if body is None:
            continue  # left to the words walk, which names it and keeps its alt text
        if note:
            notes.append(f"{token}: {note}")
        claimed.update(shape._element for shape in shapes)
        placed.append(
            Recovered(
                component=next(iter(body)),
                body=body,
                box=box,
                key=origin.placement,
                id=_id(origin, position=len(placed) + 1),
                morph=origin.placement if origin.morph else None,
                words=words,
                media=media,
                painted=_painted(shapes) if origin.component in _PAIRED else None,
                inked=inked,
                sizes=sizes,
                token=token,
                still=built is not None and not any(int(sh.shape_id) in built for sh in shapes),
                bleed=token in bleeds or _off_canvas(box),
            )
        )
    wanted, unsaid = read_reveals(slide, keys={item.key for item in placed})
    tokens = {origin.placement: token for token, origin in origins.items()}
    notes += [f"{tokens[key]}: {why}" for key, why in unsaid]
    return Recovery(
        placed=tuple(_revealed(item, wanted) for item in placed),
        notes=tuple(notes),
        claimed=frozenset(claimed),
        animate=animate,
        transition=read_transition(slide),
    )


def _off_canvas(box: Box) -> bool:
    """Whether a box in canvas fractions runs past any edge of the slide."""
    x, y, w, h = box
    return min(x, y) < -_SLACK or max(x + w, y + h) > 1 + _SLACK


def _revealed(item: Recovered, wanted: dict[str, str]) -> Recovered:
    """``item`` with its ``reveals:``, and the id a trigger must be named by."""
    if item.key in wanted.values() and item.id is None:
        item = replace(item, id=item.key)
    if item.key in wanted:
        item = replace(item, reveals=wanted[item.key])
    return item


def _built(shapes, depth: int = 0):
    """Every shape in tree order — the build's own — with a group's children in its place."""
    built, _ = members(shapes)
    for shape in built:
        if shape_type(shape) is MSO_SHAPE_TYPE.GROUP and depth < MAX_GROUP_DEPTH:
            yield from _built(shape.shapes, depth + 1)
        else:
            yield shape


def _id(origin: Origin, *, position: int) -> str | None:
    """The ``id:`` that keeps the placement's shapes named as they were: an author's own id,
    or a ``pN`` the draft's order would no longer give it."""
    if origin.morph:
        return None
    if not _AUTO.fullmatch(origin.placement) or origin.placement != f"p{position}":
        return origin.placement
    return None


def _image(
    shapes: list[Any], *, name: str
) -> tuple[dict[str, Any] | None, str | None, tuple[str, ...], tuple[tuple[str, bytes], ...]]:
    """An image placement's picture, as the file the draft will write beside it, and the
    lines set over it."""
    found = next((s for s in shapes if is_picture(s)), None)
    got = picture(found, name=name) if found is not None else None
    if not isinstance(got, Picture):
        return None, None, (), ()
    body = got.image()
    texts = [s for s in shapes if s.has_text_frame and text_lines(s)]
    over = tuple(line for shape in texts for line in text_lines(shape))
    if over:
        body["over"] = list(over)
    scrim = any(s is not found and s not in texts for s in shapes)
    note = "its scrim is solved afresh" if scrim else None
    return {"image": body}, note, over, ((got.name, got.blob),)


def _over_sizes(shapes: list[Any]) -> tuple[float | None, ...]:
    """The point size of each line over an image, in the order ``_image`` reads them."""
    return tuple(
        next((run.font.size.pt for run in p.runs if run.font.size is not None), None)
        for shape in shapes
        if shape.has_text_frame and text_lines(shape)
        for p in shape.text_frame.paragraphs
        if flat(linked_text(p))
    )


def _painted(shapes: list[Any]) -> tuple[str, str | None] | None:
    """The plate's solid fill, and the ink of the first run on it that names a colour."""
    plate = next((s for s in shapes if shape_type(s) is MSO_SHAPE_TYPE.AUTO_SHAPE), None)
    if plate is None or plate.fill.type != MSO_FILL.SOLID:
        return None
    runs = (
        run
        for shape in shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
    )
    ink = next(
        (str(r.font.color.rgb) for r in runs if r.font.color.type == MSO_COLOR_TYPE.RGB), None
    )
    return str(plate.fill.fore_color.rgb), ink


def _safe(key: str) -> str:
    return re.sub(r"[^A-Za-z0-9_-]+", "-", key).strip("-") or "image"


def _icon(
    shapes: list[Any], box: Box, canvas: tuple[int, int], icons: tuple[Path, ...]
) -> tuple[dict[str, Any] | None, str | None]:
    """An icon placement as the glyph it draws, and the fill it was painted; nothing for a
    drawing no glyph name reaches, which the words walk names with its alt text."""
    name = glyph_name(shapes[0], icons) if len(shapes) == 1 else None
    if name is None:
        return None, None
    shape = shapes[0]
    fields: dict[str, Any] = {"name": name}
    drawn, short = slide_box(shape), min(box[2] * canvas[0], box[3] * canvas[1])
    size = round(drawn[2] / short, 2) if drawn and short else 1.0
    if size != 1.0:
        fields["size"] = size
    alt, decorative = described(shape)
    if decorative:
        fields["decorative"] = True
    elif alt and not is_file_name(alt):
        fields["alt"] = alt
    fill = str(shape.fill.fore_color.rgb) if shape.fill.type == MSO_FILL.SOLID else None
    return {"icon": fields}, fill


def _union(shapes: list[Any], canvas: tuple[int, int]) -> Box:
    width, height = canvas
    edges = [
        (box[0], box[1], box[0] + box[2], box[1] + box[3])
        for box in map(slide_box, shapes)
        if box is not None
    ]
    left = min(e[0] for e in edges)
    top = min(e[1] for e in edges)
    return (
        left / width,
        top / height,
        (max(e[2] for e in edges) - left) / width,
        (max(e[3] for e in edges) - top) / height,
    )
