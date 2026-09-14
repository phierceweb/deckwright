"""Walk a slide's shape tree into leaves, in the order a human reads them.

An element the tree holds that is not a shape at all is named rather than walked, so a
deck carrying one is still harvested.
"""

from __future__ import annotations

from dataclasses import dataclass

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from deckwright.utils.links import LINK, spans

# A group nested deeper than this is named in `dropped` rather than walked, so a
# pathological file cannot exhaust the stack.
MAX_GROUP_DEPTH = 8

# Two shapes read as one row while each overlaps the other by more than this much of
# its own height. A fraction, so it holds inside a group's own child space too.
_ROW_OVERLAP = 0.35

# The children of a `p:spTree` or `p:grpSp` python-pptx has a shape proxy for. Any
# other content tag — `p:contentPart` for ink, `mc:AlternateContent` wrapping a 3D
# model or a newer-namespace effect — a slide's factory raises on and a group's hands
# back as a bare shape with no geometry to read.
_SHAPE_TAGS = frozenset(
    qn(tag) for tag in ("p:sp", "p:grpSp", "p:graphicFrame", "p:cxnSp", "p:pic")
)

# Children of a `p:spTree` or `p:grpSp` that were never shapes to begin with.
_SCAFFOLDING = frozenset({qn("p:nvGrpSpPr"), qn("p:grpSpPr"), qn("p:extLst")})


@dataclass(frozen=True)
class Unreadable:
    """An element of the shape tree python-pptx will not hand over as a shape."""

    tag: str


def tag(element) -> str:
    local = element.tag.rpartition("}")[2]
    return f"{element.prefix}:{local}" if element.prefix else local


def shape_type(shape) -> MSO_SHAPE_TYPE | None:
    """``shape_type`` is ``None`` for SmartArt and raises for a ``<p:sp>`` python-pptx
    cannot classify."""
    try:
        return shape.shape_type
    except NotImplementedError:
        return None


def leaves(shapes, depth: int = 0):
    """Every leaf shape, in reading order, and every element that is not one.

    Only shapes carrying words are banded into rows; art follows them in its own
    top-then-left order, so a tall decoration cannot split the row its caption reads in.
    A group's children report offsets in its own child space, which is not comparable
    with slide space — so they are ordered among themselves and the run takes the
    group's place. A group that yields nothing, and one past :data:`MAX_GROUP_DEPTH`,
    are yielded whole, which lands them in ``dropped`` under their own name.
    """
    built, unreadable = members(shapes)
    worded: list = []
    art: list = []
    for shape in built:
        (worded if carries_words(shape, depth) else art).append(shape)
    ordered = _reading_order(worded) + sorted(art, key=lambda s: (s.top or 0, s.left or 0))
    for shape in ordered:
        if shape_type(shape) is not MSO_SHAPE_TYPE.GROUP or depth >= MAX_GROUP_DEPTH:
            yield shape
            continue
        yield from list(leaves(shape.shapes, depth + 1)) or [shape]
    yield from unreadable


def carries_words(shape, depth: int = 0) -> bool:
    """Whether the shape, or any child of a group, holds text or a table."""
    if shape_type(shape) is MSO_SHAPE_TYPE.GROUP:
        if depth >= MAX_GROUP_DEPTH:
            return False
        built, _ = members(shape.shapes)
        return any(carries_words(child, depth + 1) for child in built)
    if shape.has_table:
        return True
    return bool(shape.has_text_frame and shape.text_frame.text.strip())


def members(shapes) -> tuple[list, list[Unreadable]]:
    """The shapes python-pptx can build and place, and the elements it cannot.

    Iterating the tree itself is not safe: a group hands back anything it does not
    recognise as a bare shape with no geometry, and a slide raises instead. Both are
    still the author's slide, so an unreadable element is named rather than ending
    the harvest.
    """
    built: list = []
    unreadable: list[Unreadable] = []
    for element in shapes._element.iterchildren():
        if not isinstance(element.tag, str) or element.tag in _SCAFFOLDING:
            continue
        shape = _build(shapes, element) if element.tag in _SHAPE_TAGS else None
        if shape is None:
            unreadable.append(Unreadable(tag(element)))
        else:
            built.append(shape)
    return built, unreadable


def _build(shapes, element):
    """The shape ``element`` reads back as, or ``None`` when it carries no geometry."""
    try:
        shape = shapes._shape_factory(element)
        _ = (shape.top, shape.left, shape.height)
    except Exception:
        return None
    return shape


def _reading_order(shapes: list) -> list:
    return [shape for row in _rows(shapes) for shape in sorted(row, key=lambda s: s.left or 0)]


def _rows(shapes: list) -> list[list]:
    """The shapes banded into visual rows, top to bottom.

    A row is a set of shapes that all share a vertical band with one another. Sorting on
    an exact ``top`` instead reverses a row a human reads left to right, because shapes
    laid out as one band sit anywhere from thousandths of an inch to a third of a line
    apart. A grouped or layout-inherited shape can also report no offset at all.
    """
    rows: list[list] = []
    for shape in sorted(shapes, key=lambda s: (s.top or 0, s.left or 0)):
        if rows and all(_one_band(member, shape) for member in rows[-1]):
            rows[-1].append(shape)
        else:
            rows.append([shape])
    return rows


def _one_band(first, other) -> bool:
    """Whether two shapes overlap by enough of *each* height to read as one row.

    Both heights have to be covered, so a caption does not join the tall panel it
    merely starts beside.
    """
    heights = (first.height or 0, other.height or 0)
    overlap = min(_bottom(first), _bottom(other)) - max(first.top or 0, other.top or 0)
    return all(overlap > _ROW_OVERLAP * height for height in heights)


def _bottom(shape) -> int:
    return (shape.top or 0) + (shape.height or 0)


SOFT_BREAK = "\x0b"  # what python-pptx hands back for an ``<a:br/>`` inside a paragraph


def linked_text(paragraph) -> str:
    """A paragraph's text as a spec writes it: a hyperlinked run as ``[words](address)``,
    and copy that would read as a link without being one with its bracket escaped.

    A hyperlink whose markup would not parse back is written as its words alone.
    """
    parts = []
    runs = iter(paragraph.runs)  # one proxy per a:r, in document order
    for child in paragraph._p.iterchildren():
        tag = child.tag.rpartition("}")[2]
        if tag == "br":
            parts.append(SOFT_BREAK)
        elif tag in ("r", "fld"):
            run = next(runs) if tag == "r" else None
            text = "".join(t.text or "" for t in child.iter(qn("a:t")))
            address = run.hyperlink.address if run is not None else None
            markup = f"[{text}]({address})"
            if address and text.strip() and spans(markup) == [(text, address)]:
                parts.append(markup)
            else:
                parts.append(_escaped(text))
    return "".join(parts)


def _escaped(text: str) -> str:
    """``text`` with a backslash before each bracket that would open a link."""
    out, at = [], 0
    for match in LINK.finditer(text):
        out.append(text[at : match.start()] + "\\[")
        at = match.start() + 1
    return "".join(out) + text[at:]


def flat(text: str) -> str:
    """A paragraph's words on one line: a soft break is a wrap, not a new item, and
    python-pptx writes the character it becomes back out as literal ``_x000B_``.
    Nothing else moves — the author's own spacing is content."""
    return text.replace(SOFT_BREAK, " ").strip()
