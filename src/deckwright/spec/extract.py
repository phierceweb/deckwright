"""Read a deck nobody authored here into the words a draft spec can be written from.

Deliberately lossy: chrome, bullets, tables and notes convert; everything else is
named in :attr:`SlideContent.dropped` so the author knows what to put back by hand.
"""

from __future__ import annotations

from dataclasses import dataclass
import re
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from pf_core.log import get_logger

from deckwright.errors import SpecError
from deckwright.spec._chrome import (
    _Text,
    _claimed,
    _leads,
    _placeholder_type,
    _resolve,
    _run_sizes,
    _texts,
)
from deckwright.spec._tree import (
    MAX_GROUP_DEPTH,
    SOFT_BREAK,
    Unreadable,
    flat,
    leaves,
    linked_text,
    members,
    shape_type,
)
from deckwright.spec.draft import render_markdown, render_spec  # noqa: F401 — re-exported
from deckwright.utils.a11y import described, is_file_name
from deckwright.utils.deck import open_presentation

__all__ = ["MAX_GROUP_DEPTH", "SlideContent", "harvest", "render_markdown", "render_spec"]

logger = get_logger(__name__)

_Lines = tuple[str, ...]
_Table = tuple[_Lines, ...]

# `components/bullets.py` writes its dot as run text, so a deckwright deck re-ingests it
# unless it is taken off again here. Scoped to that component's own shape names: a
# foreign deck's leading dot is content.
_BULLETS_NAME = re.compile(r"^s\d+\.[^#]+\.bullets#\d+$")

_BULLET_MARKER = "•  "

# A background covers the canvas to within this; the slack absorbs a deck authored in
# inches rather than EMU.
_EDGE = 12700  # one point


@dataclass(frozen=True)
class SlideContent:
    """One slide's words, and what could not be turned into words."""

    index: int
    title: str | None = None
    kicker: str | None = None
    subtitle: str | None = None
    blocks: tuple[_Lines, ...] = ()
    tables: tuple[_Table, ...] = ()
    notes: str | None = None
    dropped: _Lines = ()
    alt: _Lines = ()  # a dropped figure's alternative text, as "picture 'Logo': the logo"
    background_rgb: str | None = None  # what the slide or a full-bleed shape paints it


def _kind(shape) -> str:
    """What to call a shape in ``dropped``."""
    kind = shape_type(shape)
    return kind.name.lower() if kind is not None else type(shape).__name__.lower()


def _nothing_was_placed(shape) -> bool:
    """A layout placeholder nobody typed into, or an empty text box. Every other
    empty shape was drawn on purpose — an arrow, a divider, an icon."""
    return _placeholder_type(shape) is not None or shape_type(shape) is MSO_SHAPE_TYPE.TEXT_BOX


def _lines(shape) -> _Lines:
    lines = (flat(linked_text(p)) for p in shape.text_frame.paragraphs)
    if _BULLETS_NAME.match(shape.name):
        lines = (line.removeprefix(_BULLET_MARKER) for line in lines)
    return tuple(line for line in lines if line)


def _grid(shape) -> _Table:
    return tuple(tuple(_cell(cell) for cell in row.cells) for row in shape.table.rows)


def _cell(cell) -> str:
    return flat("\n".join(linked_text(p) for p in cell.text_frame.paragraphs))


def _notes(slide) -> str | None:
    if not slide.has_notes_slide:
        return None
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        return None
    return frame.text.replace(SOFT_BREAK, "\n").strip() or None


def _slide_content(slide, index: int, canvas: tuple[int, int]) -> SlideContent:
    texts: list[_Text] = []
    sizes: list[float] = []
    tables: list[_Table] = []
    dropped: list[str] = []
    alt: list[str] = []
    background, covered = _backdrop(slide, canvas)
    shapes = list(leaves(slide.shapes))
    taken = _claimed(shapes)
    for shape in shapes:
        if isinstance(shape, Unreadable):
            dropped.append(f"{shape.tag} — an element python-pptx cannot read")
            continue
        if shape._element in covered:
            continue  # the background, or a full-bleed panel painted over by it
        if shape.has_table:
            tables.append(_grid(shape))
            continue
        if not shape.has_text_frame:
            dropped.append(f"{_kind(shape)} {shape.name!r}")
            text, _ = described(shape)
            if text and not is_file_name(text):
                alt.append(f"{_kind(shape)} {shape.name!r}: {' '.join(text.split())}")
            continue
        lines = _lines(shape)
        if not lines:
            if not _nothing_was_placed(shape):
                dropped.append(f"{_kind(shape)} {shape.name!r}")
                text, _ = described(shape)
                if text and not is_file_name(text):
                    alt.append(f"{_kind(shape)} {shape.name!r}: {' '.join(text.split())}")
            continue
        sizes += _run_sizes(shape)
        texts += _texts(shape, lines, taken)
    fields, bodies = _resolve(texts)
    if "title" not in fields and _leads(bodies, sizes):
        fields["title"] = bodies.pop(0).lines[0]
    return SlideContent(
        index=index,
        title=fields.get("title"),
        kicker=fields.get("kicker"),
        subtitle=fields.get("subtitle"),
        blocks=tuple(body.lines for body in bodies),
        tables=tuple(tables),
        notes=_notes(slide),
        dropped=tuple(dropped),
        alt=tuple(alt),
        background_rgb=background,
    )


def _backdrop(slide, canvas: tuple[int, int]) -> tuple[str | None, list]:
    """The colour the slide shows behind its words, and the shapes that are only that.

    The slide's own ``p:bg`` is read first, then every top-level full-bleed solid in
    z order, so the one on top wins. A full-bleed panel that carries words is still
    the backdrop, but its words are harvested like any other shape's; one that carries
    none is the background itself, or lies under it, and is left out of the walk.
    """
    own = slide._element.cSld.find(qn("p:bg"))
    background = _solid_srgb(own.find(qn("p:bgPr")) if own is not None else None)
    covered: list = []
    built, _ = members(slide.shapes)
    for shape in built:
        if not _covers(shape, canvas):
            continue
        rgb = _solid_srgb(shape._element.find(qn("p:spPr")))
        if rgb is None:
            continue
        background = rgb
        if not (shape.has_text_frame and _lines(shape)):
            covered.append(shape._element)
    return background, covered


def _covers(shape, canvas: tuple[int, int]) -> bool:
    """Whether the shape paints the whole canvas — a designer's bleed past the edges
    and a half-turn both do."""
    width, height = canvas
    box = (shape.left, shape.top, shape.width, shape.height)
    if any(v is None for v in box):
        return False
    left, top, wide, high = box
    if left > _EDGE or top > _EDGE:
        return False
    if left + wide < width - _EDGE or top + high < height - _EDGE:
        return False
    return (getattr(shape, "rotation", 0) or 0) % 180 == 0


def _solid_srgb(properties) -> str | None:
    """The explicit RGB an ``a:solidFill`` under ``properties`` paints, or None.

    A gradient, a scheme colour and a picture are art the author has to put back by
    hand; an RGB carrying a modifier — alpha, lumMod, tint — renders as some other
    colour, so it is art too.
    """
    if properties is None:
        return None
    fill = properties.find(qn("a:solidFill"))
    colour = fill.find(qn("a:srgbClr")) if fill is not None else None
    if colour is None or len(colour):
        return None
    return colour.get("val", "").upper() or None


def harvest(deck: str | Path) -> list[SlideContent]:
    """Read every slide's words, in reading order.

    Raises:
        SpecError: ``deck`` is not a readable ``.pptx``.
    """
    prs = open_presentation(deck, what="deck", error=SpecError)
    canvas = (prs.slide_width, prs.slide_height)
    out = [_slide_content(slide, index, canvas) for index, slide in enumerate(prs.slides, start=1)]
    logger.info("deck_harvested", deck=str(deck), slides=len(out))
    return out
