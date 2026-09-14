"""Which language and face each CJK run is marked with. A corpus build writes both, but reads
neither back, and never reaches the refusal."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from deckwright.compile.eastasian import mark_east_asian, script_of
from deckwright.errors import SpecError

_EA = {
    "ja": "Hiragino Sans",
    "ko": "Apple SD Gothic Neo",
    "zh-Hans": "PingFang SC",
    "zh-Hant": "PingFang TC",
}


def _slide(*texts):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    runs = []
    for i, text in enumerate(texts):
        frame = slide.shapes.add_textbox(Inches(1), Inches(1 + i), Inches(6), Inches(1)).text_frame
        run = frame.paragraphs[0].add_run()
        run.text = text
        run.font.name = "Helvetica"
        runs.append(run)
    return slide, runs


def _marks(run):
    rpr = run._r.rPr
    ea = rpr.find(qn("a:ea"))
    return rpr.get("lang"), None if ea is None else ea.get("typeface")


def test_kana_is_japanese_and_hangul_is_korean_whatever_the_deck_says():
    slide, (ja, ko, latin) = _slide("売上は伸びました", "매출이 늘었습니다", "Revenue grew")
    mark_east_asian(slide, ea=_EA, lang="zh-Hans", where="slide 1")
    assert _marks(ja) == ("ja-JP", "Hiragino Sans")
    assert _marks(ko) == ("ko-KR", "Apple SD Gothic Neo")
    assert _marks(latin) == (None, None)


def test_han_alone_takes_the_scripts_face_the_lang_names():
    slide, (traditional,) = _slide("季度營收")
    mark_east_asian(slide, ea=_EA, lang="zh-HK", where="slide 1")
    assert _marks(traditional) == ("zh-HK", "PingFang TC")


def test_han_alone_with_no_lang_is_refused_when_the_theme_sets_it_several_ways():
    slide, _ = _slide("季度收入")
    with pytest.raises(
        SpecError, match=r"slide 1: '季度收入' carries Han characters and no kana or hangul"
    ):
        mark_east_asian(slide, ea=_EA, lang=None, where="slide 1")


def test_han_alone_with_no_lang_takes_the_one_han_face_a_theme_names():
    slide, (run,) = _slide("季度收入")
    mark_east_asian(slide, ea={"zh-Hans": "PingFang SC", "ko": "X"}, lang=None, where="s")
    assert _marks(run) == ("zh-CN", "PingFang SC")


def test_a_theme_with_no_cjk_faces_still_marks_the_language():
    slide, (run,) = _slide("売上は伸びました")
    mark_east_asian(slide, ea={}, lang=None, where="s")
    assert _marks(run) == ("ja-JP", None)


def test_the_face_sits_after_latin_and_before_a_link_in_the_run_properties():
    slide, (run,) = _slide("売上は伸びました")
    run.hyperlink.address = "https://example.com"
    mark_east_asian(slide, ea=_EA, lang=None, where="s")
    mark_east_asian(slide, ea=_EA, lang=None, where="s")
    tags = [el.tag.split("}")[1] for el in run._r.rPr]
    assert tags == ["latin", "ea", "hlinkClick"]


@pytest.mark.parametrize(
    "tag, script",
    [
        ("ja", "ja"),
        ("ja-JP", "ja"),
        ("ko-KR", "ko"),
        ("zh", "zh-Hans"),
        ("zh-CN", "zh-Hans"),
        ("zh-Hant", "zh-Hant"),
        ("zh-HK", "zh-Hant"),
        ("en", None),
        (None, None),
    ],
)
def test_a_language_tag_names_its_script(tag, script):
    assert script_of(tag) == script


def test_a_build_marks_each_slides_cjk_runs_with_the_slides_language_over_the_decks(project):
    from deckwright.compile import build_deck

    (project / "c.deck.yaml").write_text(
        "theme: testtheme\nlang: ja\nout: out/C.pptx\n---\ntitle: 季度收入\n"
        "---\nlang: zh-Hant\ntitle: 季度營收\n"
    )
    built = build_deck(project / "c.deck.yaml", theme_path=project / "testtheme.yaml")
    prs = Presentation(str(built.deck))
    runs = [
        next(
            r
            for s in slide.shapes
            if s.has_text_frame
            for p in s.text_frame.paragraphs
            for r in p.runs
        )
        for slide in prs.slides
    ]
    assert [r._r.rPr.get("lang") for r in runs] == ["ja", "zh-Hant"]
    assert runs[0]._r.rPr.find(qn("a:ea")).get("typeface") == "ＭＳ Ｐゴシック"


_TEMPLATE_EA = {"ja": "游ゴシック", "ko": "맑은 고딕", "zh-Hans": "等线", "zh-Hant": "新細明體"}


@pytest.mark.parametrize("text", ["Revenue up １２％", "Delivery in 3〜5 days"])
def test_fullwidth_forms_and_cjk_punctuation_alone_decide_nothing(text):
    slide, (run,) = _slide(text)
    mark_east_asian(slide, ea=_TEMPLATE_EA, lang="en", where="s")
    mark_east_asian(slide, ea=_TEMPLATE_EA, lang=None, where="s")
    assert _marks(run) == (None, None)


def test_halfwidth_katakana_is_japanese():
    slide, (run,) = _slide("東京ｽｶｲﾂﾘｰ")
    mark_east_asian(slide, ea=_TEMPLATE_EA, lang=None, where="s")
    assert _marks(run) == ("ja-JP", "游ゴシック")


def test_a_declared_korean_deck_sets_hanja_in_the_korean_face():
    slide, (run,) = _slide("大韓民國")
    mark_east_asian(slide, ea=_TEMPLATE_EA, lang="ko", where="s")
    assert _marks(run) == ("ko", "맑은 고딕")


@pytest.mark.parametrize(
    "tag, face", [("yue-Hant-HK", "新細明體"), ("cmn-Hans-CN", "等线"), ("zh_TW", "新細明體")]
)
def test_a_chinese_tag_is_read_by_its_script_and_region_subtags(tag, face):
    slide, (run,) = _slide("季度收入")
    mark_east_asian(slide, ea=_TEMPLATE_EA, lang=tag, where="s")
    assert _marks(run)[1] == face


def test_a_language_that_only_starts_like_japanese_is_not_japanese():
    assert script_of("jav") is None
    assert script_of("kok") is None


def test_han_alone_is_not_refused_when_every_han_script_shares_one_face():
    slide, (run,) = _slide("季度收入")
    same = {"ja": "Source Han Sans", "zh-Hans": "Source Han Sans", "zh-Hant": "Source Han Sans"}
    mark_east_asian(slide, ea=same, lang=None, where="s")
    assert _marks(run) == (None, "Source Han Sans")


def test_a_paragraph_is_judged_whole_so_a_han_link_in_japanese_copy_is_japanese():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paragraph = slide.shapes.add_textbox(0, 0, Inches(6), Inches(1)).text_frame.paragraphs[0]
    link, rest = paragraph.add_run(), paragraph.add_run()
    link.text, rest.text = "東京本社", "のご案内"
    mark_east_asian(slide, ea=_TEMPLATE_EA, lang=None, where="s")
    assert _marks(link) == ("ja-JP", "游ゴシック")
