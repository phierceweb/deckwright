"""`goto:` written as click actions. A corpus build writes them but reads none back."""

from __future__ import annotations

import json

from pptx import Presentation
from pptx.enum.action import PP_ACTION

from deckwright.compile import build_deck

_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def _build(project, slides: str):
    (project / "g.deck.yaml").write_text("theme: testtheme\nout: out/G.pptx\n" + slides)
    return build_deck(project / "g.deck.yaml", theme_path=project / "testtheme.yaml")


_DECK = """---
title: Agenda
place:
  - at: {cols: left-half}
    goto: appendix
    card: {heading: Appendix, body: The numbers}
  - at: {cols: right-half}
    goto: next
    card: {heading: Onward}
---
title: Middle
---
id: appendix
title: Appendix
"""


def test_every_shape_a_goto_placement_drew_jumps_to_the_named_slide(project):
    built = _build(project, _DECK)
    prs = Presentation(str(built.deck))
    first, _, appendix = prs.slides
    card = [s for s in first.shapes if s.name.startswith("s1.p1.card")]
    assert len(card) == 2
    for shape in card:
        assert shape.click_action.action == PP_ACTION.NAMED_SLIDE
        assert shape.click_action.target_slide.slide_id == appendix.slide_id


def test_a_relative_jump_is_a_show_jump_that_relates_to_nothing(project):
    built = _build(project, _DECK)
    prs = Presentation(str(built.deck))
    onward = [s for s in prs.slides[0].shapes if s.name.startswith("s1.p2.card")]
    for shape in onward:
        hlink = shape._element._nvXxPr.cNvPr.hlinkClick
        assert hlink.get("action") == "ppaction://hlinkshowjump?jump=nextslide"
        assert hlink.get(_R) == ""


def test_the_chrome_and_other_placements_do_not_jump(project):
    built = _build(project, _DECK)
    prs = Presentation(str(built.deck))
    unlinked = [s.name for s in prs.slides[0].shapes if s.click_action.action == PP_ACTION.NONE]
    assert unlinked == ["s1.chrome.title"]


def test_the_manifest_records_where_each_shape_jumps(project):
    built = _build(project, _DECK)
    shapes = json.loads(built.manifest.read_text())["slides"][0]["shapes"]
    assert {s["name"]: s.get("goto") for s in shapes} == {
        "s1.p1.card#1": "appendix",
        "s1.p1.card#2": "appendix",
        "s1.p2.card#1": "next",
        "s1.p2.card#2": "next",
        "s1.chrome.title": None,
    }
