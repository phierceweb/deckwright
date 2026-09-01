"""A slide's layout must not depend on the slides around it.

Issue 25: a `callouts` placement with `anchor: middle`, on a slide nobody edited, lost
its centring when two *other* slides were merged, and regained it when a *third* moved.
Same document, same manifest box, different offset inside it — and `qa` clean in both
states, because every check it runs reads one slide at a time.

The original snapshots are gone, so this locks in the property rather than reproducing
the report: whatever a slide draws, it draws the same whatever precedes it.

Issue 6 is the same defect seen from inside one slide: a later placement's settle swept up
an earlier one and dragged it off the anchor point.
"""

from __future__ import annotations

import json
import textwrap

import pytest

from deckwright.compile.build import build_deck

_TARGET = """
    title: The untouched slide
    place:
      - at: {cols: full, rows: {from: 3, to: 11}}
        anchor: middle
        callouts:
          items:
            - {head: The fix, body: "A body long enough that it wraps inside its column."}
            - {head: The cost, body: "A second body, also long enough to wrap in place."}
            - {head: The risk, body: "And a third, likewise long enough to wrap here."}
"""

_EARLIER_APART = """
    title: An earlier slide
    place:
      - at: {cols: full, rows: {from: 3, to: 8}}
        bullets: {items: [One, Two, Three]}
    ---
    title: Another earlier slide
    place:
      - at: {cols: full, rows: {from: 3, to: 8}}
        bullets: {items: [Four, Five]}
"""

_EARLIER_MERGED = """
    title: The two earlier slides, merged
    place:
      - at: {cols: {from: 0, to: 6}, rows: {from: 3, to: 8}}
        bullets: {items: [One, Two, Three]}
      - at: {cols: {from: 6, to: 12}, rows: {from: 3, to: 8}}
        bullets: {items: [Four, Five]}
"""


def _tops(project, name: str, earlier: str, out: str) -> list[dict]:
    docs = [f"theme: testtheme\ntitle: T\nout: out/{out}.pptx"]
    if earlier.strip():
        docs += [d for d in textwrap.dedent(earlier).split("---") if d.strip()]
    docs.append(textwrap.dedent(_TARGET))
    spec = project / f"{name}.deck.yaml"
    spec.write_text("\n---\n".join(d.strip("\n") for d in docs) + "\n")
    build_deck(spec, theme_path=project / "testtheme.yaml")
    data = json.loads((project / "out" / f"{out}.manifest.json").read_text())
    return [s["box"] for s in data["slides"][-1]["shapes"]]


def test_merging_earlier_slides_does_not_move_a_later_one(project):
    apart = _tops(project, "apart", _EARLIER_APART, "Apart")
    merged = _tops(project, "merged", _EARLIER_MERGED, "Merged")
    assert apart, "the target slide drew nothing, so this proves nothing"
    assert apart == merged


def test_a_slide_lays_out_the_same_with_nothing_before_it(project):
    alone = _tops(project, "alone", "\n", "Alone")
    preceded = _tops(project, "preceded", _EARLIER_APART, "Preceded")
    assert alone, "the target slide drew nothing, so this proves nothing"
    assert alone == preceded


def test_pruning_keeps_the_layout_every_slide_actually_uses(project):
    """The same lxml-proxy hazard as issue 25, one part up: `prune` compared layouts by
    `id()` of a dropped proxy, so a recycled address could read as a match. A false
    positive keeps a layout nobody wants; a false negative removes one a slide is on."""
    from pptx import Presentation

    spec = project / "d.deck.yaml"
    build_deck(spec, theme_path=project / "testtheme.yaml")
    prs = Presentation(str(project / "out" / "Demo.pptx"))
    kept = {
        master_layout.part.partname for m in prs.slide_masters for master_layout in m.slide_layouts
    }
    for slide in prs.slides:
        assert slide.slide_layout.part.partname in kept


_THREE_IN_A_BAND = """
    title: Three placements sharing one row band
    place:
      - at: {cols: left-third, rows: {from: 1, to: 9}}
        anchor: middle
        prose: {paragraphs: [First.]}
      - at: {cols: mid-third, rows: {from: 1, to: 9}}
        anchor: middle
        prose: {paragraphs: [Second.]}
      - at: {cols: right-third, rows: {from: 1, to: 9}}
        anchor: middle
        prose: {paragraphs: [Third.]}
"""


def test_anchor_middle_centres_every_placement_in_its_own_row_band(project):
    """Issue 6, which is issue 25 seen from inside one slide: the author computed that a
    `rows: {from: 1, to: 9}` band should centre its content near 52% of the canvas and the
    render put it at 83-90%, so row bands were abandoned for hand-written `box:` coordinates.

    Each placement settles against the band the manifest recorded for it, so the expectation
    is read off that box rather than restated here.
    """
    from pptx import Presentation

    spec = project / "band.deck.yaml"
    spec.write_text(
        "theme: testtheme\ntitle: T\nout: out/Band.pptx\n---\n"
        + textwrap.dedent(_THREE_IN_A_BAND).strip("\n")
        + "\n"
    )
    build_deck(spec, theme_path=project / "testtheme.yaml")
    slide = json.loads((project / "out" / "Band.manifest.json").read_text())["slides"][-1]
    drawn = {
        s.name: (s.top + s.height / 2) / 914400
        for s in Presentation(str(project / "out" / "Band.pptx")).slides[-1].shapes
    }

    settled = [s for s in slide["shapes"] if ".prose" in s["name"]]
    assert len(settled) == 3, [s["name"] for s in settled]
    for record, placement in zip(settled, slide["placements"], strict=True):
        band = placement["box"]
        middle = band["y"] + band["h"] / 2
        assert drawn[record["name"]] == pytest.approx(middle, abs=0.02), (
            f"{record['name']} sits at {drawn[record['name']]:.2f}in in a band "
            f"centred on {middle:.2f}in"
        )
        assert record["box"]["y"] + record["box"]["h"] / 2 == pytest.approx(
            drawn[record["name"]], abs=0.02
        ), f"the manifest puts {record['name']} where it is not — qa reads this box"
