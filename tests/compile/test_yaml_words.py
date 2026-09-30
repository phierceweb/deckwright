"""`yes`, `no`, `on` and `off` in a built deck: copy prints them as written, and a field
that takes true or false reads the truth they spell."""

import json
import textwrap

import pytest
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

from deckwright.compile import build_deck
from deckwright.errors import LayoutError


def _build(tmp_path, monkeypatch, *slides: str):
    tmp_path.mkdir(exist_ok=True)
    monkeypatch.setenv("DECKWRIGHT_THEME_DIR", str(tmp_path / "no-theme-dir"))
    Image.new("RGB", (400, 300), (200, 180, 160)).save(tmp_path / "pic.png")
    body = "".join("---\n" + textwrap.dedent(s) for s in slides)
    spec = tmp_path / "w.deck.yaml"
    spec.write_text("theme: base\nout: W.pptx\n" + body)
    result = build_deck(spec)
    return Presentation(str(result.deck)), json.loads(result.manifest.read_text())


def _shapes(deck, slide: int, prefix: str):
    return [s for s in deck.slides[slide - 1].shapes if s.name.startswith(prefix)]


def _table(deck):
    return next(s.table for s in deck.slides[0].shapes if s.has_table)


def _chart(deck):
    return next(s.chart for s in deck.slides[0].shapes if s.has_chart)


def test_copy_written_as_a_boolean_word_prints_as_written(tmp_path, monkeypatch):
    deck, _ = _build(
        tmp_path,
        monkeypatch,
        """
        title: Words
        place:
          - at: {cols: left-half}
            table: {rows: [[yes, no], [On, OFF]]}
          - at: {cols: right-half}
            bullets: {items: [yes, no, on, off]}
        """,
        """
        title: Chart
        place:
          - at: {cols: full}
            chart:
              kind: column
              data: [{category: yes, value: 1}, {category: No, value: 2}, {category: off, value: 3}]
        """,
    )
    assert [[c.text for c in r.cells] for r in _table(deck).rows] == [["yes", "no"], ["On", "OFF"]]
    (bullets,) = _shapes(deck, 1, "s1.p2.bullets")
    assert bullets.text_frame.text == "•  yes\n•  no\n•  on\n•  off"
    chart = next(s.chart for s in deck.slides[1].shapes if s.has_chart)
    assert list(chart.plots[0].categories) == ["yes", "No", "off"]


@pytest.mark.parametrize(("word", "banded"), [("yes", True), ("no", False)])
def test_banding_written_as_a_word_takes_the_truth_it_spells(tmp_path, monkeypatch, word, banded):
    deck, _ = _build(
        tmp_path,
        monkeypatch,
        f"""
        title: T
        place:
          - at: {{cols: full}}
            table: {{banding: {word}, header: [A, B], rows: [[a, b], [c, d], [e, f]]}}
        """,
    )
    filled = [
        row.cells[0]._tc.find(f"{qn('a:tcPr')}/{qn('a:solidFill')}") is not None
        for row in _table(deck).rows
    ]
    assert filled == [True, False, banded, False]


@pytest.mark.parametrize(("word", "bold"), [("yes", True), ("no", False)])
def test_a_cell_emphasised_with_a_word_takes_the_truth_it_spells(tmp_path, monkeypatch, word, bold):
    deck, _ = _build(
        tmp_path,
        monkeypatch,
        f"""
        title: T
        place:
          - at: {{cols: full}}
            table: {{rows: [[{{text: a, emphasis: {word}}}, b]]}}
        """,
    )
    cells = _table(deck).rows[0].cells
    assert [c.text_frame.paragraphs[0].runs[0].font.bold is True for c in cells] == [bold, False]


def _point_fills(chart) -> list[str]:
    return [
        dpt.find(f".//{qn('a:srgbClr')}").get("val") for dpt in chart._chartSpace.iter(qn("c:dPt"))
    ]


@pytest.mark.parametrize("word", ["yes", "no"])
def test_a_chart_row_highlighted_with_a_word_takes_the_truth_it_spells(tmp_path, monkeypatch, word):
    deck, _ = _build(
        tmp_path,
        monkeypatch,
        f"""
        title: T
        place:
          - at: {{cols: full}}
            chart:
              kind: column
              data: [{{category: A, value: 1}}, {{category: B, value: 2, highlight: {word}}}]
        """,
    )
    fills = _point_fills(_chart(deck))
    if word == "yes":
        assert len(fills) == 2 and fills[0] != fills[1]
    else:
        assert len(set(fills)) <= 1


@pytest.mark.parametrize(("word", "hidden"), [("no", "0"), ("yes", None)])
def test_series_labels_turned_off_with_a_word(tmp_path, monkeypatch, word, hidden):
    deck, _ = _build(
        tmp_path,
        monkeypatch,
        f"""
        title: T
        place:
          - at: {{cols: full}}
            chart:
              kind: column
              labels: {{Ours: {word}}}
              data:
                - {{category: A, values: {{Ours: 1, Theirs: 2}}}}
                - {{category: B, values: {{Ours: 3, Theirs: 4}}}}
        """,
    )
    ours = next(
        ser
        for ser in _chart(deck)._chartSpace.iter(qn("c:ser"))
        if ser.find(f".//{qn('c:tx')}//{qn('c:v')}").text == "Ours"
    )
    shown = ours.find(f"{qn('c:dLbls')}/{qn('c:showVal')}")
    assert (None if shown is None else shown.get("val")) == hidden


@pytest.mark.parametrize(("word", "drawn"), [("yes", 3), ("no", 2)])
def test_an_image_scrim_written_as_a_word(tmp_path, monkeypatch, word, drawn):
    """Over text a scrim is on unless declared off: picture, scrim, text — or no scrim."""
    deck, _ = _build(
        tmp_path,
        monkeypatch,
        f"""
        title: T
        place:
          - at: {{cols: full}}
            image: {{src: pic.png, alt: A picture, scrim: {word}, over: [Caption]}}
        """,
    )
    assert len(_shapes(deck, 1, "s1.p1.image")) == drawn


def test_a_versus_side_highlighted_with_a_word(tmp_path, monkeypatch):
    deck, _ = _build(
        tmp_path,
        monkeypatch,
        """
        title: T
        place:
          - at: {cols: full}
            versus:
              left: {value: "1", label: Before, highlight: no}
              right: {value: "2", label: After, highlight: yes}
        """,
    )
    plates = [s for s in _shapes(deck, 1, "s1.p1.versus") if not s.has_text_frame or not s.text]
    left, right = (p.fill.fore_color.rgb for p in plates[:2])
    assert (str(left), str(right)) == ("5F6672", "1F5FA8")


# Each written with a word, then with an ordinary word in its place: the copy must print,
# be recorded for `qa`, and be measured, so the two lay out alike.
_OPTIONAL_COPY = [
    ("bullets", "anchor: middle\n    bullets: {{heading: {w}, items: [a]}}", "{w}"),
    ("stats label", "anchor: middle\n    stats: {{items: [{{value: '1', label: {w}}}]}}", "1\n{w}"),
    (
        "stats caption",
        "anchor: middle\n    stats: {{caption: {w}, items: [{{value: '1'}}]}}",
        "{w}",
    ),
    (
        "callouts",
        "anchor: middle\n    callouts: {{items: [{{head: Head, body: {w}}}]}}",
        "Head\n{w}",
    ),
    ("flow", "flow: {{items: [{{head: One, body: {w}}}, {{head: Two}}]}}", "One\n{w}"),
    (
        "versus",
        "versus: {{left: {{value: a, label: b, note: {w}}}, right: {{value: c, label: d}}}}",
        "a\nb\n{w}",
    ),
    ("image", "image: {{src: pic.png, alt: A picture, over: [{{text: {w}}}]}}", "{w}"),
]


def _placed(tmp_path, monkeypatch, body: str):
    deck, manifest = _build(
        tmp_path, monkeypatch, f"title: T\nplace:\n  - at: {{cols: full}}\n    {body}\n"
    )
    shapes = _shapes(deck, 1, "s1.p1")
    boxes = [(s.left, s.top, s.width, s.height) for s in shapes]
    texts = [s.text_frame.text for s in shapes if s.has_text_frame]
    recorded = [
        "\n".join(r["lines"]) if r.get("lines") else r.get("text")
        for r in manifest["slides"][0]["shapes"]
        if r["name"].startswith("s1.p1")
    ]
    return boxes, texts, recorded


@pytest.mark.parametrize(
    ("body", "expected"), [c[1:] for c in _OPTIONAL_COPY], ids=[c[0] for c in _OPTIONAL_COPY]
)
@pytest.mark.parametrize(("word", "ordinary"), [("no", "nx"), ("off", "ofx")])
def test_optional_copy_written_no_or_off_is_set_like_any_word(
    tmp_path, monkeypatch, body, expected, word, ordinary
):
    boxes, texts, recorded = _placed(tmp_path / "w", monkeypatch, body.format(w=word))
    assert expected.format(w=word) in texts
    assert expected.format(w=word) in recorded
    assert boxes == _placed(tmp_path / "o", monkeypatch, body.format(w=ordinary))[0]


def test_a_chrome_override_on_a_title_written_no(tmp_path, monkeypatch):
    """`title: no` is a title, so a `chrome:` treatment for it has a line to place."""
    deck, _ = _build(tmp_path, monkeypatch, "title: no\nchrome: {title: {align: center}}\n")
    assert [s.text_frame.text for s in _shapes(deck, 1, "s1.chrome")] == ["no"]


def _stats_in_a_short_band(tmp_path, monkeypatch, fields: str):
    return _build(
        tmp_path,
        monkeypatch,
        f"""
        title: T
        place:
          - at: {{cols: full, rows: {{from: 1, to: 5}}}}
            stats: {{{fields}items: [{{value: '1'}}]}}
        """,
    )


def test_one_row_of_tiles_fits_the_short_band(tmp_path, monkeypatch):
    deck, _ = _stats_in_a_short_band(tmp_path, monkeypatch, "")
    assert _shapes(deck, 1, "s1.p1.stats")


@pytest.mark.parametrize("caption", ["off", "ofx"])
def test_a_stats_caption_written_off_is_measured_like_any_caption(tmp_path, monkeypatch, caption):
    """A caption under the tiles no longer fits; one set but not measured would overflow."""
    with pytest.raises(LayoutError, match=r"1 items in 1 columns need .* of height"):
        _stats_in_a_short_band(tmp_path, monkeypatch, f"caption: {caption}, ")
