"""What a successful `code` build never reaches."""

import dataclasses

import pytest
from pptx.util import Pt

import deckwright.components  # noqa: F401 — registers the built-ins
from deckwright.errors import LayoutError
from deckwright.layouts.components import get_component


def test_a_wrap_that_is_no_boolean_is_refused(ctx_factory):
    """A quoted `"no"` is text, and every string is truthy: it would wrap the listing."""
    ctx = ctx_factory({"code": {"lines": ["x = 1"], "wrap": "no"}})
    with pytest.raises(LayoutError, match=r"'wrap' must be true or false, got 'no'"):
        get_component("code")(ctx)


@pytest.mark.parametrize(
    "body, match",
    [
        ({"lines": ["x"], "text": "x"}, r"give 'lines' or 'text', not both"),
        ({"lines": []}, r"needs 'lines' \(a list\) or 'text' \(a block scalar\)"),
        ({"lines": "x = 1"}, r"needs 'lines' \(a list\) or 'text' \(a block scalar\)"),
        ({"lines": ["x"], "size": "big"}, r"'size' is a point size, got 'big'"),
        ({"lines": ["x"], "size": 1}, r"'size' 1\.0pt is below the theme's \d+\.\dpt minimum"),
        (
            {"lines": [f"line {i}" for i in range(400)]},
            r"400 lines need \d+\.\d\din but only \d+\.\d\din is left",
        ),
        (
            {"lines": ["x = 1", "y" * 200]},
            r"line 2 needs \d+\.\d\din but the plate holds \d+\.\d\din .* set 'wrap: true'",
        ),
        (
            {"lines": ["word " * 60] * 40, "wrap": True},
            r"40 lines wrap to \d+ rows and need \d+\.\d\din but only",
        ),
        (
            {"lines": ["\t" * 60 + "x"]},
            r"line 1 needs \d+\.\d\din but the plate holds",
        ),
    ],
)
def test_code_refuses_by_name(ctx_factory, body, match):
    ctx = ctx_factory({"code": body})
    with pytest.raises(LayoutError, match=match):
        get_component("code")(ctx)


def test_code_text_is_split_into_its_lines_without_a_trailing_blank(ctx_factory):
    """A block scalar ends in a newline; kept, it draws a fourth, empty line."""
    ctx = ctx_factory({"code": {"text": "a\nb\nc\n"}})
    get_component("code")(ctx)
    assert "a\nb\nc" in [s.text_frame.text for s in ctx.slide.shapes if s.has_text_frame]


def _plate_h(ctx_factory, theme, **body) -> int:
    courier = dataclasses.replace(theme, mono="Courier New")
    ctx = ctx_factory({"code": {"size": 14, **body}}, theme_override=courier)
    get_component("code")(ctx)
    return next(s.height for s in ctx.slide.shapes if s.has_text_frame)


@pytest.mark.parametrize(
    "line",
    [
        " " * 8 + "a" * 92,
        "a" * 45 + " " * 10 + "b" * 45,
        "a" * 100,
    ],
    ids=["indent", "aligned", "one-word"],
)
def test_a_wrapped_line_deepens_the_plate_by_the_rows_it_wraps_to(ctx_factory, theme, line):
    """A two-row plate wraps at 99.1 Courier columns. Each line is 100, its spaces charged
    where they show, so its plate is two rows deep; 92 characters alone hold one row, so the
    indent is what wraps."""
    one, two = (
        _plate_h(ctx_factory, theme, lines=["a"]),
        _plate_h(ctx_factory, theme, lines=["a", "b"]),
    )
    assert _plate_h(ctx_factory, theme, lines=["a" * 92], wrap=True) == one
    assert _plate_h(ctx_factory, theme, lines=[line], wrap=True) == two


@pytest.mark.parametrize("line", ["a" * 98, "a" * 40 + " " + "b" * 57], ids=["word", "words"])
def test_a_line_that_fits_its_row_is_one_row_when_wrapped(ctx_factory, theme, line):
    """98 columns of the 99.1 a one-row plate wraps at: charged ``_MARGIN``, they would be
    101.9, whether one run breaks where it runs out or the second word moves down whole."""
    one = _plate_h(ctx_factory, theme, lines=["a"])
    assert _plate_h(ctx_factory, theme, lines=[line], wrap=True) == one


def test_a_plate_sets_its_lines_inside_its_rounded_text_rectangle(ctx_factory, theme):
    """Fifteen rows make a plate 4.9in deep, which insets its text rectangle 0.06in a side:
    98.3 columns are left of the 99.3 its padding alone would leave. Wrapped, the 99 take a
    second row, and sixteen rows hold 98.2."""
    filler = ["x"] * 14
    _plate_h(ctx_factory, theme, lines=["a" * 98, *filler])
    with pytest.raises(LayoutError, match=r"line 1 needs \d+\.\d\din but the plate holds"):
        _plate_h(ctx_factory, theme, lines=["a" * 99, *filler])
    sixteen = _plate_h(ctx_factory, theme, lines=["x"] * 16)
    assert _plate_h(ctx_factory, theme, lines=["a" * 99, *filler], wrap=True) == sixteen


def test_a_wrapped_listing_is_set_as_wide_as_a_plate_of_its_own_depth(ctx_factory, theme):
    """99 columns fit the 99.1 a one-row plate sets, though not the 98.2 of one filling the
    5.3in left: wrapped or not, the line is one row."""
    one = _plate_h(ctx_factory, theme, lines=["a"])
    assert _plate_h(ctx_factory, theme, lines=["a" * 99]) == one
    assert _plate_h(ctx_factory, theme, lines=["a" * 99], wrap=True) == one


def test_a_file_tree_is_charged_the_columns_courier_draws_it_in(ctx_factory, theme):
    """Courier carries the box-drawing glyphs a tree is drawn in, one column each. Charged
    the widest glyph measured, this 98-column line would need 76 ems of the plate's 59."""
    tree = "│   " * 24 + "├─"
    _plate_h(ctx_factory, theme, lines=[tree, "└── café → λ ≤ Ж"])


def test_a_mono_face_named_like_a_proportional_family_is_charged_a_monospaced_advance(
    ctx_factory, theme
):
    """ "Helvetica Monospaced" carries "helvetica", which routes to Arial's table: 120 ``i``
    would be 27 ems there and build, though they set 74."""
    face = dataclasses.replace(theme, mono="Helvetica Monospaced")
    ctx = ctx_factory({"code": {"lines": ["i" * 120], "size": 14}}, theme_override=face)
    with pytest.raises(LayoutError, match=r"'Helvetica Monospaced' has no width table"):
        get_component("code")(ctx)


def test_a_tab_is_drawn_as_spaces_to_the_next_four_column_stop(ctx_factory):
    ctx = ctx_factory({"code": {"lines": ["\tx", "ab\tc"]}})
    get_component("code")(ctx)
    assert "    x\nab  c" in [s.text_frame.text for s in ctx.slide.shapes if s.has_text_frame]


def test_an_unmeasured_mono_face_is_charged_a_monospaced_advance(ctx_factory):
    """The fixture's Consolas has no table. Charged ``CEILING``, 120 narrow ``i`` fit in
    41 ems, and a listing that runs off its plate builds."""
    ctx = ctx_factory({"code": {"lines": ["i" * 120], "size": 14}})
    with pytest.raises(LayoutError, match=r"line 1 needs .*'Consolas' has no width table"):
        get_component("code")(ctx)


def test_trailing_spaces_are_not_charged_to_an_unwrapped_line(ctx_factory):
    ctx = ctx_factory({"code": {"lines": ["a" * 80 + " " * 200], "size": 14}})
    get_component("code")(ctx)
    assert any(s.text_frame.text.startswith("a" * 80) for s in ctx.slide.shapes if s.has_text_frame)


def test_a_wrapped_line_breaks_after_a_hyphen_as_the_renderer_does(ctx_factory, theme):
    """Each hyphenated run fits a row alone but not beside the one before it, so the line is
    three rows; broken only where its 185 characters run out, it would be two."""
    three = _plate_h(ctx_factory, theme, lines=["a", "b", "c"])
    line = "a" * 10 + "-" + "b" * 88 + "-" + "c" * 85
    assert _plate_h(ctx_factory, theme, lines=[line], wrap=True) == three


def test_a_listing_is_set_at_the_line_advance_its_plate_is_sized_for(ctx_factory, theme):
    """14pt rows exactly 21.7pt apart in every renderer, so three of them and the plate's
    margins are its whole depth; left at single spacing, a renderer sets them closer. The
    4.9pt a row sets over single spacing lands above it, so the top margin gives it up."""
    ctx = ctx_factory(
        {"code": {"lines": ["a", "b", "c"], "size": 14}},
        theme_override=dataclasses.replace(theme, mono="Courier New"),
    )
    get_component("code")(ctx)
    plate = next(s for s in ctx.slide.shapes if s.has_text_frame)
    frame = plate.text_frame
    assert [p.line_spacing.pt for p in frame.paragraphs] == [21.7] * 3
    assert (frame.margin_top.pt, frame.margin_bottom.pt) == (
        pytest.approx(9.5),
        pytest.approx(19.3),
    )
    assert plate.height == pytest.approx(
        3 * Pt(21.7) + frame.margin_top + frame.margin_bottom, abs=12700
    )


@pytest.mark.parametrize("wrap", [False, True], ids=["unwrapped", "wrapped"])
def test_a_placement_too_narrow_for_any_plate_is_refused_by_its_width(ctx_factory, wrap):
    """A plate's padding takes 0.52in across: in a 0.5in placement no line has room, and the
    empty first line is not what is wrong."""
    from deckwright.theme.model import Rect

    ctx = ctx_factory({"code": {"lines": ["", "x"], "wrap": wrap}})
    ctx.rect = Rect(left=1.0, top=1.0, width=0.5, height=3.0)
    with pytest.raises(LayoutError, match=r"the placement is 0\.50in wide.* widen the placement"):
        get_component("code")(ctx)
