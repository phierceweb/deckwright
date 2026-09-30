"""`equation`: Office Math behind a one-line fallback, one shape wrapped whole."""

from __future__ import annotations

import pytest
from pptx.oxml.ns import qn

import deckwright.components  # noqa: F401 — registers the built-ins
from deckwright.errors import LayoutError
from deckwright.layouts.components import get_component
from deckwright.theme.model import Rect

_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
QUADRATIC = r"x = \frac{-b \pm \sqrt{b^2 - 4ac}}{2a}"
LINE = "x = (−b ± √(b² − 4ac))/(2a)"
WHERE = r"slide 1 \(component 'equation'\)"


def _set(ctx_factory, body, **kwargs):
    ctx = ctx_factory({"equation": body}, **kwargs)
    get_component("equation")(ctx)
    return ctx


def _alternate(ctx):
    (alternate,) = ctx.slide.shapes._spTree.findall(f"{{{_MC}}}AlternateContent")
    return alternate


def test_the_equation_is_one_shape_wrapped_whole(ctx_factory):
    """PowerPoint's own storage: the Choice holds the maths, the Fallback the same shape as
    text. A reader of neither extension sees one text box."""
    ctx = _set(ctx_factory, {"tex": QUADRATIC})
    alternate = _alternate(ctx)
    choice, fallback = list(alternate)
    assert choice.tag == f"{{{_MC}}}Choice" and choice.get("Requires") == "a14"
    assert fallback.tag == f"{{{_MC}}}Fallback"
    (maths,) = choice.iter(f"{{{_A14}}}m")
    assert maths.find(f"{{{_M}}}oMathPara") is not None
    (words,) = fallback.iter(qn("a:t"))
    assert words.text == LINE
    assert list(ctx.slide.shapes) == []  # python-pptx sees neither branch; utils/mce resolves it


def test_both_branches_are_the_same_shape(ctx_factory):
    ctx = _set(ctx_factory, {"tex": QUADRATIC})
    choice, fallback = list(_alternate(ctx))
    ids = [el.get("id") for el in (*choice.iter(qn("p:cNvPr")), *fallback.iter(qn("p:cNvPr")))]
    boxes = [el.attrib for el in (*choice.iter(qn("a:off")), *fallback.iter(qn("a:off")))]
    assert ids[0] == ids[1] and len(ids) == 2
    assert boxes[0] == boxes[1]


def test_the_manifest_records_the_line_every_check_reads(ctx_factory):
    ctx = _set(ctx_factory, {"tex": QUADRATIC})
    (record,) = ctx.manifest.slides[0].shapes
    assert (record.lines, record.text, record.alt) == ([LINE], LINE, LINE)
    assert record.font_pt == pytest.approx(19.0)
    assert record.fg and record.bg


def test_an_authors_alt_text_is_kept_and_decorative_leaves_none(ctx_factory):
    ctx = _set(ctx_factory, {"tex": "x", "alt": "x, the unknown"})
    assert ctx.manifest.slides[0].shapes[0].alt == "x, the unknown"
    quiet = _set(ctx_factory, {"tex": "x", "decorative": True})
    assert (
        quiet.manifest.slides[0].shapes[0].alt,
        quiet.manifest.slides[0].shapes[0].decorative,
    ) == (
        None,
        True,
    )


def test_on_an_inverse_slide_the_maths_and_its_fallback_share_the_ink(ctx_factory):
    ctx = _set(ctx_factory, {"tex": "x"}, background="inverse")
    choice, fallback = list(_alternate(ctx))
    ink = ctx.manifest.slides[0].shapes[0].fg
    assert ink == ctx.theme.palette.pair("inverse").fg
    assert {c.get("val") for c in choice.iter(qn("a:srgbClr"))} == {ink}
    assert {c.get("val") for c in fallback.iter(qn("a:srgbClr"))} == {ink}


def test_the_rung_sets_the_size_of_both_branches(ctx_factory):
    ctx = _set(ctx_factory, {"tex": "x", "rung": "title"})
    choice, fallback = list(_alternate(ctx))
    assert {r.get("sz") for r in choice.iter(qn("a:rPr"))} == {"3200"}
    assert {r.get("sz") for r in fallback.iter(qn("a:rPr"))} == {"3200"}


def test_a_stacked_fraction_reserves_a_second_row(ctx_factory):
    """In PowerPoint the fraction stacks; the box is sized for that, not for the one line."""
    flat = _set(ctx_factory, {"tex": "a + b"}).manifest.slides[0].shapes[0].box
    stacked = _set(ctx_factory, {"tex": r"\frac{a}{b}"}).manifest.slides[0].shapes[0].box
    assert round(stacked.h / flat.h, 2) == 2.0


@pytest.mark.parametrize(
    "body, message",
    [
        ({}, "'tex' is required — the equation, written in the LaTeX subset"),
        ({"tex": "  "}, "'tex' is required — the equation, written in the LaTeX subset"),
        ({"tex": ["x"]}, "'tex' is the equation as one string, got a list"),
        ({"tex": "x", "rung": "huge"}, "'rung' 'huge' is not a type rung this theme has; it has "),
        ({"tex": r"\begin{x}"}, r"'tex' uses \\begin, which is not supported"),
    ],
)
def test_an_equation_the_build_cannot_set_is_refused_by_name(ctx_factory, body, message):
    with pytest.raises(LayoutError, match=f"^{WHERE}: {message}"):
        _set(ctx_factory, body)


def test_an_equation_wider_than_its_placement_is_refused_not_wrapped(ctx_factory):
    ctx = ctx_factory({"equation": {"tex": "a + b + c + d + e + f + g + h + i + j + k"}})
    ctx.rect = Rect(ctx.rect.left, ctx.rect.top, 1.0, ctx.rect.height)
    with pytest.raises(
        LayoutError,
        match=rf"^{WHERE}: the equation needs \d+\.\d\din on one line but the placement",
    ):
        get_component("equation")(ctx)


def test_an_equation_taller_than_its_placement_is_refused(ctx_factory):
    ctx = ctx_factory({"equation": {"tex": r"\frac{\frac{a}{b}}{\frac{c}{d}}"}})
    ctx.rect = Rect(ctx.rect.left, ctx.rect.top, ctx.rect.width, 0.5)
    with pytest.raises(LayoutError, match=rf"^{WHERE}: the equation stacks 4 rows, \d+\.\d\din"):
        get_component("equation")(ctx)
