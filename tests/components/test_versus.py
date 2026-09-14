"""What a successful `versus` build never reaches."""

from __future__ import annotations

import dataclasses

import pytest

import deckwright.components  # noqa: F401 — registers the built-ins
from deckwright.components.versus import _MIDDLE, _side_floor
from deckwright.errors import LayoutError
from deckwright.layouts.components import get_component, shape_id

LEFT = {"value": "2 days", "label": "by post"}
RIGHT = {"value": "4 hours", "label": "collected in person", "highlight": True}


def _ctx(ctx_factory, **body):
    return ctx_factory({"versus": {"left": LEFT, "right": RIGHT, **body}})


def test_the_glyph_is_inside_a_reveal_group(ctx_factory):
    """Outside every group it hangs between two plates that have not arrived yet."""
    ctx = _ctx(ctx_factory)
    groups = get_component("versus")(ctx).groups
    grouped = {shape_id(spid) for group in groups for spid in group}
    glyphs = [s.shape_id for s in ctx.slide.shapes if s.name.startswith("Icon ")]
    assert glyphs, "no glyph was drawn, so this proves nothing"
    assert set(glyphs) <= grouped


def test_one_reveal_group_per_side(ctx_factory):
    assert len(get_component("versus")(_ctx(ctx_factory)).groups) == 2


def test_the_highlighted_side_is_painted_in_the_accent(ctx_factory):
    """Without this the two plates are indistinguishable and the pair states nothing."""
    ctx = _ctx(ctx_factory)
    get_component("versus")(ctx)
    plates = sorted((s for s in ctx.slide.shapes if not s.text_frame.text), key=lambda s: s.left)
    fills = [s.fill.fore_color.rgb for s in plates if s.fill.type is not None]
    accent = ctx.theme.palette.role("accent-1")
    assert str(fills[-1]) == accent, (fills, accent)
    assert str(fills[0]) != accent


def test_a_side_needs_a_value_and_a_label(ctx_factory):
    ctx = ctx_factory({"versus": {"left": {"value": "1"}, "right": RIGHT}})
    with pytest.raises(LayoutError, match=r"'left' needs a 'value' and a 'label'"):
        get_component("versus")(ctx)


def test_a_missing_side_names_it(ctx_factory):
    ctx = ctx_factory({"versus": {"left": LEFT}})
    with pytest.raises(LayoutError, match=r"'right' needs a 'value' and a 'label'"):
        get_component("versus")(ctx)


def test_an_unknown_side_key_is_refused(ctx_factory):
    ctx = ctx_factory({"versus": {"left": {**LEFT, "colour": "red"}, "right": RIGHT}})
    with pytest.raises(LayoutError, match=r"has the unknown field 'colour'"):
        get_component("versus")(ctx)


def test_highlighting_both_sides_is_refused(ctx_factory):
    ctx = ctx_factory({"versus": {"left": {**LEFT, "highlight": True}, "right": RIGHT}})
    with pytest.raises(LayoutError, match=r"both sides set 'highlight'"):
        get_component("versus")(ctx)


def test_every_returned_id_is_a_real_shape(ctx_factory):
    ctx = _ctx(ctx_factory)
    groups = get_component("versus")(ctx).groups
    ids = {s.shape_id for s in ctx.slide.shapes}
    assert all(shape_id(spid) in ids for group in groups for spid in group)


def test_versus_is_registered():
    from deckwright.layouts.components import registered_components

    assert "versus" in registered_components()


def _plates(ctx):
    """The two rounded plates, left to right, as (left_in, width_in)."""
    from pptx.util import Emu

    return sorted(
        (Emu(s.left).inches, Emu(s.width).inches)
        for s in ctx.slide.shapes
        if s.name.startswith("Rounded Rectangle")
    )


def test_the_plates_are_sized_in_proportion_to_their_values(ctx_factory):
    """Length is the encoding, not the printed number."""
    ctx = ctx_factory(
        {
            "versus": {
                "left": {"value": "11.2%", "label": "ours", "highlight": True},
                "right": {"value": "7.0%", "label": "theirs"},
            }
        }
    )
    get_component("versus")(ctx)
    (_, wide), (_, narrow) = _plates(ctx)
    assert wide / narrow == pytest.approx(11.2 / 7.0, rel=0.02)


def test_values_that_are_not_numbers_keep_the_even_split(ctx_factory):
    """`before` and `after` carry no magnitude, so there is nothing to encode."""
    ctx = ctx_factory(
        {
            "versus": {
                "left": {"value": "before", "label": "a"},
                "right": {"value": "after", "label": "b"},
            }
        }
    )
    get_component("versus")(ctx)
    (_, one), (_, two) = _plates(ctx)
    assert one == pytest.approx(two)


def test_values_in_different_units_keep_the_even_split(ctx_factory):
    """2 against 4 would draw the longer span smaller — the opposite of the claim."""
    ctx = ctx_factory({"versus": {"left": LEFT, "right": RIGHT}})
    get_component("versus")(ctx)
    (_, one), (_, two) = _plates(ctx)
    assert one == pytest.approx(two)


def test_a_shared_unit_is_still_sized_in_proportion(ctx_factory):
    ctx = ctx_factory(
        {
            "versus": {
                "left": {"value": "2 days", "label": "a"},
                "right": {"value": "8 days", "label": "b"},
            }
        }
    )
    get_component("versus")(ctx)
    (_, one), (_, four) = _plates(ctx)
    assert four / one == pytest.approx(4.0, rel=0.02)


def test_a_different_order_of_magnitude_is_not_a_shared_unit(ctx_factory):
    """`$480K` is not 400x `$1.2M`; the suffix is part of the number."""
    ctx = ctx_factory(
        {
            "versus": {
                "left": {"value": "$1.2M", "label": "a"},
                "right": {"value": "$480K", "label": "b"},
            }
        }
    )
    get_component("versus")(ctx)
    (_, one), (_, two) = _plates(ctx)
    assert one == pytest.approx(two)


def test_a_lopsided_pair_keeps_the_smaller_plate_legible(ctx_factory):
    """1 against 99 would draw a sliver that cannot hold its own label."""
    ctx = ctx_factory(
        {"versus": {"left": {"value": 1, "label": "a"}, "right": {"value": 99, "label": "b"}}}
    )
    get_component("versus")(ctx)
    assert min(w for _, w in _plates(ctx)) >= 1.2


def test_the_small_plate_grows_to_hold_its_own_type(ctx_factory):
    """A flat floor sets `$1.2M` as `$1.` over `2M`; the floor has to measure the text."""
    ctx = ctx_factory(
        {
            "versus": {
                "left": {"value": "$1.2M", "label": "Manual effort per year"},
                "right": {"value": "$48M", "label": "automated"},
            }
        }
    )
    get_component("versus")(ctx)
    small = min(w for _, w in _plates(ctx))
    # The literal width `$1.2M` needs at this fixture's 36pt stat rung, plus both insets.
    # Not recomputed with `text_em`: that is the floor's own measure and would agree with
    # any answer it gave.
    assert small == pytest.approx(1.6702, abs=0.001)


def test_a_short_pair_keeps_the_absolute_floor(ctx_factory):
    """A measured floor below 1.2in must not licence a sliver."""
    ctx = ctx_factory(
        {"versus": {"left": {"value": "1", "label": "a"}, "right": {"value": "99", "label": "b"}}}
    )
    get_component("versus")(ctx)
    assert min(w for _, w in _plates(ctx)) == pytest.approx(1.2, abs=0.01)


def test_two_sides_that_cannot_both_be_set_are_refused(ctx_factory):
    ctx = ctx_factory(
        {
            "versus": {
                "left": {"value": "1", "label": "Reconciliation"},
                "right": {"value": "2", "label": "Reconciliation"},
            }
        }
    )
    ctx = dataclasses.replace(ctx, rect=dataclasses.replace(ctx.body_rect, width=2.0))
    with pytest.raises(LayoutError) as e:
        get_component("versus")(ctx)
    assert "widen the placement" in str(e.value)


def test_the_glyph_sits_between_the_plates_however_they_are_sized(ctx_factory):
    """Positioned off one plate's width, an uneven pair would strand it inside the other."""
    from pptx.util import Emu

    ctx = ctx_factory(
        {
            "versus": {
                "left": {"value": "11.2%", "label": "ours"},
                "right": {"value": "7.0%", "label": "theirs"},
            }
        }
    )
    get_component("versus")(ctx)
    (l1, w1), (l2, _) = _plates(ctx)
    glyph = next(s for s in ctx.slide.shapes if s.name.startswith("Icon "))
    gx = Emu(glyph.left).inches
    assert l1 + w1 <= gx and gx + Emu(glyph.width).inches <= l2


def test_an_even_split_still_respects_a_side_that_needs_more(ctx_factory):
    """Mixed units take the even split, which on a narrow placement can sit below what
    one side's own type needs."""
    left_side = {"value": "2 days", "label": "Reconciliation"}
    right_side = {"value": "9", "label": "x"}
    ctx = ctx_factory({"versus": {"left": left_side, "right": right_side}})
    stat, body, caption = ctx.style("stat"), ctx.style("body"), ctx.style("caption")
    rungs = {
        "value_pt": stat.size * 0.86,
        "label_pt": body.size,
        "note_pt": caption.size,
        "value_face": ctx.theme.font_for(stat),
        "text_face": ctx.theme.face,
    }
    wide, narrow = _side_floor(left_side, **rungs), _side_floor(right_side, **rungs)
    total = wide + narrow + 0.1
    assert total / 2 < wide, "premise: an even split has to fall short of the wide side"

    ctx = dataclasses.replace(ctx, rect=dataclasses.replace(ctx.body_rect, width=total + _MIDDLE))
    get_component("versus")(ctx)
    left, right = (w for _, w in _plates(ctx))
    assert left == pytest.approx(wide), "the long side did not take the room its type needs"
    assert left > right


def test_a_singular_against_its_own_plural_keeps_the_even_split(ctx_factory):
    """No string rule separates `hrs`/`hr` from `ms`/`m`, so only an exact match sizes."""
    ctx = ctx_factory(
        {
            "versus": {
                "left": {"value": "1 day", "label": "a"},
                "right": {"value": "4 days", "label": "b"},
            }
        }
    )
    get_component("versus")(ctx)
    one, two = (w for _, w in _plates(ctx))
    assert one == pytest.approx(two)


def test_an_abbreviation_is_never_folded_into_another_unit(ctx_factory):
    ctx = ctx_factory(
        {
            "versus": {
                "left": {"value": "200 ms", "label": "a"},
                "right": {"value": "3 m", "label": "b"},
            }
        }
    )
    get_component("versus")(ctx)
    one, two = (w for _, w in _plates(ctx))
    assert one == pytest.approx(two)


def test_a_versus_refusal_on_an_unmeasured_face_says_the_estimate_errs_wide(ctx_factory, theme):
    unmeasured = dataclasses.replace(theme, face="Segoe UI", heading_face="Segoe UI")
    ctx = ctx_factory(
        {
            "versus": {
                "left": {"value": "1", "label": "Reconciliation"},
                "right": {"value": "2", "label": "Reconciliation"},
            }
        },
        theme_override=unmeasured,
    )
    ctx = dataclasses.replace(ctx, rect=dataclasses.replace(ctx.body_rect, width=2.0))
    with pytest.raises(LayoutError, match=r"'Segoe UI' has no width table"):
        get_component("versus")(ctx)


def test_a_long_link_address_does_not_widen_the_side_it_sits_in(ctx_factory):
    long = "https://example.com/documentation/" + "onboarding-" * 20 + "guide"
    ctx = ctx_factory(
        {
            "versus": {
                "left": {"value": "2 days", "label": f"read [the setup guide]({long}) first"},
                "right": RIGHT,
            }
        }
    )
    get_component("versus")(ctx)
