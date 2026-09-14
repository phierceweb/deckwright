"""What the ``image`` component refuses, and the geometry a mask forces on it: a build reaches
the squaring but never reads it back, and a circle drawn on an oblong is an oval."""

from __future__ import annotations

import pytest
from pptx.util import Inches

import deckwright.components  # noqa: F401 — registers the built-in components
from deckwright.errors import LayoutError
from deckwright.layouts.components import get_component


def _draw(ctx):
    get_component("image")(ctx)
    return list(ctx.slide.shapes)


def test_a_circle_mask_squares_the_placement(ctx_factory, white_wide):
    """A 16:9 placement masked to an ellipse would draw an oval, so the box is squared."""
    ctx = ctx_factory({"image": {"src": str(white_wide), "mask": "circle"}})
    picture = _draw(ctx)[0]
    assert picture.width == picture.height


def test_an_unmasked_picture_keeps_the_whole_placement(ctx_factory, white_wide):
    """The contrast with the test above: nothing else squares a placement."""
    ctx = ctx_factory({"image": {"src": str(white_wide)}})
    picture = _draw(ctx)[0]
    assert picture.width > picture.height


def test_a_circle_mask_refuses_a_contain_fit(ctx_factory, white_wide):
    """``contain`` re-oblongs the squared box to the source's aspect — an oval again."""
    ctx = ctx_factory({"image": {"src": str(white_wide), "mask": "circle", "fit": "contain"}})
    with pytest.raises(LayoutError, match="circle mask needs 'fit: cover'"):
        _draw(ctx)


def test_a_rounded_mask_accepts_a_contain_fit(ctx_factory, black_tall):
    """The refusal is specific to circles: a rounded rectangle of any aspect is fine."""
    ctx = ctx_factory({"image": {"src": str(black_tall), "mask": "rounded", "fit": "contain"}})
    assert _draw(ctx)


def test_a_radius_beyond_a_half_is_refused(ctx_factory, white_wide):
    ctx = ctx_factory({"image": {"src": str(white_wide), "mask": "rounded", "radius": 0.75}})
    with pytest.raises(LayoutError, match="fraction of the picture's short side"):
        _draw(ctx)


def test_a_missing_src_is_refused(ctx_factory):
    with pytest.raises(LayoutError, match="'src' must name an image file"):
        _draw(ctx_factory({"image": {"mask": "circle"}}))


def test_an_unknown_field_names_the_known_ones(ctx_factory, white_wide):
    ctx = ctx_factory({"image": {"src": str(white_wide), "opacity": 0.5}})
    with pytest.raises(LayoutError, match="unknown field 'opacity'"):
        _draw(ctx)


def test_an_over_line_without_text_is_refused(ctx_factory, white_wide):
    ctx = ctx_factory({"image": {"src": str(white_wide), "over": [{"rung": "title"}]}})
    with pytest.raises(LayoutError, match="every 'over' line needs a 'text'"):
        _draw(ctx)


def test_an_over_line_with_an_unknown_key_names_the_known_ones(ctx_factory, white_wide):
    ctx = ctx_factory(
        {"image": {"src": str(white_wide), "over": [{"text": "Hi", "colour": "red"}]}}
    )
    with pytest.raises(LayoutError, match="has no key 'colour'"):
        _draw(ctx)


def test_a_crop_that_is_not_an_aspect_is_refused(ctx_factory, white_wide):
    ctx = ctx_factory({"image": {"src": str(white_wide), "crop": "widescreen"}})
    with pytest.raises(LayoutError, match="write it as '16:9'"):
        _draw(ctx)


def test_an_inset_wider_than_the_picture_is_refused(ctx_factory, white_wide):
    """Silently clamping would stack the text outside the picture it belongs to."""
    ctx = ctx_factory({"image": {"src": str(white_wide), "inset": 0.9, "over": [{"text": "Hi"}]}})
    with pytest.raises(LayoutError, match="leaves the text no width"):
        _draw(ctx)


def test_text_over_a_white_photograph_gets_a_scrim_nobody_asked_for(ctx_factory, white_wide):
    """The failure the component exists to prevent: white ink straight onto white."""
    ctx = ctx_factory(
        {"image": {"src": str(white_wide), "over": [{"text": "Unreadable without one"}]}}
    )
    shapes = _draw(ctx)
    assert len(shapes) == 3, "expected picture, scrim, textbox"
    scrim = shapes[1]
    assert scrim.width == Inches(ctx.body_rect.width)


def test_what_is_behind_a_scrimmed_picture_is_the_picture_under_its_scrim(ctx_factory, white_wide):
    """Text laid over the picture reads against the scrim covering it, not its bare pixels."""
    from deckwright.utils.color import relative_luminance

    ctx = ctx_factory(
        {"image": {"src": str(white_wide), "scrim": {"pair": "inverse", "opacity": 0.9}}}
    )
    _draw(ctx)
    ground = ctx.behind(ctx.body_rect, ink="FFFFFF")
    assert relative_luminance(ground) < 0.2


def test_a_picture_carries_the_alt_it_was_given_and_never_its_file_name(ctx_factory, white_wide):
    """python-pptx writes the file name as `descr`; a screen reader would read it aloud."""
    described = _draw(ctx_factory({"image": {"src": str(white_wide), "alt": "A white sheet"}}))[0]
    bare = _draw(ctx_factory({"image": {"src": str(white_wide)}}))[0]
    assert described._element.nvPicPr.cNvPr.get("descr") == "A white sheet"
    assert bare._element.nvPicPr.cNvPr.get("descr") is None


def test_a_background_image_is_marked_decorative(theme, white_wide):
    """The backdrop has no component to take `alt:`, and is what the slide sits on."""
    from pptx import Presentation

    from deckwright.compile.manifest import ManifestRecorder
    from deckwright.imagery.paint import paint_backdrop
    from deckwright.layouts.registry import SlideCtx
    from deckwright.spec.model import Background, SlideSpec
    from deckwright.utils.a11y import described

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    spec = SlideSpec(index=1, background=Background(kind="image", image=str(white_wide)))
    manifest = ManifestRecorder(deck="d", theme="t")
    manifest.begin_slide(1, background=spec.background.pair)
    ctx = SlideCtx(
        slide=prs.slides.add_slide(prs.slide_layouts[6]), theme=theme, spec=spec, manifest=manifest
    )
    paint_backdrop(ctx)
    picture = next(s for s in ctx.slide.shapes if s.shape_type == 13)
    assert described(picture) == (None, True)
