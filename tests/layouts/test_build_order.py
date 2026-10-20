"""What a slide's build leaves out and what it adds last: a placement kept on screen, a
chrome line that arrives last, and a chart's bars followed by its claim."""

from __future__ import annotations

import json

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from deckwright.compile import build_deck
from deckwright.errors import LayoutError, SpecError

_HEAD = "theme: base\ntitle: T\nout: out/B.pptx\n---\n"
_CHART = (
    "chart: {kind: column, data: [{category: A, value: 1}, {category: B, value: 2}, "
    "{category: C, value: 3}]}"
)


def _build(tmp_path, slide: str):
    spec = tmp_path / "b.deck.yaml"
    spec.write_text(_HEAD + slide)
    built = build_deck(spec)
    shapes = {s.shape_id: s.name for s in Presentation(str(built.deck)).slides[0].shapes}
    slide_xml = Presentation(str(built.deck)).slides[0]._element
    targets = [
        shapes[int(t.get("spid"))]
        for t in slide_xml.iter(qn("p:spTgt"))
        if t.getparent().getparent().tag == qn("p:cBhvr")
        and t.getparent().getparent().getparent().tag == qn("p:set")
    ]
    data = json.loads(built.manifest.read_text())["slides"][0]
    return targets, data, slide_xml


def test_a_placement_with_animate_none_is_on_screen_and_out_of_the_build(tmp_path):
    targets, data, _ = _build(
        tmp_path,
        "title: Slots\nanimate: one_at_a_time\nplace:\n"
        "  - {at: {cols: full, rows: {from: 0, to: 3}}, animate: none, "
        "table: {header: [a, b], rows: [['1', '2']]}}\n"
        "  - {at: {cols: left-half, rows: {from: 4, to: 8}}, card: {heading: cat}}\n"
        "  - {at: {cols: right-half, rows: {from: 4, to: 8}}, card: {heading: dog}}\n",
    )
    assert not any("table" in t for t in targets)
    assert [a["kind"] for a in data["animations"]] == ["click_sequence"]
    assert data["animations"][0]["clicks"] == 2


def test_a_chrome_line_with_build_last_arrives_on_the_slides_last_click(tmp_path):
    targets, data, _ = _build(
        tmp_path,
        "title: The claim, after its evidence\nanimate: one_at_a_time\n"
        "chrome:\n  title: {at: {cols: full, rows: {from: 9, to: 11}}, build: last}\n"
        "place:\n  - {at: {cols: full, rows: {from: 0, to: 8}}, card: {heading: Evidence}}\n",
    )
    assert targets[-1] == "s1.chrome.title"
    assert data["animations"][0]["clicks"] == 2


def test_a_together_slide_with_a_line_built_last_records_both_as_written(tmp_path):
    """One timeline, two kinds: the beats view, the beat-size check and rhythm each read the
    kind, and none of them should take this slide for one_at_a_time."""
    _, data, _ = _build(
        tmp_path,
        "title: T\nsubtitle: The claim\nanimate: together\n"
        "chrome:\n  subtitle: {at: {box: {x: 6%, y: 85%, w: 60%, h: auto}}, build: last}\n"
        "place:\n  - {at: {cols: left-half}, card: {heading: A}}\n"
        "  - {at: {cols: right-half}, card: {heading: B}}\n",
    )
    assert [(a["kind"], a["clicks"]) for a in data["animations"]] == [
        ("click_build", 1),
        ("line_build", 1),
    ]


def test_a_chrome_line_that_builds_on_a_slide_with_no_animate_is_its_only_click(tmp_path):
    targets, data, _ = _build(
        tmp_path,
        "title: Hold the claim\n"
        "chrome:\n  title: {at: {cols: full, rows: {from: 9, to: 11}}, build: last}\n"
        "place:\n  - {at: {cols: full, rows: {from: 0, to: 8}}, card: {heading: Evidence}}\n",
    )
    assert targets == ["s1.chrome.title"]
    assert [(a["kind"], a["clicks"]) for a in data["animations"]] == [("line_build", 1)]


def test_a_chart_built_by_category_is_followed_by_the_claim_under_it(tmp_path):
    targets, data, slide = _build(
        tmp_path,
        f"title: Did it work?\nanimate: by_category\nplace:\n"
        f"  - at: {{cols: full, rows: {{from: 0, to: 9}}}}\n    {_CHART}\n"
        f"  - {{at: {{cols: full, rows: {{from: 9, to: 12}}}}, "
        f"callouts: {{items: [{{head: It did}}]}}}}\n",
    )
    assert [a["kind"] for a in data["animations"]] == ["chart_build", "click_sequence"]
    assert data["animations"][0]["clicks"] == 4 and data["animations"][1]["clicks"] == 1
    builds = slide.find(qn("p:timing")).find(qn("p:bldLst"))
    assert {child.tag.split("}")[1] for child in builds} == {"bldGraphic", "bldP"}


def test_a_chart_build_and_a_line_built_last_record_what_each_is(tmp_path):
    _, data, _ = _build(
        tmp_path,
        f"title: Bars\nsubtitle: The claim\nanimate: by_category\n"
        f"chrome:\n  subtitle: {{at: {{box: {{x: 6%, y: 88%, w: 60%, h: auto}}}}, build: last}}\n"
        f"place:\n  - at: {{cols: full, rows: {{from: 0, to: 9}}}}\n    {_CHART}\n",
    )
    assert [(a["kind"], a["clicks"]) for a in data["animations"]] == [
        ("chart_build", 4),
        ("line_build", 1),
    ]


def test_a_note_kept_on_screen_beside_a_chart_that_builds_alone(tmp_path):
    _, data, _ = _build(
        tmp_path,
        f"title: Bars\nanimate: by_category\nplace:\n"
        f"  - at: {{cols: full, rows: {{from: 0, to: 9}}}}\n    {_CHART}\n"
        f"  - {{at: {{cols: full, rows: {{from: 9, to: 12}}}}, animate: none, "
        f"callouts: {{items: [{{head: Read the axis}}]}}}}\n",
    )
    assert [a["kind"] for a in data["animations"]] == ["chart_build"]


def test_a_chart_kept_still_stays_out_of_the_build_beside_one_that_builds(tmp_path):
    """errors.md's way out of two charts asking to build: one of them `animate: none`."""
    targets, data, slide = _build(
        tmp_path,
        f"title: Two\nanimate: by_category\nplace:\n"
        f"  - at: {{cols: left-half}}\n    animate: none\n    {_CHART}\n"
        f"  - at: {{cols: right-half}}\n    id: moving\n    {_CHART}\n",
    )
    built = [b.get("spid") for b in slide.iter(qn("p:bldGraphic"))]
    names = {
        str(s.shape_id): s.name for s in Presentation(str(tmp_path / "out/B.pptx")).slides[0].shapes
    }
    assert {names[spid] for spid in built} == {"s1.moving.chart#1"}
    assert [a["kind"] for a in data["animations"]] == ["chart_build"]


@pytest.mark.parametrize(
    ("slide", "error", "match"),
    [
        (
            "title: S\nplace:\n  - {at: {cols: full}, animate: none, card: {heading: H}}\n",
            SpecError,
            "this slide does not build",
        ),
        (
            "title: S\nanimate: one_at_a_time\nplace:\n"
            "  - {at: {cols: full}, animate: together, card: {heading: H}}\n",
            SpecError,
            "a placement's 'animate' takes only none",
        ),
        (
            "title: S\nanimate: one_at_a_time\nplace:\n"
            "  - {at: {cols: left-half}, id: q, card: {heading: Q}}\n"
            "  - {at: {cols: right-half}, reveals: q, animate: none, card: {heading: A}}\n",
            SpecError,
            "hides it until that is clicked",
        ),
        (
            "title: S\nchrome:\n  title: {build: last}\nplace:\n"
            "  - {at: {cols: full}, card: {heading: H}}\n",
            LayoutError,
            "sits in the stack",
        ),
        (
            "title: S\nchrome:\n  title: {at: {cols: full, rows: {from: 9, to: 11}}, build: last}\n"
            "place:\n  - {at: {cols: left-half, rows: {from: 0, to: 8}}, id: q, card: {heading: Q}}\n"
            "  - {at: {cols: right-half, rows: {from: 0, to: 8}}, reveals: q, card: {heading: A}}\n",
            LayoutError,
            "a chrome line that builds and a 'reveals:' trigger",
        ),
        (
            f"title: S\nanimate: by_series\nplace:\n"
            f"  - at: {{cols: left-half}}\n    {_CHART}\n"
            f"  - at: {{cols: right-half}}\n    {_CHART}\n",
            LayoutError,
            "a second chart asks to build",
        ),
        (
            f"title: S\nanimate: by_category\nplace:\n"
            f"  - at: {{cols: left-half}}\n    animate: none\n    {_CHART}\n"
            f"  - {{at: {{cols: right-half}}, card: {{heading: H}}}}\n",
            LayoutError,
            "its only chart is 'animate: none'",
        ),
        (
            f"title: S\nanimate: by_category\nplace:\n"
            f"  - at: {{cols: left-half}}\n    animate: none\n    {_CHART}\n"
            f"  - at: {{cols: right-half}}\n    animate: none\n    {_CHART}\n",
            LayoutError,
            "its charts are all 'animate: none'",
        ),
        (
            "title: S\nanimate: by_category\nplace:\n"
            "  - {at: {cols: full}, animate: none, card: {heading: H}}\n",
            LayoutError,
            "animate 'by_category' only applies to a native chart",
        ),
    ],
)
def test_a_build_that_cannot_be_written_is_refused_by_name(tmp_path, slide, error, match):
    spec = tmp_path / "b.deck.yaml"
    spec.write_text(_HEAD + slide)
    with pytest.raises(error, match=match):
        build_deck(spec)
