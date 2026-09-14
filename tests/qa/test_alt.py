"""The `alt-text` check reads the saved package. Every case is a real saved deck."""

from __future__ import annotations

import io
import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from deckwright.qa.alt import check_alt_text
from deckwright.qa.model import Severity
from deckwright.utils.a11y import describe, description


@pytest.fixture
def png(tmp_path):
    path = tmp_path / "photo.png"
    Image.new("RGB", (8, 8)).save(path)
    return path


def _saved(tmp_path, build) -> str:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    build(slide)
    path = tmp_path / "d.pptx"
    prs.save(str(path))
    return str(path)


def test_a_picture_with_no_alt_text_is_warned_about(tmp_path, png):
    def build(slide):
        describe(slide.shapes.add_picture(str(png), 0, 0))

    findings = check_alt_text(_saved(tmp_path, build))
    assert [(f.slide, f.check, f.severity) for f in findings] == [(1, "alt-text", Severity.WARN)]
    assert findings[0].detail.startswith("picture ")
    assert "has no alternative text" in findings[0].detail


def test_python_pptxs_file_name_does_not_count_as_alt_text(tmp_path, png):
    """add_picture writes `descr="photo.png"`; a screen reader reads that aloud."""
    findings = check_alt_text(_saved(tmp_path, lambda s: s.shapes.add_picture(str(png), 0, 0)))
    assert len(findings) == 1
    assert "only its file name 'photo.png'" in findings[0].detail


def test_a_picture_with_alt_text_is_clean(tmp_path, png):
    def build(slide):
        describe(slide.shapes.add_picture(str(png), 0, 0), alt="A dark square")

    assert check_alt_text(_saved(tmp_path, build)) == []


def test_a_picture_marked_decorative_is_clean(tmp_path, png):
    def build(slide):
        describe(slide.shapes.add_picture(str(png), 0, 0), decorative=True)

    assert check_alt_text(_saved(tmp_path, build)) == []


def test_a_chart_is_named_as_a_chart(tmp_path):
    def build(slide):
        data = CategoryChartData()
        data.categories = ["A", "B"]
        data.add_series("s", (1, 2))
        slide.shapes.add_chart(XL_CHART_TYPE.PIE, 0, 0, Inches(4), Inches(3), data)

    findings = check_alt_text(_saved(tmp_path, build))
    assert len(findings) == 1
    assert findings[0].detail.startswith("chart ")


def test_a_table_is_text_and_is_not_asked_for_alt(tmp_path):
    def build(slide):
        slide.shapes.add_table(2, 2, 0, 0, Inches(4), Inches(2))

    assert check_alt_text(_saved(tmp_path, build)) == []


def test_a_picture_inside_a_group_is_checked(tmp_path, png):
    def build(slide):
        group = slide.shapes.add_group_shape()
        describe(group.shapes.add_picture(str(png), 0, 0))

    assert len(check_alt_text(_saved(tmp_path, build))) == 1


def test_a_markup_compatibility_fallback_picture_is_not_a_second_figure(tmp_path, png):
    """The Fallback is what an older reader draws in the Choice's place."""
    path = _saved(tmp_path, lambda s: describe(s.shapes.add_picture(str(png), 0, 0), alt="x"))
    part = "ppt/slides/slide1.xml"
    with zipfile.ZipFile(path) as z:
        entries = {n: z.read(n) for n in z.namelist()}
    xml = entries[part].decode()
    pic = xml[xml.index("<p:pic>") : xml.index("</p:pic>") + len("</p:pic>")]
    bare = pic.replace(' descr="x"', "")
    wrapped = (
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
        f'<mc:Choice Requires="p14">{pic}</mc:Choice><mc:Fallback>{bare}</mc:Fallback>'
        "</mc:AlternateContent>"
    )
    entries[part] = xml.replace(pic, wrapped).encode()
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as z:
        for name, data in entries.items():
            z.writestr(name, data)
    rebuilt = tmp_path / "wrapped.pptx"
    rebuilt.write_bytes(buffer.getvalue())

    assert check_alt_text(rebuilt) == []


def test_an_unreadable_file_is_left_to_the_package_check(tmp_path):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a zip")
    assert check_alt_text(bad) == []


def test_describe_without_alt_removes_an_existing_description(tmp_path, png):
    prs = Presentation()
    picture = prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_picture(str(png), 0, 0)
    assert picture._element.nvPicPr.cNvPr.get("descr") == "photo.png"
    describe(picture)
    assert picture._element.nvPicPr.cNvPr.get("descr") is None


@pytest.mark.parametrize(
    "val, decorative", [('val="1"', True), ('val="true"', True), ("", True), ('val="0"', False)]
)
def test_the_decorative_flag_is_read_by_its_val(val, decorative):
    from lxml import etree

    c_nv_pr = etree.fromstring(
        '<p:cNvPr xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" id="2" name="x">'
        '<a:extLst><a:ext uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}">'
        '<adec:decorative xmlns:adec="http://schemas.microsoft.com/office/drawing/2017/decorative" '
        f"{val}/></a:ext></a:extLst></p:cNvPr>"
    )
    assert description(c_nv_pr) == (None, decorative)


@pytest.mark.parametrize(
    "name", ["Screen Shot 2026-09-01 at 10.00.00.png", "team photo.JPG", "C:\\Users\\me\\photo.png"]
)
def test_a_file_name_with_spaces_or_a_path_is_not_alt_text(name):
    from deckwright.utils.a11y import is_file_name

    assert is_file_name(name)
    assert not is_file_name("The team at the summit")


def test_the_preview_picture_inside_an_embedded_object_is_not_a_figure_of_its_own(tmp_path, png):
    from pptx.enum.shapes import PROG_ID

    xlsx = tmp_path / "book.xlsx"
    xlsx.write_bytes(b"PK\x05\x06" + b"\x00" * 18)

    def build(slide):
        frame = slide.shapes.add_ole_object(str(xlsx), PROG_ID.XLSX, 0, 0, icon_file=str(png))
        describe(frame, alt="Quarterly figures workbook")

    assert check_alt_text(_saved(tmp_path, build)) == []
