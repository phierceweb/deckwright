"""What a deck's chart parts say about the charts: a claim the render cannot verify, and a
chart too thin for the treatment it was given. Every case is a real saved deck."""

from __future__ import annotations

import pathlib
import re

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from deckwright.qa.charts import check_charts
from deckwright.qa.model import Severity


def _charted(tmp_path, chart_type, values, name="c.pptx", *, categories=("Up", "Down")):
    """A saved deck holding one chart of ``chart_type``."""
    from pptx.chart.data import CategoryChartData

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = list(categories)
    data.add_series("s", values)
    slide.shapes.add_chart(chart_type, Inches(1), Inches(1), Inches(6), Inches(4), data)
    path = tmp_path / name
    prs.save(str(path))
    return path


# --- charts the render cannot verify ----------------------------------------


def test_a_bar_chart_with_a_negative_value_is_flagged_unverifiable(tmp_path):
    """The file is correct; the render this suite and `deckwright render` both go through plots the
    absolute value, so nothing automated can check the chart. Reproduced with bare python-pptx."""
    deck = _charted(tmp_path, XL_CHART_TYPE.COLUMN_CLUSTERED, (271, -146))
    findings = [f for f in check_charts(deck) if f.check == "chart-negative"]
    assert len(findings) == 1
    assert findings[0].severity is Severity.WARN
    assert findings[0].slide == 1
    assert "diverge" in findings[0].detail


def test_an_all_positive_bar_chart_is_clean(tmp_path):
    deck = _charted(tmp_path, XL_CHART_TYPE.COLUMN_CLUSTERED, (271, 146))
    assert [f for f in check_charts(deck) if f.check == "chart-negative"] == []


def test_a_line_chart_with_a_negative_value_is_not_flagged(tmp_path):
    """Line series render their negatives correctly, so warning about them is noise."""
    deck = _charted(tmp_path, XL_CHART_TYPE.LINE, (271, -146))
    assert [f for f in check_charts(deck) if f.check == "chart-negative"] == []


# --- charts too thin for the treatment they were given -----------------------


def test_a_two_bar_chart_is_flagged_and_the_finding_quotes_the_rule(tmp_path):
    """The class of slide both blind judges flagged, in two separate runs: a two-bar chart
    under a headline that already states both numbers."""
    deck = _charted(tmp_path, XL_CHART_TYPE.COLUMN_CLUSTERED, (36, 64))
    findings = [f for f in check_charts(deck) if f.check == "chart-datapoints"]
    assert len(findings) == 1
    assert findings[0].severity is Severity.WARN
    assert findings[0].slide == 1
    assert "plots 2 datapoints" in findings[0].detail


def test_the_quoted_rule_is_the_sentence_choosing_md_carries(tmp_path):
    """The finding quotes the doc. Reword one side only and it quotes nothing — worse than a
    paraphrase, because it still reads as sourced."""
    choosing = (pathlib.Path(__file__).resolve().parents[2] / "docs/choosing.md").read_text(
        encoding="utf-8"
    )
    deck = _charted(tmp_path, XL_CHART_TYPE.COLUMN_CLUSTERED, (36, 64))
    detail = next(f for f in check_charts(deck) if f.check == "chart-datapoints").detail
    quoted = re.search(r'"(A chart with[^"]+)"', detail)
    assert quoted is not None, detail
    assert quoted.group(1) in choosing


def test_a_four_bar_chart_is_not_flagged(tmp_path):
    """Four is the floor the rule states, so four passes it."""
    deck = _charted(
        tmp_path,
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        (18, 36, 52, 64),
        categories=("Remote", "Urban", "Suburban", "Rural"),
    )
    assert [f for f in check_charts(deck) if f.check == "chart-datapoints"] == []


def test_values_are_counted_across_series_not_bars_per_category(tmp_path):
    """Two categories against two series is four numbers, and reads as a comparison. Counting
    categories instead flags real charts deckwright's own feature tour builds."""
    from pptx.chart.data import CategoryChartData

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["Up", "Down"]
    data.add_series("before", (36, 64))
    data.add_series("after", (41, 59))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(4), data
    )
    deck = tmp_path / "two.pptx"
    prs.save(str(deck))
    assert [f for f in check_charts(deck) if f.check == "chart-datapoints"] == []


def test_a_thin_line_chart_is_exempt(tmp_path):
    """A line's claim is the direction between ordered points, which no `stats` row can make."""
    deck = _charted(tmp_path, XL_CHART_TYPE.LINE, (36, 64))
    assert [f for f in check_charts(deck) if f.check == "chart-datapoints"] == []


def test_a_thin_pie_is_exempt(tmp_path):
    """The parts sum to the whole, so a slice count is the composition being described."""
    deck = _charted(tmp_path, XL_CHART_TYPE.PIE, (36, 64))
    assert [f for f in check_charts(deck) if f.check == "chart-datapoints"] == []
