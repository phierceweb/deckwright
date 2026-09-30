"""`qa` after `render` checks the render already on disk instead of converting the deck again,
and converts again once the deck's bytes change. Driven through the CLI with no `--outdir`, so
the two commands' default directories have to coincide for the reuse to happen at all."""

from __future__ import annotations

import subprocess
from pathlib import Path

from pptx import Presentation
from typer.testing import CliRunner

from deckwright.cli import app
from deckwright.qa import runner as qa_runner
from deckwright.services import render as render_mod

cli = CliRunner()


def _fake_converters(argv, **kwargs):
    if "--convert-to" in argv:
        outdir = Path(argv[argv.index("--outdir") + 1])
        (outdir / f"{Path(argv[-1]).stem}.pdf").write_bytes(b"%PDF-1.4\n%%EOF\n")
    elif "-jpeg" in argv:
        (Path(argv[-1]).parent / "slide-1.jpg").write_bytes(b"\xff\xd8\xff")
    return subprocess.CompletedProcess(argv, 0, "", "")


def test_qa_reuses_the_render_of_these_bytes_and_renders_again_after_an_edit(
    monkeypatch, clean_manifest_deck
):
    deck, _ = clean_manifest_deck
    monkeypatch.setattr(render_mod.subprocess, "run", _fake_converters)
    converted: list[Path] = []
    monkeypatch.setattr(
        qa_runner, "render_to_images", lambda d, out, **kw: converted.append(d) or []
    )
    monkeypatch.setattr(qa_runner, "extract_pages", lambda pdf, **kw: ["hello"])
    monkeypatch.setattr(qa_runner, "check_rendered_faces", lambda *a: [])

    assert cli.invoke(app, ["render", str(deck)]).exit_code == 0
    result = cli.invoke(app, ["qa", str(deck)])
    assert result.exit_code == 0, result.output
    assert converted == []

    edited = Presentation()
    edited.slides.add_slide(edited.slide_layouts[6])
    edited.slides.add_slide(edited.slide_layouts[6])
    edited.save(str(deck))
    result = cli.invoke(app, ["qa", str(deck)])
    assert result.exit_code == 0, result.output
    assert converted == [deck]
