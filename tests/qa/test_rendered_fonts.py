"""The render's own font list: which faces it set, and whether CJK had a font to draw with.
pdffonts and fontconfig are stood in for; the rows are copied from real LibreOffice PDFs."""

from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches

import deckwright.qa.rendered_fonts as rf
from deckwright.qa.model import Severity

_BROKEN = """name                                 type              encoding         emb sub uni object ID
------------------------------------ ----------------- ---------------- --- --- --- ---------
BAAAAA+LinuxLibertineG               TrueType          WinAnsi          yes yes yes    129  0
CAAAAA+FrankRuhlHofshi-Bold          Type 1            Builtin          yes yes yes    124  0
"""
_DRAWN = """name                                 type              encoding         emb sub uni object ID
------------------------------------ ----------------- ---------------- --- --- --- ---------
BAAAAA+HiraginoSansGB-W3             Type 1            Builtin          yes yes yes    134  0
EAAAAA+Poppins-Bold                  CID TrueType      Identity-H       yes yes yes    159  0
"""


def test_pdffonts_rows_are_sliced_by_its_own_rule_even_where_a_type_has_a_space():
    assert rf.parse_pdffonts(_DRAWN) == [
        ("BAAAAA+HiraginoSansGB-W3", "Builtin"),
        ("EAAAAA+Poppins-Bold", "Identity-H"),
    ]


def _deck(tmp_path, face="Poppins"):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    run = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame.paragraphs[0].add_run()
    run.text = "売上"
    run.font.name = face
    ea = run._r.rPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    ea.set("typeface", "游ゴシック")
    run._r.rPr.append(ea)
    path = tmp_path / "d.pptx"
    prs.save(str(path))
    return path


_MANIFEST = {
    "slides": [
        {"index": 1, "shapes": [{"text": "売上"}]},
        {"index": 2, "shapes": [{"text": "Plain"}]},
    ]
}


def _stand_in(monkeypatch, table, *, cjk=frozenset({"HiraginoSansGB-W3"}), aliases=None):
    monkeypatch.setattr(rf, "embedded_fonts", lambda pdf, page=None: rf.parse_pdffonts(table))
    monkeypatch.setattr(rf, "cjk_postscript_names", lambda: cjk)
    monkeypatch.setattr(rf, "postscript_names", lambda: aliases or {})


def test_the_faces_a_deck_names_are_read_off_its_runs(tmp_path):
    assert rf.deck_faces(_deck(tmp_path)) == {"Poppins", "游ゴシック"}


def test_a_face_the_render_did_not_embed_is_reported_substituted(tmp_path, monkeypatch):
    _stand_in(monkeypatch, _BROKEN)
    findings = rf.check_rendered_faces(_deck(tmp_path), tmp_path / "d.pdf", _MANIFEST)
    substituted = [f for f in findings if f.check == "font-substituted"]
    assert [f.severity for f in substituted] == [Severity.WARN, Severity.WARN]
    assert "'Poppins'" in substituted[0].detail
    assert "FrankRuhlHofshi-Bold, LinuxLibertineG" in substituted[0].detail


def test_a_face_embedded_under_its_postscript_name_is_found(tmp_path, monkeypatch):
    """A PDF names 游ゴシック by the PostScript name fontconfig gives it."""
    table = (
        _DRAWN
        + "FAAAAA+YuGothic-Medium             Type 1            Builtin          yes yes yes    160  0\n"
    )
    _stand_in(monkeypatch, table, aliases={"游ゴシック": frozenset({"YuGothic-Medium"})})
    findings = rf.check_rendered_faces(_deck(tmp_path), tmp_path / "d.pdf", _MANIFEST)
    assert [f for f in findings if f.check == "font-substituted"] == []


def test_cjk_on_a_page_that_embeds_no_cjk_font_is_an_error_on_that_slide(tmp_path, monkeypatch):
    _stand_in(monkeypatch, _BROKEN)
    findings = rf.check_rendered_faces(_deck(tmp_path), tmp_path / "d.pdf", _MANIFEST)
    unrendered = [(f.slide, f.severity) for f in findings if f.check == "cjk-unrendered"]
    assert unrendered == [(1, Severity.ERROR)]


def test_cjk_drawn_from_a_cjk_font_is_clean_whatever_its_encoding(tmp_path, monkeypatch):
    """LibreOffice embeds a Hiragino subset as Type 1 with a one-byte encoding."""
    _stand_in(monkeypatch, _DRAWN)
    findings = rf.check_rendered_faces(_deck(tmp_path), tmp_path / "d.pdf", _MANIFEST)
    assert [f for f in findings if f.check == "cjk-unrendered"] == []


def test_without_fontconfig_the_cjk_half_says_nothing(tmp_path, monkeypatch):
    _stand_in(monkeypatch, _BROKEN, cjk=None)
    findings = rf.check_rendered_faces(_deck(tmp_path), tmp_path / "d.pdf", _MANIFEST)
    assert [f for f in findings if f.check == "cjk-unrendered"] == []


def test_without_pdffonts_the_caller_is_told_to_fall_back(tmp_path, monkeypatch):
    monkeypatch.setattr(rf, "embedded_fonts", lambda pdf, page=None: None)
    assert rf.check_rendered_faces(_deck(tmp_path), tmp_path / "d.pdf", _MANIFEST) is None
