"""Which fonts a render hands LibreOffice. fontconfig is stood in for by scripts printing fixed
answers; the one real render at the end needs LibreOffice, Poppler and fontconfig on the machine."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from deckwright.qa.rendered_fonts import embedded_fonts
from deckwright.services import render as render_mod
from deckwright.services import render_fonts as rf
from deckwright.theme.fonts import installed_families

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

#: What python-pptx's stock Office template names in its font scheme and master bullets.
_STOCK = frozenset(
    {
        "Angsana New", "Arial", "Calibri", "Cordia New", "DaunPenh", "DokChampa",
        "Estrangelo Edessa", "Euphemia", "Gautami", "Iskoola Pota", "Kalinga", "Kartika",
        "Latha", "MV Boli", "Mangal", "Microsoft Himalaya", "Microsoft Uighur",
        "Microsoft Yi Baiti", "Mongolian Baiti", "MoolBoran", "Nyala", "Plantagenet Cherokee",
        "Raavi", "Shruti", "Sylfaen", "Times New Roman", "Tunga", "Vrinda", "宋体", "新細明體",
        "맑은 고딕", "ＭＳ Ｐゴシック",
    }
)  # fmt: skip


def _run_with(paragraph, text: str, *, latin: str, **faces: str) -> None:
    run = paragraph.add_run()
    run.text = text
    run.font.name = latin
    for tag, face in faces.items():
        element = run._r.get_or_add_rPr().makeelement(f"{_A}{tag}")
        element.set("typeface", face)
        run._r.get_or_add_rPr().append(element)


def _deck(tmp_path, *, words="Revenue", category="North") -> Path:
    """A deck naming a face in each part LibreOffice draws from: slide, layout, master, chart."""
    prs = Presentation()
    master_title = prs.slide_master.placeholders[0].text_frame.paragraphs[0]
    _run_with(master_title, "Title", latin="Master Face", ea="+mn-ea")
    layout_title = prs.slide_layouts[1].placeholders[0].text_frame.paragraphs[0]
    _run_with(layout_title, "Title", latin="Layout Face")
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    body = slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text_frame.paragraphs[0]
    _run_with(body, words, latin="Poppins", ea="游ゴシック", cs="Noto Naskh Arabic")
    bullet = body._p.get_or_add_pPr().makeelement(f"{_A}buFont")
    bullet.set("typeface", "Wingdings")
    body._p.get_or_add_pPr().append(bullet)
    data = CategoryChartData()
    data.categories = [category, "South"]
    data.add_series("Sales", (1, 2))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, 0, Inches(2), Inches(4), Inches(3), data
    ).chart
    chart.font.name = "Chart Face"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return path


def test_every_face_the_deck_names_is_collected_and_theme_slots_are_not(tmp_path):
    faces, _ = rf.deck_fonts(_deck(tmp_path))
    assert faces == _STOCK | {
        "Master Face",
        "Layout Face",
        "Poppins",
        "游ゴシック",
        "Noto Naskh Arabic",
        "Wingdings",
        "Chart Face",
    }


@pytest.mark.parametrize(
    ("words", "category", "cjk"),
    [("Revenue", "North", False), ("売上", "North", True), ("Revenue", "東京", True)],
)
def test_cjk_is_read_from_the_words_and_chart_labels_not_the_theme_face_names(
    tmp_path, words, category, cjk
):
    """The stock theme names ＭＳ Ｐゴシック, so a scan of attributes would call every deck CJK."""
    assert rf.deck_fonts(_deck(tmp_path, words=words, category=category))[1] is cjk


def _stub(tmp_path, name: str, script: str) -> str:
    path = tmp_path / name
    path.write_text("#!/bin/sh\n" + script)
    path.chmod(0o755)
    return str(path)


def _fonts(tmp_path, *names: str) -> Path:
    home = tmp_path / "fonts"
    home.mkdir(exist_ok=True)
    for name in names:
        (home / name).write_bytes(b"")
    return home


def _fc_list(tmp_path, monkeypatch, listing: str) -> None:
    body = listing.replace("{home}", str(tmp_path / "fonts"))
    monkeypatch.setenv(
        "DECKWRIGHT_FC_LIST", _stub(tmp_path, "fc-list", f"cat <<'EOF'\n{body}EOF\n")
    )


def test_only_a_family_with_the_requested_name_is_linked_never_a_substitute(tmp_path, monkeypatch):
    home = _fonts(
        tmp_path, "Helvetica.ttc", "CourierNew.ttf", "CourierNewBold.ttf", "Hiragino.ttc",
        "Verdana.ttf", "Georgia.ttf",
    )  # fmt: skip
    _fc_list(
        tmp_path,
        monkeypatch,
        "Helvetica\t{home}/Helvetica.ttc\n"
        "Courier New\t{home}/CourierNew.ttf\n"
        "Courier New\t{home}/CourierNewBold.ttf\n"
        "Hiragino Sans,ヒラギノ角ゴシック\t{home}/Hiragino.ttc\n"
        "Verdana\t{home}/Verdana.ttf\n"
        "Georgia\t{home}/Georgia.ttf\n",
    )
    faces = frozenset({"helvetica", "CourierNew", "ヒラギノ角ゴシック", "Georgia Pro Light"})
    assert rf.font_files(faces, cjk=False) == [
        home / "CourierNew.ttf",
        home / "CourierNewBold.ttf",
        home / "Helvetica.ttc",
        home / "Hiragino.ttc",
    ]


@pytest.mark.parametrize("cjk", [True, False])
def test_fontconfigs_cjk_choice_is_added_only_for_a_deck_carrying_cjk(tmp_path, monkeypatch, cjk):
    home = _fonts(tmp_path, "cjk-ja.ttc", "cjk-ko.ttc", "cjk-zh-cn.ttc", "cjk-zh-tw.ttc")
    _fc_list(tmp_path, monkeypatch, "")
    asked = tmp_path / "asked"
    fc_match = f'echo "$3" >> "{asked}"\nprintf "%s" "{home}/cjk-${{3#:lang=}}.ttc"\n'
    monkeypatch.setenv("DECKWRIGHT_FC_MATCH", _stub(tmp_path, "fc-match", fc_match))
    files = rf.font_files(frozenset(), cjk=cjk)
    if cjk:
        assert [f.name for f in files] == [
            "cjk-ja.ttc", "cjk-ko.ttc", "cjk-zh-cn.ttc", "cjk-zh-tw.ttc",
        ]  # fmt: skip
    else:
        assert files == []
        assert not asked.exists(), asked.read_text()


def test_without_fontconfig_the_render_proceeds_with_no_fonts_to_link(tmp_path, monkeypatch):
    monkeypatch.setenv("DECKWRIGHT_FC_LIST", str(tmp_path / "absent-fc-list"))
    assert rf.deck_font_files(_deck(tmp_path)) == []


def test_the_render_links_the_fonts_into_the_profile_soffice_is_given(tmp_path, monkeypatch):
    home = _fonts(tmp_path, "Calibri.ttf")
    _fc_list(tmp_path, monkeypatch, "Calibri\t{home}/Calibri.ttf\n")
    deck = tmp_path / "deck.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(deck))
    seen: list[Path] = []
    real_run = subprocess.run

    def fake_run(argv, **kwargs):
        if "--convert-to" in argv:
            profile = next(a for a in argv if a.startswith("-env:UserInstallation="))
            fonts = Path(profile.split("file://", 1)[1]) / "user" / "fonts"
            seen.extend(p.resolve() for p in sorted(fonts.iterdir()))
            (tmp_path / "deck.pdf").write_bytes(b"%PDF-1.4\n%%EOF\n")
            return subprocess.CompletedProcess(argv, 0, "", "")
        if "-jpeg" in argv:
            return subprocess.CompletedProcess(argv, 0, "", "")
        return real_run(argv, **kwargs)

    monkeypatch.setattr(render_mod.subprocess, "run", fake_run)
    render_mod.render_to_images(deck, tmp_path)
    assert seen == [(home / "Calibri.ttf").resolve()]


def _can_render_helvetica() -> bool:
    tools = ("soffice", "pdftoppm", "pdffonts")
    return all(shutil.which(t) for t in tools) and "helvetica" in (installed_families() or ())


@pytest.mark.skipif(
    not _can_render_helvetica(), reason="needs LibreOffice, Poppler and an installed Helvetica"
)
def test_a_real_render_sets_the_face_the_deck_names(tmp_path):
    """LibreOffice on a fresh profile set this in Linux Libertine on macOS."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    _run_with(box.text_frame.paragraphs[0], "Quarterly revenue", latin="Helvetica")
    deck = tmp_path / "helvetica.pptx"
    prs.save(str(deck))
    render_mod.render_to_images(deck, tmp_path / "render")
    fonts = embedded_fonts(tmp_path / "render" / "helvetica.pdf")
    assert fonts is not None
    assert "Helvetica" in {name.split("+", 1)[-1] for name, _ in fonts}, fonts
