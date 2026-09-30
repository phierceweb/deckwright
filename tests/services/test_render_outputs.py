"""What the stale-page sweep may delete, and what the render reports back. ``--outdir`` is a
free-form user path, so a sweep matching every ``slide-*`` in it eats the PDF just written for a
deck named ``slide-deck.pptx``. Both subprocesses are faked; the sweep's result is a file fact."""

from __future__ import annotations

import subprocess
from pathlib import Path

import pytest

from deckwright.services import render as render_mod

PAGES = 2

JPEG = b"\xff\xd8\xff"


@pytest.fixture
def captured(monkeypatch):
    """Fake soffice (leaves a PDF) and pdftoppm (leaves PAGES pages), keeping both argv."""
    calls: list[list[str]] = []

    def fake_run(argv, **kwargs):
        calls.append(list(argv))
        if "--convert-to" in argv:
            outdir = Path(argv[argv.index("--outdir") + 1])
            (outdir / f"{Path(argv[-1]).stem}.pdf").write_bytes(b"%PDF-1.4\n%%EOF\n")
        else:
            prefix = Path(argv[-1])
            for page in range(1, PAGES + 1):
                prefix.with_name(f"{prefix.name}-{page}.jpg").write_bytes(JPEG)
        return subprocess.CompletedProcess(argv, 0, "", "")

    monkeypatch.setattr(render_mod.subprocess, "run", fake_run)
    return calls


def _render(tmp_path, outdir=None, *, deck="deck.pptx") -> list[str]:
    """Drive one render and return the image paths it reports."""
    source = tmp_path / deck
    source.write_bytes(b"not really a deck")
    return render_mod.render_to_images(source, outdir or tmp_path)


def test_notes_beside_the_pages_survive_the_sweep(captured, tmp_path):
    """Per-slide notes sit under the page glob; only their suffix keeps them."""
    notes = tmp_path / "slide-2-notes.md"
    notes.write_text("what slide 2 is for")
    _render(tmp_path)
    assert notes.exists()


def test_a_hand_named_export_of_one_slide_survives_the_sweep(captured, tmp_path):
    """A raster suffix is not enough to make a file ours: `slide-3-final.png` is
    something someone named and kept, and the sweep deletes what it matches."""
    kept = tmp_path / "slide-3-final.png"
    kept.write_bytes(JPEG)
    _render(tmp_path)
    assert kept.exists()


def test_a_deck_named_for_a_slide_keeps_its_own_pdf(captured, tmp_path):
    """The sweep runs after the conversion: this is the rasterizer's own input."""
    _render(tmp_path, deck="slide-deck.pptx")
    assert (tmp_path / "slide-deck.pdf").exists()


def test_a_page_from_a_longer_previous_render_is_removed(captured, tmp_path):
    """Two pages this time, ninety-nine last time — page 99 is nobody's."""
    stale = tmp_path / "slide-99.jpg"
    stale.write_bytes(JPEG)
    _render(tmp_path)
    assert not stale.exists()


def test_a_directory_named_like_a_page_is_not_swept(captured, tmp_path):
    """Unlinking a directory raises; the sweep skips it rather than taking the render down."""
    archive = tmp_path / "slide-9.jpg"
    archive.mkdir()
    _render(tmp_path)
    assert archive.is_dir()


def test_an_outdir_with_a_bracket_in_its_name_reports_its_pages(captured, tmp_path):
    """``[v2]`` is a character class to ``glob.glob``, which then matches no page at all."""
    images = _render(tmp_path, tmp_path / "Deck [v2]" / "render")
    assert [Path(p).name for p in images] == ["slide-1.jpg", "slide-2.jpg"]


def test_the_rasterizer_command_is_configurable(captured, tmp_path, monkeypatch):
    monkeypatch.setenv("DECKWRIGHT_PDFTOPPM", "poppler-pdftoppm")
    _render(tmp_path)
    rasterize = next(c for c in captured if "-jpeg" in c)
    assert rasterize[0] == "poppler-pdftoppm"


def test_pages_are_rasterised_to_a_whole_number_of_pixels(tmp_path, captured, monkeypatch):
    """13.333in at 110 dpi is 1466.63px. Left to round up, the last column is a third paper,
    and on a dark slide that is a light line down the right edge."""
    from pptx import Presentation
    from pptx.util import Inches

    # The fake writes a page for every command it is handed, the font query included.
    monkeypatch.chdir(tmp_path)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.slides.add_slide(prs.slide_layouts[6])
    deck = tmp_path / "wide.pptx"
    prs.save(str(deck))

    render_mod.render_to_images(deck, tmp_path / "out", dpi=110)

    raster = next(argv for argv in captured if argv[0].endswith("pdftoppm"))
    assert raster[raster.index("-scale-to-x") + 1] == "1467"
    assert raster[raster.index("-scale-to-y") + 1] == "825"


def test_a_deck_whose_size_cannot_be_read_is_rasterised_by_resolution_alone(tmp_path, captured):
    _render(tmp_path)
    raster = next(argv for argv in captured if argv[0].endswith("pdftoppm"))
    assert "-scale-to-x" not in raster
    assert "-r" in raster


def test_the_right_edge_of_a_dark_slide_is_the_slide(tmp_path, theme_file):
    """The lead itself, through the real tools: the last pixel column is the page, not paper."""
    import shutil

    from PIL import Image

    from deckwright.compile import build_deck

    if not (shutil.which("soffice") and shutil.which("pdftoppm")):
        pytest.skip("needs LibreOffice and pdftoppm")
    spec = tmp_path / "d.deck.yaml"
    spec.write_text(
        "theme: testtheme\ntitle: T\nout: out/D.pptx\n---\ntitle: Dark\nbackground: inverse\n"
    )
    built = build_deck(spec, theme_path=theme_file)
    (page,) = render_mod.render_to_images(built.deck, tmp_path / "r", fmt="png")
    with Image.open(page) as img:
        rgb = img.convert("RGB")
        edge, inside = rgb.getpixel((rgb.width - 1, 400)), rgb.getpixel((rgb.width - 20, 400))
    assert edge == inside == (0, 0, 0)
