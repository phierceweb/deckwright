"""Cards and pictures in a deck deckwright did not build, drafted as the components they are."""

from __future__ import annotations

import io

import pytest
import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from deckwright.spec.draft import draft_spec
from deckwright.spec.extract import harvest


def _png(colour="red") -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), colour).save(buf, format="PNG")
    buf.seek(0)
    return buf


def _slide():
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _plate(slide, left=1, top=1, width=4, height=2, *, text=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
    if text is not None:
        frame = shape.text_frame
        frame.text = text[0]
        for line in text[1:]:
            frame.add_paragraph().text = line
        for p in frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(18)
    return shape


def _box(slide, text, left=1.5, top=1.2, width=3, height=1.5, size=18):
    frame = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    ).text_frame
    frame.text = text[0]
    for line in text[1:]:
        frame.add_paragraph().text = line
    for p in frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
    return frame


def _save(prs, tmp_path, name="d.pptx"):
    path = tmp_path / name
    prs.save(str(path))
    return path


def _place(path, **kwargs):
    drafted = draft_spec(harvest(path), title="Draft", theme="base", **kwargs)
    return drafted, [d for d in yaml.safe_load_all(drafted.text) if isinstance(d, dict)][1:]


def test_a_filled_shape_holding_a_heading_and_a_line_is_a_card(tmp_path):
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _plate(slide, text=["The question", "What went wrong"])
    _, (doc,) = _place(_save(prs, tmp_path))
    assert doc["title"] == "A slide title"
    assert doc["place"][0]["card"] == {"heading": "The question", "body": "What went wrong"}


def test_a_filled_shape_holding_four_lines_stays_a_list(tmp_path):
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _plate(slide, text=["one", "two", "three", "four"])
    _, (doc,) = _place(_save(prs, tmp_path))
    assert doc["place"][0]["bullets"]["items"] == ["one", "two", "three", "four"]


def test_a_plate_with_one_text_box_on_it_is_one_card_and_the_plate_is_not_lost(tmp_path):
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _plate(slide)
    _box(slide, ["Heading", "Body line"])
    content = harvest(_save(prs, tmp_path))
    assert content[0].dropped == ()
    _, (doc,) = _place(tmp_path / "d.pptx")
    assert doc["place"][0]["card"] == {"heading": "Heading", "body": "Body line"}


def test_a_filled_banner_holding_the_slides_title_is_still_the_title(tmp_path):
    prs, slide = _slide()
    banner = _plate(slide, top=0.2, height=1, text=["The only big words"])
    for r in banner.text_frame.paragraphs[0].runs:
        r.font.size = Pt(40)
    _box(slide, ["a small line"], top=3, size=14)
    _, (doc,) = _place(_save(prs, tmp_path))
    assert doc["title"] == "The only big words"
    assert all("card" not in entry for entry in doc.get("place", []))


def test_a_picture_is_drafted_as_an_image_with_its_file_beside_the_draft(tmp_path):
    prs, slide = _slide()
    picture = slide.shapes.add_picture(_png(), Inches(1), Inches(1), Inches(4), Inches(2))
    picture._element._nvXxPr.cNvPr.set("descr", "A red tile")
    drafted, (doc,) = _place(_save(prs, tmp_path), media="d.media")
    (entry,) = doc["place"]
    assert entry["image"] == {"src": "d.media/slide1-1.png", "alt": "A red tile"}
    ((name, data),) = drafted.media.items()
    assert name == "d.media/slide1-1.png" and data.startswith(b"\x89PNG")


def test_a_picture_that_covers_the_slide_is_left_as_it_was(tmp_path):
    prs, slide = _slide()
    slide.shapes.add_picture(_png(), 0, 0, prs.slide_width, prs.slide_height)
    content = harvest(_save(prs, tmp_path))
    assert content[0].pictures == () and len(content[0].dropped) == 1


def test_a_linked_picture_is_named_with_its_reason(tmp_path):
    prs, slide = _slide()
    picture = slide.shapes.add_picture(_png(), Inches(1), Inches(1), Inches(4), Inches(2))
    blip = picture._element.blipFill.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    )
    r = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    blip.set(f"{r}link", blip.get(f"{r}embed"))
    del blip.attrib[f"{r}embed"]
    (dropped,) = harvest(_save(prs, tmp_path))[0].dropped
    assert dropped.endswith(" — a linked picture, whose file is not in the deck")


def test_a_group_holding_a_plate_and_its_words_is_read_in_its_own_coordinates(tmp_path):
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    group = slide.shapes.add_group_shape()
    _plate(group, left=5, top=3)
    _box(group, ["Grouped heading", "Grouped body"], left=5.5, top=3.2)
    _, (doc,) = _place(_save(prs, tmp_path))
    assert doc["place"][0]["card"] == {"heading": "Grouped heading", "body": "Grouped body"}


def test_a_panel_near_the_size_of_the_slide_is_not_a_card_plate(tmp_path):
    """Words on a panel that fills most of the slide are the slide's words, not a card's."""
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _plate(slide, left=0.2, top=1.0, width=9.4, height=6.2)
    _box(slide, ["Heading", "Body line"], left=1, top=2, size=18)
    content = harvest(_save(prs, tmp_path))
    assert content[0].cards == ()
    assert content[0].blocks == (("Heading", "Body line"),)


def test_a_text_box_on_two_plates_at_once_is_no_ones_card(tmp_path):
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _plate(slide)
    _plate(slide, left=0.8, top=0.9, width=4.5, height=2.4)
    _box(slide, ["Heading", "Body line"])
    content = harvest(_save(prs, tmp_path))
    assert content[0].cards == ()
    assert content[0].blocks == (("Heading", "Body line"),)


def test_a_text_box_running_off_the_bottom_of_a_plate_is_not_on_it(tmp_path):
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _plate(slide)
    _box(slide, ["Heading", "Body line"], height=2.5)
    assert harvest(_save(prs, tmp_path))[0].cards == ()


def test_a_plate_in_a_moved_group_is_measured_where_the_group_draws_it(tmp_path):
    """The group sits elsewhere on the slide; its plate's own numbers only look like 1in."""
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    group = slide.shapes.add_group_shape()
    _plate(group)
    off = group._element.grpSpPr.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm/"
        "{http://schemas.openxmlformats.org/drawingml/2006/main}off"
    )
    off.set("x", str(Inches(5))), off.set("y", str(Inches(4)))
    _box(slide, ["Heading", "Body line"])
    content = harvest(_save(prs, tmp_path))
    assert content[0].cards == ()
    assert content[0].blocks == (("Heading", "Body line"),)


@pytest.mark.parametrize("bold, key", [(True, "heading"), (False, "body")])
def test_a_card_with_one_line_is_a_heading_when_bold_and_copy_when_not(tmp_path, bold, key):
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    plate = _plate(slide, text=["The one line"])
    plate.text_frame.paragraphs[0].runs[0].font.bold = bold
    assert harvest(_save(prs, tmp_path))[0].cards == ({key: "The one line"},)


def _unfilled(shape):
    shape.fill.background()


def _style_fill_none(shape):
    shape._element.find(qn("p:style")).find(qn("a:fillRef")).set("idx", "0")


@pytest.mark.parametrize(
    "paint, card", [(_unfilled, False), (_style_fill_none, False), (None, True)]
)
def test_only_a_shape_that_paints_a_fill_is_a_card(tmp_path, paint, card):
    """An outline is not a plate, nor a style whose fill reference is 0; the style's own fill is."""
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(2))
    shape.text_frame.text = "Heading"
    shape.text_frame.add_paragraph().text = "Body line"
    for p in shape.text_frame.paragraphs:
        p.runs[0].font.size = Pt(18)
    if paint is not None:
        paint(shape)
    cards = harvest(_save(prs, tmp_path))[0].cards
    assert cards == (({"heading": "Heading", "body": "Body line"},) if card else ())


def _table(slide, rows, top=4):
    shape = slide.shapes.add_table(rows, 2, Inches(1), Inches(top), Inches(8), Inches(0.3 * rows))
    for r in range(rows):
        for c in range(2):
            shape.table.cell(r, c).text = f"r{r}c{c}"


def test_a_card_is_placed_with_the_words_before_the_figures(tmp_path):
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _table(slide, 2)
    _plate(slide, text=["The question", "What went wrong"])
    _, (doc,) = _place(_save(prs, tmp_path))
    assert [next(k for k in e if k != "at") for e in doc["place"]] == ["card", "table"]


def test_a_full_slide_keeps_a_cards_words_and_names_each_figure_it_has_no_room_for(tmp_path):
    """The table has no room in either layout and is named; the picture finds room beside
    the words."""
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _plate(slide, text=["The question", "What went wrong"])
    _table(slide, 40, top=3)
    picture = slide.shapes.add_picture(_png(), Inches(6), Inches(1), Inches(3), Inches(2))
    picture._element._nvXxPr.cNvPr.set("descr", "A red tile")
    drafted, (doc,) = _place(_save(prs, tmp_path))
    assert doc["place"][0]["bullets"]["items"] == ["The question", "What went wrong"]
    assert [e["at"]["cols"] for e in doc["place"] if "image" in e] == ["right-half"]
    assert "#   r39c0 | r39c1" in drafted.text
    assert list(drafted.media) == ["draft.media/slide1-1.png"]


def test_pictures_past_what_the_grid_holds_are_named_and_still_written(tmp_path):
    prs, slide = _slide()
    for n in range(3):
        slide.shapes.add_picture(_png(), Inches(0.5 + 3 * n), Inches(1), Inches(3), Inches(2))
    drafted, (doc,) = _place(_save(prs, tmp_path))
    assert [e["image"]["src"] for e in doc["place"]] == [
        "draft.media/slide1-1.png",
        "draft.media/slide1-2.png",
    ]
    assert drafted.spilled == 1 and "#   a picture — draft.media/slide1-3.png\n" in drafted.text
    assert sorted(drafted.media) == [f"draft.media/slide1-{n}.png" for n in (1, 2, 3)]


def test_the_transcript_carries_a_cards_words(tmp_path):
    from deckwright.spec.draft import render_markdown

    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _plate(slide, text=["The question", "What went wrong"])
    text = render_markdown(harvest(_save(prs, tmp_path)), title="Deck")
    assert "- The question\n- What went wrong\n" in text


def test_a_picture_filling_a_layouts_picture_placeholder_is_a_picture(tmp_path):
    """python-pptx types a filled placeholder as a placeholder, not as the picture it holds."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[8])  # Picture with Caption
    holder = next(s for s in slide.placeholders if s.placeholder_format.type == 18)
    holder.insert_picture(_png())
    content = harvest(_save(prs, tmp_path))[0]
    assert [p.name for p in content.pictures] == ["slide1-1.png"]
    assert not any("placeholder" in item for item in content.dropped)


def test_a_picture_that_would_push_words_off_goes_beside_them(tmp_path):
    """Stacked, the picture pushed lines off the slide and was named instead; beside the
    words, every line and the picture both fit."""
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    words = [f"Line {n} of the argument, long enough to wrap across the column." for n in range(9)]
    _box(slide, words, top=1.5, width=5, height=5, size=16)
    slide.shapes.add_picture(_png(), Inches(6), Inches(1.5), Inches(3), Inches(2))
    drafted, (doc,) = _place(_save(prs, tmp_path))
    assert [line for e in doc["place"] for line in e.get("bullets", {}).get("items", [])] == words
    assert [(e["at"]["cols"], next(k for k in e if k != "at")) for e in doc["place"]] == [
        ("left-half", "bullets"),
        ("right-half", "image"),
    ]
    assert drafted.spilled == 0


def test_a_picture_is_still_named_where_the_words_need_the_whole_width(tmp_path):
    """Halving the measure would push more words off than the picture does: it stays named."""
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    words = [f"Line {n} of the argument, long enough to wrap across the column." for n in range(14)]
    _box(slide, words, top=1.5, width=5, height=5, size=16)
    slide.shapes.add_picture(_png(), Inches(6), Inches(1.5), Inches(3), Inches(2))
    drafted, (doc,) = _place(_save(prs, tmp_path))
    assert [line for e in doc["place"] for line in e.get("bullets", {}).get("items", [])] == words
    assert not [e for e in doc["place"] if "image" in e]
    assert "#   a picture — draft.media/slide1-1.png\n" in drafted.text


def test_a_picture_stays_when_leaving_it_out_would_save_no_words(tmp_path):
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _box(slide, ["An argument that no band could ever hold. " * 120], top=1.5, width=5, size=16)
    slide.shapes.add_picture(_png(), Inches(6), Inches(1.5), Inches(3), Inches(2))
    _, (doc,) = _place(_save(prs, tmp_path))
    assert [e["image"]["src"] for e in doc["place"] if "image" in e] == ["draft.media/slide1-1.png"]


def test_a_text_box_where_a_moved_groups_plate_is_drawn_is_that_plates_card(tmp_path):
    """The plate's own numbers say 1in; the group draws it at 5in, 4in, and a slide-level
    text box there lies on it."""
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    group = slide.shapes.add_group_shape()
    _plate(group)
    off = group._element.grpSpPr.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm/"
        "{http://schemas.openxmlformats.org/drawingml/2006/main}off"
    )
    off.set("x", str(Inches(5))), off.set("y", str(Inches(4)))
    _box(slide, ["Heading", "Body line"], left=5.5, top=4.2)
    assert harvest(_save(prs, tmp_path))[0].cards == ({"heading": "Heading", "body": "Body line"},)


def test_a_slide_sized_panel_its_group_draws_small_is_a_card_plate(tmp_path):
    """9.4in by 6.2in by its own numbers, which no card is; its group draws it at 40%."""
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    group = slide.shapes.add_group_shape()
    _plate(group, left=0.2, top=1.0, width=9.4, height=6.2)
    _box(group, ["Heading", "Body line"], left=1, top=2, size=18)
    xfrm = group._element.grpSpPr.get_or_add_xfrm()
    xfrm.get_or_add_ext().cx = round(xfrm.get_or_add_chExt().cx * 0.4)
    xfrm.get_or_add_ext().cy = round(xfrm.get_or_add_chExt().cy * 0.4)
    assert harvest(_save(prs, tmp_path))[0].cards == ({"heading": "Heading", "body": "Body line"},)


def test_a_heading_box_and_a_body_box_on_one_plate_are_one_card(tmp_path):
    """Read one box at a time, the heading became a one-line card and the body a loose line."""
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _plate(slide)
    _box(slide, ["Body line"], top=2.0, height=0.6)
    _box(slide, ["Heading"], top=1.2, height=0.6)
    content = harvest(_save(prs, tmp_path))
    assert content[0].cards == ({"heading": "Heading", "body": "Body line"},)
    assert content[0].blocks == () and content[0].dropped == ()


def test_three_lines_across_the_boxes_on_one_plate_stay_words(tmp_path):
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _plate(slide)
    _box(slide, ["Heading"], top=1.1, height=0.5)
    _box(slide, ["One line", "and another"], top=1.8, height=1.0)
    content = harvest(_save(prs, tmp_path))
    assert content[0].cards == ()
    assert content[0].blocks == (("Heading",), ("One line", "and another"))


def test_a_box_on_two_plates_joins_neither_plates_card(tmp_path):
    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    _plate(slide)
    _plate(slide, left=0.8, top=0.9, width=4.5, height=2.4)
    _box(slide, ["Heading"], top=1.2, height=0.6)
    _box(slide, ["Body line"], top=2.0, height=0.6)
    content = harvest(_save(prs, tmp_path))
    assert content[0].cards == ()


def test_a_slide_of_cards_drafts_into_one_that_builds(tmp_path):
    """Each card is given the rows its type needs, so a slide of them does not draft into
    plates too shallow to set their words in."""
    from deckwright.compile import build_deck

    prs, slide = _slide()
    _box(slide, ["A slide title"], top=0.2, size=36)
    for n in range(5):
        _plate(
            slide,
            left=0.5,
            top=1.2 + 1.2 * n,
            width=9,
            height=1.1,
            text=[
                f"Step {n + 1}",
                "What happens at this step, what it needs from the one before it, and what "
                "it hands to the one after.",
            ],
        )
    drafted, _ = _place(_save(prs, tmp_path))
    spec = tmp_path / "draft.deck.yaml"
    spec.write_text(drafted.text)
    assert build_deck(spec).deck.is_file()


def test_a_card_drafted_with_its_icon_is_given_room_for_it(tmp_path):
    from deckwright.compile import build_deck
    from deckwright.spec.extract import SlideContent

    body = "What happens at this step, and what it hands to the one after."
    slide = SlideContent(
        index=1,
        cards=tuple({"heading": f"Step {n}", "body": body, "icon": "home"} for n in range(1, 4)),
    )
    spec = tmp_path / "draft.deck.yaml"
    spec.write_text(draft_spec([slide], title="Draft", theme="base").text)
    assert build_deck(spec).deck.is_file()


_STEP = (
    "What happens at this step, what it needs from the one before it, what it hands to the "
    "one after it, and who signs it off before anything moves on to the next of them all."
)


def _wide_slide():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _box(slide, ["Steps and pictures"], left=0.5, top=0.2, width=12, height=1, size=36)
    for n in range(2):
        slide.shapes.add_picture(_png(), Inches(7), Inches(1.5 + 2.9 * n), Inches(5), Inches(2.8))
    return prs, slide


def _three_cards(slide):
    for n in range(3):
        _plate(
            slide, left=0.5, top=1.5 + 1.93 * n, width=6, height=1.83, text=[f"Step {n + 1}", _STEP]
        )


def _a_table(slide):
    rows = slide.shapes.add_table(5, 3, Inches(0.5), Inches(1.5), Inches(6), Inches(5)).table
    for r in range(5):
        for c in range(3):
            rows.cell(r, c).text = f"Head {c}" if r == 0 else "what happens at this"


@pytest.mark.parametrize("words", [_three_cards, _a_table], ids=["cards", "table"])
def test_words_drafted_beside_the_pictures_are_sized_for_the_half_they_land_in(tmp_path, words):
    from deckwright.compile import build_deck

    prs, slide = _wide_slide()
    words(slide)
    drafted, (doc,) = _place(_save(prs, tmp_path))
    assert {e["at"]["cols"] for e in doc["place"]} == {"left-half", "right-half"}
    spec = tmp_path / "draft.deck.yaml"
    spec.write_text(drafted.text)
    for name, data in drafted.media.items():
        (tmp_path / name).parent.mkdir(parents=True, exist_ok=True)
        (tmp_path / name).write_bytes(data)
    assert build_deck(spec).deck.is_file()
