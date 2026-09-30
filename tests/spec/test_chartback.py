"""A chart in a deck deckwright did not build, read back as a chart block with its data."""

from __future__ import annotations

import pytest
import yaml
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE as T
from pptx.oxml.ns import qn
from pptx.util import Inches

from deckwright.spec._chartback import chart_block


def _deck(tmp_path, chart_type, series, *, categories=("Q1", "Q2", "Q3", "Q4"), name="c.pptx"):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Quarterly"
    data = CategoryChartData()
    data.categories = list(categories)
    for label, values in series:
        data.add_series(label, values)
    slide.shapes.add_chart(chart_type, Inches(1), Inches(2), Inches(6), Inches(4), data)
    path = tmp_path / name
    prs.save(str(path))
    return path


def _chart(path):
    return next(s for s in Presentation(str(path)).slides[0].shapes if s.has_chart)


def test_a_one_series_column_chart_is_a_value_per_category(tmp_path):
    block, why = chart_block(
        _chart(_deck(tmp_path, T.COLUMN_CLUSTERED, [("Sales", (1, 2.5, 3, 4))]))
    )
    assert why is None
    assert [type(row["value"]) for row in block["data"]] == [int, float, int, int]
    assert block == {
        "kind": "column",
        "data": [
            {"category": "Q1", "value": 1},
            {"category": "Q2", "value": 2.5},
            {"category": "Q3", "value": 3},
            {"category": "Q4", "value": 4},
        ],
    }


def test_two_series_are_named_values_per_category(tmp_path):
    series = [("Ads", (12, 20, 25, 30)), ("Organic", (8, 20, 39, 62))]
    block, _ = chart_block(_chart(_deck(tmp_path, T.COLUMN_STACKED, series)))
    assert block["kind"] == "column-stacked"
    assert block["data"][0] == {"category": "Q1", "values": {"Ads": 12, "Organic": 8}}


def test_a_pie_is_read_back_as_a_pie(tmp_path):
    block, _ = chart_block(_chart(_deck(tmp_path, T.PIE, [("Share", (5, 3, 2, 1))])))
    assert block["kind"] == "pie"


@pytest.mark.parametrize(
    "series, why",
    [
        ([("A", (1, None, 3, 4))], "a chart with an empty value"),
        ([("A", (1, 2, 3, 4)), ("A", (4, 3, 2, 1))], "a chart whose series share a name"),
    ],
)
def test_a_chart_the_spec_could_not_state_is_left_with_its_reason(tmp_path, series, why):
    assert chart_block(_chart(_deck(tmp_path, T.COLUMN_CLUSTERED, series))) == (None, why)


def test_a_scatter_charts_points_are_not_read_back(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = XyChartData()
    data.add_series("S").add_data_point(1, 2)
    slide.shapes.add_chart(T.XY_SCATTER, Inches(1), Inches(1), Inches(4), Inches(3), data)
    path = tmp_path / "xy.pptx"
    prs.save(str(path))
    assert chart_block(_chart(path)) == (None, "a xy-scatter chart, whose points are not read back")


def test_a_strangers_chart_is_drafted_as_a_chart_and_builds_with_the_same_numbers(tmp_path):
    from deckwright.compile import build_deck
    from deckwright.spec.draft import draft_spec
    from deckwright.spec.extract import harvest

    series = [("Ads", (12, 20, 25, 30)), ("Organic", (8, 20, 39, 62))]
    source = _deck(tmp_path, T.COLUMN_STACKED, series)
    content = harvest(source)
    assert content[0].dropped == ()
    text = draft_spec(content, title="Charts", theme="base").text
    (slide,) = [d for d in yaml.safe_load_all(text) if isinstance(d, dict)][1:]
    (entry,) = slide["place"]
    assert entry["chart"]["kind"] == "column-stacked"

    spec = tmp_path / "charts.deck.yaml"
    spec.write_text(text.replace("out: out/charts/Charts v1.pptx", "out: out/Charts.pptx"))
    rebuilt = build_deck(spec)
    again = _chart(rebuilt.deck).chart
    assert [list(s.values) for s in again.plots[0].series] == [
        [12.0, 20.0, 25.0, 30.0],
        [8.0, 20.0, 39.0, 62.0],
    ]


def test_a_chart_that_cannot_be_read_back_is_named_with_its_reason(tmp_path):
    from deckwright.spec.extract import harvest

    source = _deck(tmp_path, T.COLUMN_CLUSTERED, [("A", (1, None, 3, 4))])
    (dropped,) = harvest(source)[0].dropped
    assert dropped.startswith("chart 'Chart ") and dropped.endswith(
        " — a chart with an empty value"
    )


def test_charts_the_slide_has_no_room_for_are_named_in_the_draft(tmp_path):
    """Each takes half the grid; the third is named, with its numbers, rather than dropped."""
    from deckwright.spec.draft import draft_spec
    from deckwright.spec.extract import harvest

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for n in range(3):
        data = CategoryChartData()
        data.categories = ["a", "b"]
        data.add_series("S", (n + 1, n + 2))
        slide.shapes.add_chart(
            T.COLUMN_CLUSTERED, Inches(1 + 3 * n), Inches(1), Inches(3), Inches(3), data
        )
    path = tmp_path / "three.pptx"
    prs.save(str(path))

    drafted = draft_spec(harvest(path), title="Three", theme="base")
    (doc,) = [d for d in yaml.safe_load_all(drafted.text) if isinstance(d, dict)][1:]
    assert [entry["chart"]["data"][0]["value"] for entry in doc["place"]] == [1, 2]
    assert drafted.spilled == 1
    assert "#   a column chart — a: 3, b: 4" in drafted.text


def test_a_chart_whose_numbers_are_only_formulas_is_named_not_drafted_empty(tmp_path):
    shape = _chart(_deck(tmp_path, T.COLUMN_CLUSTERED, [("Sales", (1, 2, 3, 4))]))
    for cache in list(shape.chart._chartSpace.iter(qn("c:numCache"))):
        cache.getparent().remove(cache)
    assert chart_block(shape) == (None, "a chart whose cached values do not cover its categories")
