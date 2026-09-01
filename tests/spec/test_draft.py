"""Turning harvested words into a draft spec that builds."""

from __future__ import annotations

from pathlib import PurePosixPath

import pytest
import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from deckwright.compile import build_deck
from deckwright.compile.scaffold import new_deck
from deckwright.spec.draft import draft_spec, render_markdown, render_spec
from deckwright.spec.extract import SlideContent, harvest


def _docs(text: str) -> list:
    return list(yaml.safe_load_all(text))


def _built(tmp_path, text: str):
    spec = tmp_path / "draft.deck.yaml"
    spec.write_text(text, encoding="utf-8")
    return build_deck(spec)


def test_the_draft_config_names_the_theme_the_title_and_where_it_builds_to():
    config = _docs(render_spec([SlideContent(index=1, title="Hi")], title="Deck"))[0]
    assert list(config) == ["theme", "title", "out"]
    assert (config["theme"], config["title"]) == ("base", "Deck")


def test_a_block_becomes_full_width_bullets():
    slides = [SlideContent(index=1, title="Hi", blocks=(("one", "two"),))]
    slide = _docs(render_spec(slides, title="Deck"))[1]
    assert slide["title"] == "Hi"
    assert slide["place"] == [{"at": {"cols": "full"}, "bullets": {"items": ["one", "two"]}}]


def test_a_table_keeps_its_first_row_as_the_header():
    slides = [SlideContent(index=1, tables=((("Region", "Growth"), ("EMEA", "4%")),))]
    slide = _docs(render_spec(slides, title="Deck"))[1]
    assert slide["place"][0]["table"] == {
        "header": ["Region", "Growth"],
        "rows": [["EMEA", "4%"]],
    }


def test_each_placement_on_a_crowded_slide_takes_its_own_rows():
    slides = [SlideContent(index=1, blocks=(("one",), ("two",), ("three",)))]
    place = _docs(render_spec(slides, title="Deck"))[1]["place"]
    assert [p["at"]["rows"] for p in place] == [
        {"from": 0, "to": 4},
        {"from": 4, "to": 8},
        {"from": 8, "to": 12},
    ]


def test_more_blocks_than_rows_run_together_rather_than_into_empty_bands():
    slides = [SlideContent(index=1, blocks=tuple((f"line {i}",) for i in range(13)))]
    place = _docs(render_spec(slides, title="Deck"))[1]["place"]
    assert len(place) == 1
    assert place[0]["bullets"]["items"] == [f"line {i}" for i in range(13)]


def test_a_title_yaml_would_read_as_syntax_is_dumped_rather_than_interpolated():
    """`title: Q4: the year` is not a mapping value — an f-string here writes a spec
    that will not parse."""
    nasty = "Q4: the year - in review #1"
    text = render_spec([SlideContent(index=1, title=nasty)], title="Deck")
    assert f"title: {nasty}" not in text
    assert _docs(text)[1]["title"] == nasty


def test_what_was_dropped_is_written_into_the_draft_as_a_comment():
    slides = [SlideContent(index=1, title="Hi", dropped=("picture 'Logo'",))]
    text = render_spec(slides, title="Deck")
    assert "# not converted: picture 'Logo'" in text
    assert _docs(text)[1] == {"title": "Hi"}, "it must be a comment, not a field"


def test_the_transcript_carries_slides_titles_bullets_tables_notes_and_losses():
    slides = [
        SlideContent(
            index=1,
            title="Revenue is up",
            blocks=(("North America grew 12%",),),
            tables=((("Region", "Growth"), ("EMEA", "4%")),),
            notes="Lead with the number.",
            dropped=("chart 'Chart 3'",),
        )
    ]
    text = render_markdown(slides, title="Deck")
    for needle in (
        "## Slide 1",
        "### Revenue is up",
        "- North America grew 12%",
        "| Region | Growth |",
        "> Lead with the number.",
        "chart 'Chart 3'",
    ):
        assert needle in text, needle


@pytest.fixture
def six_row_theme(tmp_path, monkeypatch):
    """A theme declaring half the built-in grid's rows — ``scale: {rows: N}``."""
    (tmp_path / "six.theme.yaml").write_text("name: six\nscale: {rows: 6}\n", encoding="utf-8")
    monkeypatch.setenv("DECKWRIGHT_THEME_DIR", str(tmp_path))
    return "six"


def test_the_draft_bands_over_the_rows_the_named_theme_declares(six_row_theme):
    slides = [SlideContent(index=1, blocks=(("one",), ("two",)))]
    place = _docs(render_spec(slides, title="Deck", theme=six_row_theme))[1]["place"]
    assert [p["at"]["rows"] for p in place] == [{"from": 0, "to": 3}, {"from": 3, "to": 6}]


def test_a_draft_against_a_theme_of_fewer_rows_builds(tmp_path, six_row_theme):
    """Spans indexing twelve rows are out of range on a six-row grid."""
    slides = [
        SlideContent(
            index=1,
            blocks=(("North America grew 12%",), ("EMEA grew 4%",)),
            tables=((("Region", "Growth"), ("EMEA", "4%")),),
        )
    ]
    built = _built(tmp_path, render_spec(slides, title="Deck", theme=six_row_theme))
    assert built.deck.is_file()


def test_a_list_longer_than_its_band_takes_a_second_column():
    """Twenty-eight bullets need 9.64in of the base theme's 5.40in content band."""
    slides = [SlideContent(index=1, blocks=(tuple(f"line {i}" for i in range(28)),))]
    place = _docs(render_spec(slides, title="Deck"))[1]["place"]
    assert place[0]["bullets"]["columns"] == 2
    assert place[0]["bullets"]["items"] == [f"line {i}" for i in range(28)]


def test_bullets_past_three_columns_spill_into_the_draft_as_text():
    slides = [SlideContent(index=1, blocks=(tuple(f"line {i}" for i in range(60)),))]
    text = render_spec(slides, title="Deck")
    bullets = _docs(text)[1]["place"][0]["bullets"]
    assert bullets["columns"] == 3
    assert bullets["items"] == [f"line {i}" for i in range(42)]
    assert "#   line 42\n" in text
    assert "#   line 59\n" in text


def test_the_count_the_command_prints_is_the_number_of_spilled_lines():
    """What stdout reports has to be what the file holds as a comment."""
    slides = [SlideContent(index=1, blocks=(tuple(f"line {i}" for i in range(60)),))]
    drafted = draft_spec(slides, title="Deck")
    assert drafted.spilled == 18
    assert drafted.text.count("#   line ") == 18


def test_a_draft_of_a_slide_of_more_words_than_the_band_holds_builds(tmp_path):
    slides = [SlideContent(index=1, blocks=(tuple(f"line {i}" for i in range(28)),))]
    assert _built(tmp_path, render_spec(slides, title="Deck")).deck.is_file()


@pytest.mark.parametrize(
    "deep, kept",
    [
        (2, [{"from": i, "to": i + 2} for i in range(0, 12, 2)]),
        (3, [{"from": i, "to": i + 3} for i in range(0, 12, 3)]),
        (4, [{"from": i, "to": i + 4} for i in range(0, 12, 4)]),
    ],
)
def test_a_table_is_given_the_rows_its_own_height_asks_for(deep, kept):
    """A base-theme row is 0.90in for two, so a three-row table needs a third."""
    grids = tuple(tuple((f"head {i}", f"r{r}") for r in range(deep)) for i in range(13))
    text = render_spec([SlideContent(index=1, tables=grids)], title="Deck")
    place = _docs(text)[1]["place"]
    assert [p["at"]["rows"] for p in place] == kept
    assert [p["table"]["header"] for p in place] == [[f"head {i}", "r0"] for i in range(len(kept))]
    assert f"#   head {len(kept)} | r0\n" in text


@pytest.mark.parametrize("deep", [2, 3, 4])
def test_a_draft_of_a_slide_of_more_tables_than_rows_builds(tmp_path, deep):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Thirteen tables"
    for i in range(13):
        frame = slide.shapes.add_table(deep, 1, Inches(1), Inches(1), Inches(2), Inches(0.4))
        for row in range(deep):
            frame.table.cell(row, 0).text = f"cell {i}.{row}"
    source = tmp_path / "tables.pptx"
    prs.save(source)

    text = render_spec(harvest(source), title="Deck")
    assert "cell 12.0" in text
    assert _built(tmp_path, text).deck.is_file()


def test_a_header_only_table_becomes_bullets_rather_than_an_empty_body():
    """``rows:`` must be a non-empty list, and a one-row grid has no body rows to give it."""
    slides = [SlideContent(index=1, tables=((("Region", "Growth", "Share"),),))]
    place = _docs(render_spec(slides, title="Deck"))[1]["place"]
    assert place == [{"at": {"cols": "full"}, "bullets": {"items": ["Region", "Growth", "Share"]}}]


def test_a_draft_of_a_slide_carrying_a_one_row_table_builds(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "A label band"
    table = slide.shapes.add_table(1, 3, Inches(1), Inches(2), Inches(6), Inches(0.4)).table
    for column, text in enumerate(("Region", "Growth", "Share")):
        table.cell(0, column).text = text
    source = tmp_path / "band.pptx"
    prs.save(source)

    built = _built(tmp_path, render_spec(harvest(source), title="Deck"))
    assert built.deck.is_file()


@pytest.fixture
def deck_of_control_characters(tmp_path):
    """A Shift+Enter soft break and a two-paragraph cell, both far enough down the
    slide to land in the spilled tail rather than in a placement."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Spilled tail"
    frame = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(4)).text_frame
    frame.text = "line 0"
    for i in range(1, 60):
        frame.add_paragraph().text = f"line {i}\x0bsoft {i}" if i == 50 else f"line {i}"
    for i in range(13):
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(0.4)).table
        table.cell(0, 0).text = f"head {i}"
        table.cell(0, 1).text = "b"
        table.cell(1, 0).text = f"cell {i}"
        cell = table.cell(1, 1)
        cell.text = "para one"
        cell.text_frame.add_paragraph().text = "para two"
    path = tmp_path / "controls.pptx"
    prs.save(path)
    return path


def test_a_spilled_soft_break_and_multi_paragraph_cell_leave_a_draft_that_builds(
    tmp_path, deck_of_control_characters
):
    """``\\x0b`` and ``\\n`` both end a YAML comment where they stand."""
    text = render_spec(harvest(deck_of_control_characters), title="Deck")
    assert "#   line 50 soft 50\n" in text
    assert "#   cell 12 | para one para two\n" in text
    assert _built(tmp_path, text).deck.is_file()


@pytest.mark.parametrize(
    "name, directory",
    [
        ("Q3 Review: Plan", "q3-review-plan"),
        ("ML-pipeline", "ml-pipeline"),
        ("Ünïcode Deck", "ncode-deck"),
    ],
)
def test_a_deck_name_lands_in_the_same_directory_whichever_command_wrote_it(
    tmp_path, name, directory
):
    spec = new_deck(name, root=tmp_path, build=False).spec.read_text(encoding="utf-8")
    scaffolded = yaml.safe_load(spec.split("\n---\n")[0])
    drafted = _docs(render_spec([SlideContent(index=1)], title=name))[0]
    assert PurePosixPath(scaffolded["out"]).parent.name == directory
    assert PurePosixPath(drafted["out"]).parent.name == directory


def test_the_draft_names_an_output_path_so_it_can_be_built():
    config = _docs(render_spec([SlideContent(index=1)], title="Deck"))[0]
    assert config["out"] == "out/deck/Deck v1.pptx"


def test_a_kicker_reaches_the_draft_and_the_transcript():
    slides = [SlideContent(index=1, kicker="PART ONE", title="Revenue is up")]
    assert _docs(render_spec(slides, title="Deck"))[1]["kicker"] == "PART ONE"
    assert "**PART ONE**" in render_markdown(slides, title="Deck")


def test_the_bullet_budget_shrinks_under_the_chrome_the_slide_carries():
    """Stacked chrome pushes the content band down, so fewer bullets fit beneath it."""
    items = tuple(f"line {n}" for n in range(60))
    bare = _docs(render_spec([SlideContent(index=1, blocks=(items,))], title="Deck"))[1]
    crowned = _docs(
        render_spec(
            [
                SlideContent(
                    index=1,
                    kicker="PART ONE",
                    title=(
                        "Revenue is up across every region we sell into, and the number "
                        "that matters is the one nobody reported last quarter"
                    ),
                    subtitle=(
                        "And the line beneath it, running long enough that it wraps onto "
                        "a second line of its own as well as pushing the band down"
                    ),
                    blocks=(items,),
                )
            ],
            title="Deck",
        )
    )[1]
    assert len(crowned["place"][0]["bullets"]["items"]) < len(bare["place"][0]["bullets"]["items"])


def test_a_draft_of_a_slide_carrying_chrome_and_a_long_list_builds(tmp_path):
    """The band the budget measures has to be the one the compiler will hand the
    bullets, or a full list overflows it."""
    slides = [
        SlideContent(
            index=1,
            kicker="CI",
            title="Six gates",
            subtitle=(
                "Diff-based against the merge-base, so only lines your branch adds can "
                "fail. One arch-ignore comment opts a line out."
            ),
            blocks=(tuple(f"gate {n}" for n in range(15)),),
        )
    ]
    built = _built(tmp_path, render_spec(slides, title="Deck", out=str(tmp_path / "rt.pptx")))
    assert built.deck.is_file()


def test_a_deck_deckwright_built_extracts_its_own_chrome_and_builds_again(tmp_path):
    """The round trip that only the shape-name convention can carry: a built deck has
    no title placeholder anywhere in it."""
    spec = tmp_path / "source.deck.yaml"
    spec.write_text(
        "theme: base\ntitle: Source\nout: source.pptx\n"
        "---\n"
        "kicker: PART ONE\n"
        "title: Revenue is up\n"
        "subtitle: And the line beneath it\n"
        "place: [{at: {cols: full}, bullets: {items: [North America grew 12%]}}]\n",
        encoding="utf-8",
    )
    built = build_deck(spec)

    slide = harvest(built.deck)[0]
    assert (slide.kicker, slide.title, slide.subtitle) == (
        "PART ONE",
        "Revenue is up",
        "And the line beneath it",
    )
    assert _built(
        tmp_path, render_spec([slide], title="Round Trip", out=str(tmp_path / "rt.pptx"))
    ).deck.is_file()


def test_a_colour_no_pair_owns_is_named_in_a_comment_and_not_counted_as_spilled():
    """Painted in a colour this theme never declared: a note beside the slide, not a
    line that failed to fit."""
    drafted = draft_spec([SlideContent(index=1, title="T", background_rgb="ABCDEF")], title="Deck")
    assert "# this slide was painted ABCDEF; name a background pair\n" in drafted.text
    assert "did not fit" not in drafted.text
    assert drafted.spilled == 0
    assert _docs(drafted.text)[1] == {"title": "T"}


@pytest.mark.parametrize("rgb", ["FFFFFF", "#ffffff"])
def test_the_page_colour_is_the_default_and_writes_no_background(rgb):
    """``page-muted`` shares the page's surface; naming it would set the whole slide's
    ink to the muted colour."""
    drafted = draft_spec([SlideContent(index=1, title="T", background_rgb=rgb)], title="Deck")
    assert _docs(drafted.text)[1] == {"title": "T"}
    assert "was painted" not in drafted.text


def test_a_foreign_slide_painted_white_drafts_as_an_ordinary_page(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Hello"
    source = tmp_path / "white.pptx"
    prs.save(source)

    text = render_spec(harvest(source), title="Deck")
    assert "background" not in text
    assert _docs(text)[1]["place"][0]["bullets"]["items"] == ["Hello"]
