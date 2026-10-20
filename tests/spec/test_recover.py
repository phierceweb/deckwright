"""A deckwright deck drafted back as the placements that built it, not as its words alone."""

from __future__ import annotations

import re

import pytest
import yaml

from deckwright.compile import build_deck
from deckwright.compile.readback import read_back
from deckwright.spec.draft import draft_spec, render_markdown
from deckwright.spec.extract import harvest

SOURCE = r"""theme: testtheme
title: Source
out: out/Source.pptx
---
title: Cards and a list
place:
  - at: {cols: left-half, rows: {from: 0, to: 5}}
    id: q
    card: {heading: The question, body: What went wrong}
  - at: {cols: right-half, rows: {from: 0, to: 5}}
    card: {heading: The answer, body: One bug.}
  - at: {cols: full, rows: {from: 6, to: 12}}
    bullets: {heading: Why, items: [alpha, beta, gamma]}
---
title: A quote under a rule
animate: together
place:
  - at: {cols: full, rows: {from: 0, to: 1}}
    rule: {}
  - at: {cols: full, rows: {from: 1, to: 9}}
    prose: {paragraphs: [First paragraph., Second paragraph.], cite: Someone}
---
title: A table
place:
  - at: {cols: full}
    table: {header: [A, B], rows: [['1', '2'], ['3', '4']]}
---
title: A chart built by category
animate: by_category
place:
  - at: {cols: full}
    chart:
      kind: column
      data:
        - {category: Q1, value: 12}
        - {category: Q2, value: 34}
        - {category: Q3, value: 58}
        - {category: Q4, value: 91}
---
title: Two series
animate: by_series
place:
  - at: {cols: full}
    chart:
      kind: column-stacked
      data:
        - {category: Q1, values: {Ads: 12, Organic: 8}}
        - {category: Q2, values: {Ads: 20, Organic: 20}}
        - {category: Q3, values: {Ads: 25, Organic: 39}}
        - {category: Q4, values: {Ads: 30, Organic: 62}}
---
title: Staged
animate: one_at_a_time
place:
  - at: {cols: full}
    bullets: {items: [one, two, three]}
---
title: Click the question
place:
  - at: {cols: left-half, rows: {from: 0, to: 6}}
    id: ask
    card: {heading: Q, body: Click me}
  - at: {cols: right-half, rows: {from: 0, to: 6}}
    reveals: ask
    card: {heading: A, body: Shown}
---
title: Numbers
place:
  - at: {cols: full, rows: {from: 0, to: 6}}
    stats: {items: [{value: '12', label: things}, {value: '34', label: more things}]}
---
title: Before
place:
  - at: {cols: {from: 0, to: 5}, rows: {from: 0, to: 6}}
    morph: fig
    card: {heading: Revenue, body: Up}
---
title: After
transition: morph
place:
  - at: {cols: {from: 3, to: 12}, rows: {from: 0, to: 8}}
    morph: fig
    card: {heading: Revenue, body: Up}
---
title: Single lines, two columns, a panel
place:
  - at: {cols: {from: 0, to: 4}, rows: {from: 0, to: 4}}
    card: {heading: Only a heading}
  - at: {cols: {from: 4, to: 8}, rows: {from: 0, to: 4}}
    card: {body: Only a body}
  - at: {cols: {from: 8, to: 9}, rows: {from: 0, to: 5}}
    rule: {orient: vertical}
  - at: {cols: {from: 9, to: 12}, rows: {from: 0, to: 4}}
    panel: {}
  - at: {cols: full, rows: {from: 6, to: 12}}
    bullets: {columns: 2, items: [a, b, c, d]}
---
title: A morph trigger
place:
  - at: {cols: left-half, rows: {from: 0, to: 6}}
    id: go
    morph: shared
    card: {heading: Go, body: Click}
  - at: {cols: right-half, rows: {from: 0, to: 6}}
    reveals: go
    card: {heading: Gone, body: Shown}
"""


@pytest.fixture
def source(tmp_path, theme_file):
    spec = tmp_path / "source.deck.yaml"
    spec.write_text(SOURCE)
    return build_deck(spec, theme_path=theme_file)


@pytest.fixture
def drafted(source, tmp_path):
    text = draft_spec(
        harvest(source.deck), title="Draft", theme=str(tmp_path / "testtheme.yaml")
    ).text
    docs = [d for d in yaml.safe_load_all(text) if isinstance(d, dict)]
    return text, docs[1:]


def _component(entry: dict) -> str:
    return next(k for k in entry if k not in ("at", "id", "morph", "reveals", "bleed"))


def test_each_placement_comes_back_as_its_component_with_its_fields(drafted):
    _, slides = drafted
    first = slides[0]["place"]
    assert [_component(e) for e in first] == ["card", "card", "bullets"]
    assert first[0]["id"] == "q" and first[0]["card"] == {
        "heading": "The question",
        "body": "What went wrong",
    }
    assert "id" not in first[1]
    assert first[2]["bullets"] == {"heading": "Why", "items": ["alpha", "beta", "gamma"]}
    assert all(set(e["at"]) == {"box"} for e in first)
    second = slides[1]["place"]
    assert [_component(e) for e in second] == ["rule", "prose"]
    assert second[1]["prose"] == {
        "paragraphs": ["First paragraph.", "Second paragraph."],
        "cite": "Someone",
    }
    assert slides[2]["place"][0]["table"] == {
        "header": ["A", "B"],
        "rows": [["1", "2"], ["3", "4"]],
    }


def test_a_chart_comes_back_with_its_kind_and_its_data(drafted):
    _, slides = drafted
    assert slides[3]["place"][0]["chart"] == {
        "kind": "column",
        "data": [
            {"category": "Q1", "value": 12},
            {"category": "Q2", "value": 34},
            {"category": "Q3", "value": 58},
            {"category": "Q4", "value": 91},
        ],
    }
    assert slides[4]["place"][0]["chart"]["kind"] == "column-stacked"
    assert slides[4]["place"][0]["chart"]["data"][3] == {
        "category": "Q4",
        "values": {"Ads": 30, "Organic": 62},
    }


def test_the_slides_motion_comes_back(drafted):
    _, slides = drafted
    assert slides[1]["animate"] == "together"
    assert slides[4]["animate"] == "by_series"
    assert slides[3]["animate"] == "by_category"
    assert slides[5]["animate"] == "one_at_a_time"
    click = slides[6]["place"]
    assert click[0]["id"] == "ask" and click[1]["reveals"] == "ask" and "animate" not in slides[6]
    assert slides[9]["transition"] == "morph"
    assert slides[8]["place"][0]["morph"] == "fig" and slides[9]["place"][0]["morph"] == "fig"


def test_a_component_whose_fields_are_not_in_the_file_keeps_its_words_and_says_so(drafted):
    text, slides = drafted
    assert _component(slides[7]["place"][0]) == "bullets"
    assert "12" in slides[7]["place"][0]["bullets"]["items"]
    assert "# s8.p1.stats: drafted as bullets" in text


def test_the_draft_rebuilds_into_the_deck_it_was_read_from(source, drafted, tmp_path, theme_file):
    """Read back against the source's own manifest, the rebuild differs only where the draft
    says it could not recover a component."""
    text, _ = drafted
    spec = tmp_path / "draft.deck.yaml"
    spec.write_text(text.replace("out: out/draft/Draft v1.pptx", "out: out/Rebuilt.pptx"))
    rebuilt = build_deck(spec, theme_path=theme_file)
    drift = read_back(rebuilt.deck, manifest=source.manifest)
    unexpected = [c for c in drift.changes if not c.shape.startswith("s8.p1.")]
    assert unexpected == []


def test_single_line_cards_columns_a_vertical_rule_and_a_panel_come_back(drafted):
    _, slides = drafted
    place = slides[10]["place"]
    assert [_component(e) for e in place] == ["card", "card", "rule", "panel", "bullets"]
    assert place[0]["card"] == {"heading": "Only a heading"}
    assert place[1]["card"] == {"body": "Only a body"}
    assert place[2]["rule"] == {"orient": "vertical"}
    assert place[3]["panel"] == {}
    assert place[4]["bullets"] == {"columns": 2, "items": ["a", "b", "c", "d"]}


def test_a_morph_placement_that_triggers_a_reveal_is_given_the_id_it_is_named_by(drafted):
    _, slides = drafted
    trigger, target = slides[11]["place"]
    assert (trigger["morph"], trigger["id"], target["reveals"]) == ("shared", "shared", "shared")


def test_a_placement_after_a_gap_keeps_the_number_it_was_drawn_under(source, tmp_path):
    """Delete the first card by hand: the second is still `p2`, so the draft names it."""
    from pptx import Presentation

    prs = Presentation(str(source.deck))
    tree = prs.slides[0].shapes._spTree
    for shape in list(prs.slides[0].shapes):
        if shape.name.startswith("s1.q."):
            tree.remove(shape._element)
    edited = tmp_path / "edited.pptx"
    prs.save(str(edited))

    placed = harvest(edited)[0].placed
    assert [(item.key, item.id) for item in placed] == [("p2", "p2"), ("p3", "p3")]


def test_a_placement_read_back_is_not_also_harvested_as_loose_words(source):
    """Its words would come back twice: in the placement, and as lines no placement drew."""
    assert [slide.blocks for slide in harvest(source.deck)] == [()] * 12


def test_the_transcript_carries_what_the_placements_read_back_hold(source):
    text = render_markdown(harvest(source.deck), title="Source")
    for words in ("- The question", "- What went wrong", "- alpha", "- First paragraph.", "- 12"):
        assert words in text, words
    assert "| A | B |" in text and "| 3 | 4 |" in text
    assert "| Q4 | 91 |" in text


PICTURES = r"""theme: base
title: Pictures
out: out/Pictures.pptx
---
title: Two pictures
place:
  - at: {cols: left-half}
    image: {src: pale.png, alt: A pale field, over: [Words on it]}
  - at: {cols: right-half}
    id: plain
    image: {src: pale.png, decorative: true}
"""


def test_an_image_comes_back_with_its_file_and_rebuilds_from_the_draft(tmp_path):
    from PIL import Image
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    Image.new("RGB", (800, 450), (250, 250, 245)).save(tmp_path / "pale.png")
    spec = tmp_path / "pictures.deck.yaml"
    spec.write_text(PICTURES)
    built = build_deck(spec)
    shapes = Presentation(str(built.deck)).slides[0].shapes
    assert len([s for s in shapes if s.name.startswith("s1.p1.")]) == 3  # picture, scrim, words

    drafted = draft_spec(harvest(built.deck), title="Draft", theme="base")
    (slide,) = [d for d in yaml.safe_load_all(drafted.text) if isinstance(d, dict)][1:]
    first, second = slide["place"]
    assert first["image"] == {
        "src": "draft.media/slide1-p1.png",
        "alt": "A pale field",
        "over": ["Words on it"],
    }
    assert second["id"] == "plain"
    assert second["image"] == {"src": "draft.media/slide1-plain.png", "decorative": True}
    note = "# s1.p1.image: its scrim is solved afresh"
    assert note in drafted.text
    assert sorted(drafted.media) == ["draft.media/slide1-p1.png", "draft.media/slide1-plain.png"]
    transcript = render_markdown(harvest(built.deck), title="Pictures")
    assert "*picture: A pale field*\n\n- Words on it\n" in transcript
    assert "*picture, with no alt text*" in transcript

    home = tmp_path / "again"
    for name, blob in drafted.media.items():
        (home / name).parent.mkdir(parents=True, exist_ok=True)
        (home / name).write_bytes(blob)
    (home / "draft.deck.yaml").write_text(
        drafted.text.replace("out: out/draft/Draft v1.pptx", "out: out/Draft.pptx")
    )
    rebuilt = build_deck(home / "draft.deck.yaml").deck
    assert read_back(rebuilt, manifest=built.manifest).changes == ()
    original = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    again = Presentation(str(rebuilt)).slides[0].shapes
    pictures = [s for s in again if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert [p.image.blob for p in pictures] == [p.image.blob for p in original]


def test_a_card_a_chart_and_a_picture_added_by_hand_are_named_for_the_author(source, tmp_path):
    """No band is free on a slide its placements already fill, so each is named instead."""
    import io

    from PIL import Image
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    prs = Presentation(str(source.deck))
    slide = prs.slides[0]
    plate = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1))
    plate.fill.solid()
    plate.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
    plate.text_frame.text = "Added"
    plate.text_frame.add_paragraph().text = "by hand"
    for p in plate.text_frame.paragraphs:
        p.runs[0].font.size = Pt(14)
    data = CategoryChartData()
    data.categories = ["a", "b"]
    data.add_series("S", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1), Inches(3), Inches(2), data
    )
    png = io.BytesIO()
    Image.new("RGB", (40, 20), "red").save(png, format="PNG")
    slide.shapes.add_picture(png, Inches(9), Inches(1), Inches(2), Inches(1))
    edited = tmp_path / "edited.pptx"
    prs.save(str(edited))

    drafted = draft_spec(harvest(edited)[:1], title="Draft", theme="base")
    for line in (
        "Added",
        "by hand",
        "a column chart — a: 1, b: 2",
        "a picture — draft.media/slide1-1.png",
    ):
        assert f"#   {line}\n" in drafted.text, line
    assert list(drafted.media) == ["draft.media/slide1-1.png"]


def test_a_placement_pasted_twice_under_one_name_is_not_merged_into_one(source, tmp_path):
    """PowerPoint keeps a copy's name, so one name now covers two cards' shapes."""
    import copy

    from pptx import Presentation

    prs = Presentation(str(source.deck))
    tree = prs.slides[0].shapes._spTree
    for shape in [s for s in prs.slides[0].shapes if s.name.startswith("s1.q.")]:
        twin = copy.deepcopy(shape._element)
        twin.xpath("./*[1]/p:cNvPr")[0].set("id", str(900 + len(tree)))
        off = twin.xpath(".//a:xfrm/a:off")[0]
        off.set("y", str(int(off.get("y")) + 3_000_000))
        tree.append(twin)
    edited = tmp_path / "pasted.pptx"
    prs.save(str(edited))

    content = harvest(edited)[0]
    assert "q" not in [item.key for item in content.placed]
    question = {"heading": "The question", "body": "What went wrong"}
    assert content.cards == (question, question)


def test_a_placement_grouped_by_hand_is_still_read_back(source, tmp_path):
    from pptx import Presentation

    prs = Presentation(str(source.deck))
    slide = prs.slides[0]
    group = slide.shapes.add_group_shape()
    for shape in [s for s in slide.shapes if s.name.startswith("s1.q.")]:
        group._element.append(shape._element)
    edited = tmp_path / "grouped.pptx"
    prs.save(str(edited))

    (question,) = [item for item in harvest(edited)[0].placed if item.key == "q"]
    assert question.body == {"card": {"heading": "The question", "body": "What went wrong"}}


PAIRS = r"""theme: base
title: Pairs
out: out/Pairs.pptx
---
title: Painted plates
place:
  - at: {cols: {from: 0, to: 4}, rows: {from: 0, to: 6}}
    card: {pair: accent-1, heading: Blue, body: On accent one}
  - at: {cols: {from: 4, to: 8}, rows: {from: 0, to: 6}}
    card: {heading: Plain, body: On the default}
  - at: {cols: {from: 8, to: 12}, rows: {from: 0, to: 6}}
    panel: {pair: inverse}
  - at: {cols: {from: 0, to: 4}, rows: {from: 7, to: 12}}
    card: {pair: page-muted, body: Grey words on white}
"""


def test_a_plates_pair_comes_back_when_the_theme_declares_its_colours(tmp_path):
    spec = tmp_path / "pairs.deck.yaml"
    spec.write_text(PAIRS)
    built = build_deck(spec)
    drafted = draft_spec(harvest(built.deck), title="Draft", theme="base")
    (slide,) = [d for d in yaml.safe_load_all(drafted.text) if isinstance(d, dict)][1:]
    blue, plain, panel, muted = slide["place"]
    assert blue["card"] == {"pair": "accent-1", "heading": "Blue", "body": "On accent one"}
    assert plain["card"] == {"heading": "Plain", "body": "On the default"}
    assert panel["panel"] == {"pair": "inverse"}
    assert muted["card"] == {"pair": "page-muted", "body": "Grey words on white"}

    again = tmp_path / "draft.deck.yaml"
    again.write_text(drafted.text.replace("out: out/draft/Draft v1.pptx", "out: out/Draft.pptx"))
    assert read_back(build_deck(again).deck, manifest=built.manifest).changes == ()


def test_a_plate_filled_by_hand_with_no_solid_colour_comes_back_without_a_pair(tmp_path):
    from pptx import Presentation

    spec = tmp_path / "pairs.deck.yaml"
    spec.write_text(PAIRS)
    prs = Presentation(str(build_deck(spec).deck))
    plate = next(s for s in prs.slides[0].shapes if s.name == "s1.p1.card#1")
    plate.fill.gradient()
    edited = tmp_path / "gradient.pptx"
    prs.save(str(edited))

    (blue, *_) = harvest(edited)[0].placed
    assert blue.painted is None
    assert blue.body == {"card": {"heading": "Blue", "body": "On accent one"}}


def test_a_placement_whose_file_holds_only_its_picture_comes_back_as_that_picture(tmp_path):
    """A `document:` is rendered to a picture: its source is gone, and the picture is not."""
    from PIL import Image
    from pptx import Presentation

    Image.new("RGB", (800, 450), (250, 250, 245)).save(tmp_path / "pale.png")
    spec = tmp_path / "pictures.deck.yaml"
    spec.write_text(PICTURES)
    prs = Presentation(str(build_deck(spec).deck))
    for shape in prs.slides[0].shapes:
        if shape.name.startswith("s1.plain."):
            shape.name = shape.name.replace(".image#", ".document#")
    edited = tmp_path / "document.pptx"
    prs.save(str(edited))

    drafted = draft_spec(harvest(edited), title="Draft", theme="base")
    (slide,) = [d for d in yaml.safe_load_all(drafted.text) if isinstance(d, dict)][1:]
    assert slide["place"][1]["image"] == {
        "src": "draft.media/slide1-plain.png",
        "decorative": True,
    }
    note = "# s1.plain.document: drafted as an image: a document placement's source is not"
    assert note in drafted.text


@pytest.mark.parametrize(
    "whole, note",
    [
        (False, "# s7.ask.card: part of it is revealed by p2, and a spec reveals whole placements"),
        (
            True,
            "# s7.ask.card: its reveal by p2 is left out, since it would close a ring of reveals",
        ),
    ],
)
def test_a_reveal_a_spec_cannot_state_is_named_and_the_draft_still_builds(
    source, tmp_path, theme_file, whole, note
):
    """Clicking the revealed card shows the question again: only its words, or all of it."""
    from pptx import Presentation

    prs = Presentation(str(source.deck))
    plate = "s7.ask.card#1" if whole else "s7.ask.card#2"
    reveal = {"s7.p2.card#1": plate, "s7.p2.card#2": "s7.ask.card#2"}
    _retrigger(prs.slides[6], trigger="s7.ask.card#1", now="s7.p2.card#2", reveal=reveal)
    edited = tmp_path / "ring.pptx"
    prs.save(str(edited))

    text = draft_spec(harvest(edited), title="Draft", theme=str(tmp_path / "testtheme.yaml")).text
    slide7 = [d for d in yaml.safe_load_all(text) if isinstance(d, dict)][7]
    ask, answer = slide7["place"]
    assert "reveals" not in ask and answer["reveals"] == "ask"
    assert note in text
    spec = tmp_path / "draft.deck.yaml"
    spec.write_text(text.replace("out: out/draft/Draft v1.pptx", "out: out/Rebuilt.pptx"))
    build_deck(spec, theme_path=theme_file)


def _retrigger(slide, *, trigger: str, now: str, reveal: dict[str, str]) -> None:
    """Copy each click sequence ``trigger`` starts, started by ``now`` instead and showing
    ``reveal[shape]`` for each shape it showed — a trigger added in PowerPoint."""
    import copy

    from pptx.oxml.ns import qn

    ids = {s.name: s.shape_id for s in slide.shapes}
    names = {str(spid): name for name, spid in ids.items()}
    for seq in list(slide._element.iter(qn("p:seq"))):
        timing = seq.find(qn("p:cTn"))
        start = timing.find(qn("p:stCondLst")).find(".//" + qn("p:spTgt"))
        if timing.get("nodeType") != "interactiveSeq" or start.get("spid") != str(ids[trigger]):
            continue
        back = copy.deepcopy(seq)
        back_timing = back.find(qn("p:cTn"))
        back_timing.find(qn("p:stCondLst")).find(".//" + qn("p:spTgt")).set("spid", str(ids[now]))
        for target in back_timing.find(qn("p:childTnLst")).iter(qn("p:spTgt")):
            target.set("spid", str(ids[reveal[names[target.get("spid")]]]))
        seq.addnext(back)


TRIGGERS = r"""theme: base
title: Triggers
out: out/Triggers.pptx
---
title: Two ways in
place:
  - at: {cols: {from: 0, to: 3}, rows: {from: 0, to: 6}}
    id: a
    card: {heading: A, body: Click}
  - at: {cols: {from: 3, to: 6}, rows: {from: 0, to: 6}}
    id: b
    card: {heading: B, body: Or click}
  - at: {cols: {from: 6, to: 8}, rows: {from: 0, to: 6}}
    id: go
    icon: {name: shield, decorative: true}
  - at: {cols: {from: 8, to: 12}, rows: {from: 0, to: 6}}
    reveals: a
    card: {heading: C, body: Shown}
"""


@pytest.mark.parametrize(
    "now, note",
    [
        ("s1.b.card#1", "# s1.p4.card: it is also revealed by b, and a spec takes one trigger"),
        ("s1.go.icon#1", "# s1.p4.card: it is also revealed by go, and a spec takes one trigger"),
    ],
)
def test_a_placement_with_a_second_trigger_keeps_the_first(tmp_path, now, note):
    from pptx import Presentation

    spec = tmp_path / "triggers.deck.yaml"
    spec.write_text(TRIGGERS)
    prs = Presentation(str(build_deck(spec).deck))
    same = {"s1.p4.card#1": "s1.p4.card#1", "s1.p4.card#2": "s1.p4.card#2"}
    _retrigger(prs.slides[0], trigger="s1.a.card#1", now=now, reveal=same)
    edited = tmp_path / "twice.pptx"
    prs.save(str(edited))

    text = draft_spec(harvest(edited), title="Draft", theme="base").text
    (slide,) = [d for d in yaml.safe_load_all(text) if isinstance(d, dict)][1:]
    assert slide["place"][-1]["reveals"] == "a"
    assert note in text
    again = tmp_path / "draft.deck.yaml"
    again.write_text(text.replace("out: out/draft/Draft v1.pptx", "out: out/Draft.pptx"))
    build_deck(again)


def test_a_hand_grouped_placement_moved_with_its_group_is_drafted_where_it_now_sits(
    source, tmp_path
):
    """Read from its shapes alone, with no manifest beside the copy: the group's move of
    (1in, 0.5in) is where the placement now is, not its shapes' own numbers."""
    from pptx import Presentation
    from pptx.util import Inches

    def drafted_box(move):
        prs = Presentation(str(source.deck))
        slide = prs.slides[0]
        group = slide.shapes.add_group_shape()
        for shape in [s for s in slide.shapes if s.name.startswith("s1.q.")]:
            group._element.append(shape._element)
        if move:
            xfrm = group._element.grpSpPr.get_or_add_xfrm()
            off, child = xfrm.get_or_add_off(), xfrm.get_or_add_chOff()
            off.x, off.y, child.x, child.y = Inches(1), Inches(0.5), 0, 0
            for size in (xfrm.get_or_add_ext(), xfrm.get_or_add_chExt()):
                size.cx, size.cy = prs.slide_width, prs.slide_height
        edited = tmp_path / f"grouped-{move}.pptx"
        prs.save(str(edited))
        (question,) = [item for item in harvest(edited)[0].placed if item.key == "q"]
        return question.box, (prs.slide_width, prs.slide_height)

    still, _ = drafted_box(False)
    moved, (width, height) = drafted_box(True)
    shift = (Inches(1) / width, Inches(0.5) / height, 0.0, 0.0)
    assert moved == pytest.approx(tuple(a + b for a, b in zip(still, shift, strict=True)))


ICONS = r"""theme: base
title: Icons
out: out/Icons.pptx
---
title: An icon and a card
place:
  - at: {cols: {from: 0, to: 4}, rows: {from: 0, to: 8}}
    id: key
    icon: {name: shield, ink: accent-2, size: 0.5, alt: A shield}
  - at: {cols: {from: 6, to: 12}, rows: {from: 0, to: 8}}
    reveals: key
    card: {heading: Locked, body: Safe, icon: shield}
"""


def _drafted_icons(tmp_path, spec_text=ICONS, theme=None):
    spec = tmp_path / "icons.deck.yaml"
    spec.write_text(spec_text)
    built = build_deck(spec)
    text = draft_spec(harvest(built.deck, theme=theme), title="Draft", theme="base").text
    return text, [d for d in yaml.safe_load_all(text) if isinstance(d, dict)][1:]


def test_an_icon_comes_back_as_the_glyph_it_draws(tmp_path):
    """Its name from its drawing, its ink from the role its fill is, its size against the
    placement's short side, and its alt text as a picture's is read."""
    _, (slide,) = _drafted_icons(tmp_path)
    (icon,) = [entry for entry in slide["place"] if "icon" in entry]
    assert icon["icon"] == {"name": "shield", "ink": "accent-2", "size": 0.5, "alt": "A shield"}
    assert icon["id"] == "key"


def test_a_cards_icon_comes_back_by_name_and_its_reveal_with_it(tmp_path):
    text, (slide,) = _drafted_icons(tmp_path)
    (card,) = [entry for entry in slide["place"] if "card" in entry]
    assert card["card"] == {"heading": "Locked", "body": "Safe", "icon": "shield"}
    assert card["reveals"] == "key"
    assert "not in the file by name" not in text and "not converted" not in text


def test_an_aliased_glyph_comes_back_under_the_first_of_its_names(tmp_path):
    """``mode_cool`` and ``ac_unit`` draw the same path; either rebuilds the same icon."""
    spec = ICONS.replace("name: shield, ink: accent-2", "name: mode_cool, ink: accent-2")
    _, (slide,) = _drafted_icons(tmp_path, spec)
    (icon,) = [entry for entry in slide["place"] if "icon" in entry]
    assert icon["icon"]["name"] == "ac_unit"


_MARK = (
    '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24">'
    '<path d="M2 2 L22 2 L12 22 Z"/></svg>'
)


def test_a_glyph_from_a_folder_the_draft_is_given_comes_back_and_one_it_is_not_stays_named(
    tmp_path, monkeypatch
):
    import dataclasses

    from deckwright.theme.load import load_theme

    folder = tmp_path / "brand-icons"
    folder.mkdir()
    (folder / "brandmark.svg").write_text(_MARK)
    monkeypatch.setenv("DECKWRIGHT_ICON_DIR", str(folder))
    spec = ICONS.replace(
        "name: shield, ink: accent-2, size: 0.5, alt: A shield", "name: brandmark, alt: Our mark"
    )
    spec_path = tmp_path / "brand.deck.yaml"
    spec_path.write_text(spec)
    built = build_deck(spec_path)
    monkeypatch.delenv("DECKWRIGHT_ICON_DIR")

    unknown = draft_spec(harvest(built.deck), title="Draft", theme="base").text
    assert "# alt text on freeform 's1.key.icon#1': Our mark" in unknown

    theme = dataclasses.replace(load_theme("base"), icons=folder)
    known = draft_spec(harvest(built.deck, theme=theme), title="Draft", theme="base").text
    (slide,) = [d for d in yaml.safe_load_all(known) if isinstance(d, dict)][1:]
    (icon,) = [entry for entry in slide["place"] if "icon" in entry]
    assert icon["icon"]["name"] == "brandmark"


def _pictures_draft(tmp_path, spec_text):
    from PIL import Image

    Image.new("RGB", (800, 450), (250, 250, 245)).save(tmp_path / "pale.png")
    spec = tmp_path / "pictures.deck.yaml"
    spec.write_text(spec_text)
    built = build_deck(spec)
    return built, draft_spec(harvest(built.deck), title="Draft", theme="base")


def test_an_over_line_on_another_rung_comes_back_on_it_and_rebuilds_with_no_drift(tmp_path):
    """Drafted at the default rung, a caption line rebuilt 14pt → 34pt."""
    built, drafted = _pictures_draft(
        tmp_path,
        PICTURES.replace(
            "over: [Words on it]", "over: [{text: Small, rung: caption}, Words on it]"
        ),
    )
    # Two lines share one frame; recorded once each, the unedited deck read back as edited.
    assert read_back(built.deck, manifest=built.manifest).changes == ()
    (slide,) = [d for d in yaml.safe_load_all(drafted.text) if isinstance(d, dict)][1:]
    assert slide["place"][0]["image"]["over"] == [
        {"text": "Small", "rung": "caption"},
        "Words on it",
    ]
    assert "default rung" not in drafted.text

    home = tmp_path / "again"
    for name, blob in drafted.media.items():
        (home / name).parent.mkdir(parents=True, exist_ok=True)
        (home / name).write_bytes(blob)
    (home / "draft.deck.yaml").write_text(
        drafted.text.replace("out: out/draft/Draft v1.pptx", "out: out/Draft.pptx")
    )
    rebuilt = build_deck(home / "draft.deck.yaml").deck
    assert read_back(rebuilt, manifest=built.manifest).changes == ()


def test_an_over_line_set_at_a_size_no_rung_has_is_named_in_the_note(tmp_path):
    from pptx import Presentation
    from pptx.util import Pt

    built, _ = _pictures_draft(tmp_path, PICTURES)
    prs = Presentation(str(built.deck))
    (words,) = [
        s
        for s in prs.slides[0].shapes
        if s.name.startswith("s1.p1.") and s.has_text_frame and s.text_frame.text
    ]
    for run in words.text_frame.paragraphs[0].runs:
        run.font.size = Pt(23)
    edited = tmp_path / "edited.pptx"
    prs.save(str(edited))

    drafted = draft_spec(harvest(edited), title="Draft", theme="base")
    (slide,) = [d for d in yaml.safe_load_all(drafted.text) if isinstance(d, dict)][1:]
    assert slide["place"][0]["image"]["over"] == ["Words on it"]
    assert (
        "# s1.p1.image: 'Words on it' is set at 23pt, which no rung is, so it comes back at the default rung"
        in drafted.text
    )


def test_a_size_two_rungs_share_takes_the_default_else_the_first_and_says_so():
    import dataclasses

    from deckwright.spec._themeback import _rung
    from deckwright.theme.load import load_theme

    base = load_theme("base")
    per_rung = base.ramp["title"].size / base.ramp["title"].rung

    def at(name, pt):
        return dataclasses.replace(base.ramp[name], rung=pt / per_rung)

    ramp = {"title": at("title", 24), "head": at("head", 24), "body": at("body", 16)}
    ramp["caption"] = at("caption", 16)
    theme = dataclasses.replace(base, ramp=ramp)
    assert _rung(24.2, theme, default="title") == ("title", None)
    assert _rung(16.0, theme, default="title") == (
        "body",
        "which body and caption share, so it is drafted as body",
    )
    assert _rung(30.0, theme, default="title") == (
        None,
        "which no rung is, so it comes back at the default rung",
    )


def _title_treatments(tmp_path, extra=""):
    from pathlib import Path

    example = Path(__file__).parents[2] / "examples" / "title-treatments.deck.yaml"
    text = re.sub(r"(?m)^out: .*$", "out: out/TT.pptx", example.read_text()) + extra
    spec = tmp_path / "tt.deck.yaml"
    spec.write_text(text)
    return build_deck(spec)


def test_a_slides_chrome_overrides_come_back_and_rebuild_where_they_were(tmp_path):
    """The tour's slide 1 drifted `moved` and `restyled` on its chrome and nowhere else."""
    built = _title_treatments(tmp_path, "---\ntitle: Only a title\n")
    drafted = draft_spec(harvest(built.deck), title="Draft", theme="base")
    slides = [d for d in yaml.safe_load_all(drafted.text) if isinstance(d, dict)][1:]
    assert ["chrome" in s for s in slides] == [False, True, True, True, True, True, False]
    assert slides[2]["chrome"]["title"]["rung"] == "display"
    assert slides[2]["chrome"]["kicker"]["anchor"] == "bottom"
    assert slides[1]["chrome"]["title"]["align"] == "center"

    again = tmp_path / "again"
    again.mkdir()
    (again / "d.deck.yaml").write_text(
        drafted.text.replace("out: out/draft/Draft v1.pptx", "out: out/Draft.pptx")
    )
    rebuilt = build_deck(again / "d.deck.yaml").deck
    drift = read_back(rebuilt, manifest=built.manifest).changes
    assert [c for c in drift if ".chrome" in c.shape] == []


def test_sections_come_back_from_the_manifest_and_not_without_it(tmp_path, theme_file):
    spec = tmp_path / "sectioned.deck.yaml"
    sectioned = SOURCE.replace("out: out/Source.pptx", "out: out/Sectioned.pptx\nsections: [A, B]")
    sectioned = sectioned.replace("title: Cards and a list", "section: A\ntitle: Cards and a list")
    sectioned = sectioned.replace("title: A table", "section: B\ntitle: A table")
    spec.write_text(sectioned)
    built = build_deck(spec, theme_path=theme_file)
    theme = str(tmp_path / "testtheme.yaml")

    deck, *slides = [
        d
        for d in yaml.safe_load_all(draft_spec(harvest(built.deck), title="D", theme=theme).text)
        if isinstance(d, dict)
    ]
    assert deck["sections"] == ["A", "B"]
    assert [s.get("section") for s in slides[:3]] == ["A", None, "B"]

    built.manifest.rename(tmp_path / "elsewhere.json")
    alone = draft_spec(harvest(built.deck), title="D", theme=theme).text
    deck, *slides = [d for d in yaml.safe_load_all(alone) if isinstance(d, dict)]
    assert "sections" not in deck and not any("section" in s for s in slides)


_STILL_AND_LAST = """theme: base
title: Builds
out: out/Builds.pptx
---
title: Slots, then keys
animate: one_at_a_time
place:
  - at: {cols: full, rows: {from: 0, to: 3}}
    animate: none
    table: {header: [a, b], rows: [['1', '2']]}
  - at: {cols: left-half, rows: {from: 4, to: 8}}
    card: {heading: cat}
  - at: {cols: right-half, rows: {from: 4, to: 8}}
    card: {heading: dog}
---
title: The claim, last
chrome:
  title: {at: {cols: full, rows: {from: 9, to: 11}}, build: last}
place:
  - at: {cols: full, rows: {from: 0, to: 8}}
    card: {heading: Evidence}
---
title: Bars, then the claim
animate: by_category
place:
  - at: {cols: full, rows: {from: 0, to: 9}}
    chart: {kind: column, data: [{category: A, value: 1}, {category: B, value: 2}]}
  - at: {cols: full, rows: {from: 9, to: 12}}
    card: {heading: B wins}
"""


def test_a_still_placement_a_title_built_last_and_a_chart_then_its_claim_round_trip(tmp_path):
    spec = tmp_path / "b.deck.yaml"
    spec.write_text(_STILL_AND_LAST)
    built = build_deck(spec)
    drafted = draft_spec(harvest(built.deck), title="Draft", theme="base")
    first, second, third = [d for d in yaml.safe_load_all(drafted.text) if isinstance(d, dict)][1:]
    assert first["place"][0]["animate"] == "none" and "animate" not in first["place"][1]
    assert second["chrome"]["title"]["build"] == "last" and "animate" not in second
    assert third["animate"] == "by_category"

    again = tmp_path / "again"
    again.mkdir()
    (again / "d.deck.yaml").write_text(
        drafted.text.replace("out: out/draft/Draft v1.pptx", "out: out/Draft.pptx")
    )
    rebuilt = build_deck(again / "d.deck.yaml").deck
    assert read_back(rebuilt, manifest=built.manifest).changes == ()
