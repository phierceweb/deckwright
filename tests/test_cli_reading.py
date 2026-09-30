"""`diff` and `inspect` as a terminal shows them."""

from __future__ import annotations

import pytest
from typer.testing import CliRunner

from deckwright.cli import app

runner = CliRunner()


@pytest.fixture
def built(tmp_path, theme_file):
    from deckwright.compile import build_deck

    spec = tmp_path / "d.deck.yaml"
    spec.write_text(
        "theme: testtheme\nout: out/D.pptx\n---\ntitle: A title\n"
        "place:\n  - at: {cols: full}\n    bullets: {items: [alpha, beta]}\n"
    )
    return build_deck(spec, theme_path=theme_file)


def _retitled(built):
    from pptx import Presentation

    prs = Presentation(str(built.deck))
    title = next(s for s in prs.slides[0].shapes if s.name == "s1.chrome.title")
    title.text_frame.paragraphs[0].runs[0].text = "Another title"
    prs.save(str(built.deck))


def test_diff_on_the_deck_that_was_built_says_there_is_nothing_to_carry_back(built):
    result = runner.invoke(app, ["diff", str(built.deck)])
    assert result.exit_code == 0, result.output
    assert result.stdout == "D.pptx is the deck that was built — nothing to carry back.\n"


def test_diff_lists_each_change_and_writes_the_report(built, tmp_path):
    _retitled(built)
    result = runner.invoke(app, ["diff", str(built.deck), "--out", str(tmp_path / "r")])
    assert result.exit_code == 0, result.output
    lines = result.stdout.splitlines()
    assert lines[0].startswith("D.pptx was edited after its build from ")
    assert lines[1] == "  slide 1  retyped  s1.chrome.title  'A title' → 'Another title'"
    assert lines[2] == f"report -> {tmp_path / 'r' / 'readback.md'}"
    assert "**retyped** `s1.chrome.title`" in (tmp_path / "r" / "readback.md").read_text()


def test_diff_says_so_when_the_file_changed_and_no_shape_did(built):
    from pptx import Presentation

    prs = Presentation(str(built.deck))
    prs.core_properties.comments = "resaved"
    prs.save(str(built.deck))
    result = runner.invoke(app, ["diff", str(built.deck)])
    assert result.stdout.splitlines()[1] == (
        "  no shape differs — a resave, or a change this cannot see"
    )


def test_inspect_lists_each_slides_shapes_with_their_boxes(built):
    result = runner.invoke(app, ["inspect", str(built.deck)])
    assert result.exit_code == 0, result.output
    lines = result.stdout.splitlines()
    assert lines[0] == "D.pptx: 1 slide(s)"
    assert lines[1].startswith("slide 1 (")
    assert any(line.endswith("'s1.chrome.title'") and line.startswith("  id=") for line in lines)
    assert all("×" in line or "no box" in line for line in lines[2:])
