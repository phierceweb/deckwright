"""End-to-end shots through a real headless browser. Skipped when none is installed."""

from __future__ import annotations

import pytest
from PIL import Image

from deckwright.errors import RenderError
from deckwright.services.htmlcard import markdown_card
from deckwright.services import htmlshot
from deckwright.services.chrome import resolve_chrome
from deckwright.services.htmlshot import render_html_to_png


def _chrome_or_skip() -> str:
    try:
        return resolve_chrome(None)
    except RenderError:
        pytest.skip("no Chrome/Chromium binary available")


CARD = markdown_card(
    "# Doc\n\n" + "\n\n".join(f"Paragraph {i}." for i in range(6)), filename="doc.md"
)

# python-markdown passes a raw HTML block straight through, so <plaintext> reaches the
# browser and swallows the rest of the parse — the appended height probe included.
UNMEASURABLE_CARD = markdown_card(
    "# Release notes\n\n<plaintext>\n"
    + "\n".join(f"line {i} of the changelog" for i in range(150)),
    filename="CHANGELOG.md",
)


def test_a_document_taller_than_the_canvas_raises_instead_of_truncating(tmp_path, monkeypatch):
    _chrome_or_skip()
    monkeypatch.setenv("DECKWRIGHT_SHOT_CANVAS_H", "200")
    with pytest.raises(RenderError, match=r"the browser clipped it"):
        render_html_to_png(CARD, tmp_path / "clipped.png", width=600, scale=1)


def test_the_same_document_renders_once_the_canvas_is_raised(tmp_path, monkeypatch):
    _chrome_or_skip()
    monkeypatch.setenv("DECKWRIGHT_SHOT_CANVAS_H", "2000")
    out = render_html_to_png(CARD, tmp_path / "ok.png", width=600, scale=1)
    assert Image.open(out).height < 2000  # cropped to content, not the canvas


def test_a_document_that_swallows_the_height_probe_is_caught_by_its_pixels(tmp_path, monkeypatch):
    """With no height to read, a card running off the last row is the only evidence."""
    _chrome_or_skip()
    monkeypatch.setenv("DECKWRIGHT_SHOT_CANVAS_H", "400")
    with pytest.raises(RenderError, match=r"the height probe did not run"):
        render_html_to_png(UNMEASURABLE_CARD, tmp_path / "swallowed.png", width=600, scale=1)


FRAMED_FILE_CARD = markdown_card(
    '# Notes\n\nBody.\n\n<iframe src="file:///etc/hosts" width="600" height="130"></iframe>\n',
    filename="notes.md",
)


def _ink(path) -> int:
    """Pixels far enough from white to be content."""
    grey = Image.open(path).convert("L")
    return sum(1 for px in grey.get_flattened_data() if px < 200)


def test_a_file_url_frame_in_someone_elses_markdown_never_reaches_the_render(tmp_path, monkeypatch):
    _chrome_or_skip()
    monkeypatch.setattr(htmlshot, "CSP_META", "")
    leaked = _ink(render_html_to_png(FRAMED_FILE_CARD, tmp_path / "off.png", width=700, scale=1))
    monkeypatch.undo()
    if leaked < 50_000:
        pytest.skip("this browser does not render file:// frames — nothing to contain")

    contained = _ink(render_html_to_png(FRAMED_FILE_CARD, tmp_path / "on.png", width=700, scale=1))
    assert contained < leaked * 0.3


def test_the_height_probe_still_runs_under_the_content_policy(tmp_path, monkeypatch):
    """Allowed by hash; if that breaks, cards clip silently — pin the loud variant."""
    _chrome_or_skip()
    monkeypatch.setenv("DECKWRIGHT_SHOT_CANVAS_H", "200")
    with pytest.raises(RenderError, match=r"content is \d+px tall"):
        render_html_to_png(CARD, tmp_path / "clipped.png", width=600, scale=1)


def test_a_card_is_placed_on_a_slide_without_its_file_name_as_alt_text(tmp_path):
    """python-pptx describes a picture by its file name; a temp file's name is not alt text."""
    from pptx import Presentation

    from deckwright.services.htmlshot import card_to_slide
    from deckwright.utils.a11y import described

    chrome = _chrome_or_skip()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    picture = card_to_slide(slide, CARD, left=1.0, top=1.0, width=4.0, chrome=chrome)
    assert [s.shape_type for s in slide.shapes] == [picture.shape_type]
    assert round(picture.width / 914400, 2) == 4.0
    assert described(picture) == (None, False)


def test_a_transparent_render_is_clear_outside_the_card_and_solid_inside(tmp_path):
    png = render_html_to_png(
        CARD, tmp_path / "card.png", chrome=_chrome_or_skip(), width=600, pad=20, transparent=True
    )
    with Image.open(png) as img:
        assert img.mode == "RGBA"
        assert img.getpixel((0, 0))[3] == 0
        assert img.getpixel((img.width // 2, img.height // 2))[3] == 255


def test_a_plain_render_is_as_opaque_as_it_was(tmp_path):
    png = render_html_to_png(
        CARD, tmp_path / "card.png", chrome=_chrome_or_skip(), width=600, pad=20
    )
    with Image.open(png) as img:
        assert img.convert("RGBA").getpixel((0, 0))[3] == 255
