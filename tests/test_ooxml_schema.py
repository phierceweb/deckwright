"""The raw OOXML deckwright writes, against ISO/IEC 29500-4:2016: one writer at a time on
small decks, then every part of a built deck (`docs/testing.md`)."""

from __future__ import annotations

import contextlib
import copy
import io
import re
import zipfile

import pytest
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from deckwright.motion import (
    add_chart_build,
    add_click_build,
    add_click_reveals,
    add_click_sequence,
    add_transition,
)
from deckwright.motion.transition import EFFECTS, SPEEDS
from tests import ooxml

SCHEMA = ooxml.SCHEMAS / "pml.xsd"
_PML, _DML, _MC, _OMML = ooxml.PML, ooxml.DML, ooxml.MC, ooxml.OMML
_A14 = "http://schemas.microsoft.com/office/drawing/2010/main"
_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"


@pytest.fixture(scope="module")
def schema():
    return ooxml.schema()


def _deck():
    """A slide carrying the three shape kinds the builds are pointed at."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    box.text_frame.text = "text"
    blank = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1))
    data = CategoryChartData()
    data.categories = ["A", "B", "C"]
    data.add_series("S1", (1, 2, 3))
    data.add_series("S2", (3, 2, 1))
    frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4), Inches(1), Inches(4), Inches(3), data
    )
    return prs, slide, box.shape_id, blank.shape_id, frame.shape_id


def _slide_xml(prs) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return zipfile.ZipFile(buf).read("ppt/slides/slide1.xml")


def _validate(schema, prs) -> list[str]:
    doc = etree.fromstring(_slide_xml(prs))
    if schema.validate(doc):
        return []
    return [e.message for e in schema.error_log]


def _equation_slide(prefix_on_root: bool = False):
    """A text box stored as PowerPoint stores an equation: maths in the Choice, plain text in
    the Fallback; ``prefix_on_root`` declares `a14` on the root and spells it `x` on the maths."""
    prs, slide, text, blank, chart = _deck()
    root = etree.fromstring(_slide_xml(prs))
    if prefix_on_root:
        declared = etree.Element(root.tag, attrib=root.attrib, nsmap={**root.nsmap, "a14": _A14})
        declared.extend(list(root))
        root = declared
    sp = next(
        el
        for el in root.iter(f"{{{_PML}}}sp")
        if el.find(f".//{{{_PML}}}cNvPr").get("id") == str(text)
    )
    alternate = etree.Element(f"{{{_MC}}}AlternateContent", nsmap={"mc": _MC})
    sp.addprevious(alternate)
    choice = etree.SubElement(
        alternate,
        f"{{{_MC}}}Choice",
        nsmap=None if prefix_on_root else {"a14": _A14},
        Requires="a14",
    )
    maths = copy.deepcopy(sp)
    paragraph = maths.find(f".//{{{_DML}}}p")
    for child in list(paragraph):
        paragraph.remove(child)
    paragraph.append(
        etree.fromstring(
            f'<a14:m xmlns:a14="{_A14}" xmlns:m="{_OMML}"><m:oMathPara><m:oMath>'
            "<m:r><m:t>x²</m:t></m:r></m:oMath></m:oMathPara></a14:m>"
        )
    )
    choice.append(maths)
    etree.SubElement(alternate, f"{{{_MC}}}Fallback").append(sp)
    if not prefix_on_root:
        return root
    data = etree.tostring(root).replace(b"a14:m", b"x:m")
    data = data.replace(b"<p:sld ", f'<p:sld xmlns:x="{_A14}" '.encode(), 1)
    assert b"<x:m" in data and b"<a14:m" not in data
    return etree.fromstring(data)


@pytest.mark.parametrize("which", ["Choice", "Fallback"])
def test_each_branch_of_markup_compatibility_validates_on_its_own(schema, which):
    """`pml.xsd` has no `mc:AlternateContent`, so the wrapper itself never validates; a reader
    only ever sees one branch, and each has to be valid on its own."""
    root = _equation_slide()
    assert not schema.validate(root)

    assert schema.validate(ooxml.branch(root, which)), [e.message for e in schema.error_log]


@pytest.mark.parametrize("broken, sound", [("Choice", "Fallback"), ("Fallback", "Choice")])
def test_the_branch_gate_catches_a_fault_inside_one_branch(schema, broken, sound):
    """The negative control for the case above. A harness that dropped the wrapper, or read
    one branch for both, would pass a fault confined to a single branch."""
    root = _equation_slide()
    sp = root.find(f".//{{{_MC}}}{broken}/{{{_PML}}}sp")
    sp.remove(sp.find(f"{{{_PML}}}spPr"))

    assert schema.validate(ooxml.branch(root, sound)), [e.message for e in schema.error_log]
    assert not schema.validate(ooxml.branch(root, broken))
    assert "spPr" in schema.error_log[0].message


def _morph_slide():
    """A slide that arrives by morph and also carries a build, so the wrapper has a
    `p:timing` to sit in front of."""
    prs, slide, text, blank, chart = _deck()
    add_click_build(slide, [text])
    add_transition(slide, "morph", speed="slow")
    return etree.fromstring(_slide_xml(prs))


@pytest.mark.parametrize("which", ["Choice", "Fallback"])
def test_each_branch_of_a_morph_transition_validates_on_its_own(schema, which):
    assert schema.validate(ooxml.branch(_morph_slide(), which)), [
        e.message for e in schema.error_log
    ]


def test_the_gate_catches_a_morph_wrapper_after_the_timing_tree(schema):
    """The negative control: `CT_Slide` orders the transition before the timing, and a branch
    keeps the place its wrapper had."""
    root = _morph_slide()
    wrapper = root.find(f"{{{_MC}}}AlternateContent")
    root.remove(wrapper)
    root.append(wrapper)

    assert not schema.validate(ooxml.branch(root, "Fallback"))
    assert "transition" in schema.error_log[0].message


def test_the_vendored_schema_is_present():
    """A skipped schema gate proves nothing; this fails rather than skips."""
    assert SCHEMA.is_file(), f"missing vendored schema at {SCHEMA}"


def test_a_click_build_validates(schema):
    prs, slide, text, blank, chart = _deck()
    add_click_build(slide, [text, blank])
    assert _validate(schema, prs) == []


def test_a_staggered_click_build_validates(schema):
    prs, slide, text, blank, chart = _deck()
    add_click_build(slide, [text, blank], 80)
    assert _validate(schema, prs) == []


@pytest.mark.parametrize("kind", ["fade", "wipeup", "wiperight"])
def test_every_click_sequence_effect_kind_validates(schema, kind):
    prs, slide, text, blank, chart = _deck()
    add_click_sequence(slide, [[(text, kind)], [(blank, kind)]], 60)
    assert _validate(schema, prs) == []


@pytest.mark.parametrize("kind", ["fade", "wipeup", "wiperight"])
@pytest.mark.parametrize("by", ["category", "series", "element", "all"])
def test_every_chart_build_validates(schema, by, kind):
    """`bldStep` is required on `<a:chart>`. Every
    entrance a theme's `datum` role may bind is built the same way."""
    prs, slide, text, blank, chart = _deck()
    add_chart_build(slide, chart, by, parts=3, kind=kind)
    assert _validate(schema, prs) == []


@pytest.mark.parametrize("kind", ["fade", "wipeup", "wiperight"])
def test_every_after_previous_chain_validates(schema, kind):
    """`afterEffect` group starts and their beat gate, the auto-advance path."""
    prs, slide, text, blank, chart = _deck()
    add_click_sequence(slide, [[(text, kind)], [(blank, kind)]], beat_ms=250)
    assert _validate(schema, prs) == []


def test_a_build_over_only_text_free_shapes_validates(schema):
    """No bldP is legal for these, so the build list is omitted — and `<p:bldLst/>`
    with no children would be invalid."""
    prs, slide, text, blank, chart = _deck()
    add_click_build(slide, [blank])
    assert _validate(schema, prs) == []


def test_a_group_holding_both_shape_kinds_validates(schema):
    """A worded shape and a text-free one on one click: the first entrance node carries `grpId` and
    the second does not. The XSD accepts either — it is `use="optional"` on
    `CT_TLCommonTimeNodeData` — so this gates only spelling and position;
    `tests/test_pptx_helpers.py` gates which node."""
    prs, slide, text, blank, chart = _deck()
    add_click_sequence(slide, [[text, blank]])
    assert _validate(schema, prs) == []


def test_an_interactive_reveal_validates(schema):
    prs, slide, text, blank, chart = _deck()
    add_click_reveals(slide, [(text, blank)])
    assert _validate(schema, prs) == []


def test_two_triggers_onto_one_target_stay_siblings_under_the_root(schema):
    """Every shape a trigger placement drew listens, so a two-shape trigger writes two `<p:seq>`.
    Both hang off the tmRoot: burying the second inside the first validates, but it would only arm
    once the outer sequence had run."""
    prs, slide, text, blank, chart = _deck()
    add_click_reveals(slide, [(text, blank), (chart, blank)])

    assert _validate(schema, prs) == []
    seqs = list(etree.fromstring(_slide_xml(prs)).iter(f"{{{_PML}}}seq"))
    assert len(seqs) == 2
    assert {s.getparent().getparent().get("nodeType") for s in seqs} == {"tmRoot"}


# PowerPoint for Mac 16's own trigger, saved onto the confirmation kit's deck 04: ids dropped,
# trigger spid `T`, target spid `R`, and without the `grpId`/`bldP` it adds for a worded target.
_POWERPOINT_TRIGGER = (
    f'<p:seq xmlns:p="{_PML}" concurrent="1" nextAc="seek">'
    '<p:cTn restart="whenNotActive" fill="hold" evtFilter="cancelBubble" nodeType="interactiveSeq">'
    '<p:stCondLst><p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="T"/></p:tgtEl></p:cond></p:stCondLst>'
    '<p:endSync evt="end" delay="0"><p:rtn val="all"/></p:endSync>'
    '<p:childTnLst><p:par><p:cTn fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
    '<p:par><p:cTn fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
    '<p:par><p:cTn presetID="1" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect">'
    '<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr>'
    '<p:cTn dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
    '<p:tgtEl><p:spTgt spid="R"/></p:tgtEl>'
    "<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>"
    '</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par>'
    "</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn>"
    '<p:nextCondLst><p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="T"/></p:tgtEl></p:cond></p:nextCondLst>'
    "</p:seq>"
)


def _skeleton(node, spids: dict[str, str]) -> list[tuple[str, tuple[tuple[str, str], ...], str]]:
    """Every element under ``node`` in document order as (tag, attributes, text), with ``id``
    dropped and each ``spid`` renamed through ``spids``."""
    out = []
    for el in node.iter():
        attrs = tuple(
            sorted(
                (k, spids.get(v, v) if k == "spid" else v)
                for k, v in el.attrib.items()
                if k != "id"
            )
        )
        out.append((etree.QName(el).localname, attrs, (el.text or "").strip()))
    return out


def test_a_trigger_is_shaped_like_powerpoints_own(schema):
    """`cancelBubble`, an `endSync`, a next-condition on the trigger's own click and no
    previous-condition: a click anywhere else on the slide advances it instead of stepping
    every trigger forward. The shape is PowerPoint's, learned back from its save."""
    prs, slide, text, blank, chart = _deck()
    add_click_reveals(slide, [(text, blank)])

    assert _validate(schema, prs) == []
    seq = etree.fromstring(_slide_xml(prs)).find(f".//{{{_PML}}}seq")
    expected = etree.fromstring(_POWERPOINT_TRIGGER.encode())
    assert _skeleton(seq, {str(text): "T", str(blank): "R"}) == _skeleton(expected, {})


def test_the_gate_catches_an_end_sync_out_of_order(schema):
    """The negative control for the case above: `endSync` after `childTnLst` breaks
    `CT_TLCommonTimeNodeData`'s sequence, which LibreOffice repairs without a word."""
    prs, slide, text, blank, chart = _deck()
    add_click_reveals(slide, [(text, blank)])
    root = etree.fromstring(_slide_xml(prs))
    ctn = root.find(f".//{{{_PML}}}seq/{{{_PML}}}cTn")
    sync = ctn.find(f"{{{_PML}}}endSync")
    ctn.remove(sync)
    ctn.append(sync)

    assert not schema.validate(root)
    assert "endSync" in schema.error_log[0].message


@pytest.mark.parametrize("kind", sorted(EFFECTS))
def test_every_transition_validates(schema, kind):
    prs, slide, text, blank, chart = _deck()
    add_transition(slide, kind)
    assert _validate(schema, prs) == []


@pytest.mark.parametrize(
    "kind, direction",
    [(k, d) for k, dirs in sorted(EFFECTS.items()) for d in dirs],
)
def test_every_transition_direction_validates(schema, kind, direction):
    """Direction vocabularies are per element — a shared list is invalid at `strips`."""
    prs, slide, text, blank, chart = _deck()
    add_transition(slide, kind, direction=direction)
    assert _validate(schema, prs) == []


@pytest.mark.parametrize("speed", SPEEDS)
def test_every_transition_speed_validates(schema, speed):
    prs, slide, text, blank, chart = _deck()
    add_transition(slide, "wipe", speed=speed)
    assert _validate(schema, prs) == []


def test_a_transition_beside_an_animation_validates_and_keeps_child_order(schema):
    """`CT_Slide` is an xsd:sequence. LibreOffice repairs a wrong order on import, so
    this is the only mechanical check that the two writers do not collide."""
    prs, slide, text, blank, chart = _deck()
    add_click_sequence(slide, [[text], [blank]])
    add_transition(slide, "push", direction="u")

    assert _validate(schema, prs) == []
    tags = [el.tag.split("}")[1] for el in slide._element]
    assert tags.index("transition") < tags.index("timing")


def test_the_gate_catches_a_missing_required_attribute(schema):
    """The negative control. Without it, a validator that silently passed everything
    would look identical to one that works."""
    prs, slide, text, blank, chart = _deck()
    add_chart_build(slide, chart, "category", parts=3)
    sabotaged = re.sub(rb' bldStep="[^"]+"', b"", _slide_xml(prs))

    assert not schema.validate(etree.fromstring(sabotaged))
    assert "bldStep" in schema.error_log[0].message


def test_a_click_sequence_revealing_paragraphs_of_one_shape_validates(schema):
    """`one_at_a_time` on a single bullets column reveals one paragraph per click."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(2))
    box.text_frame.text = "one"
    for line in ("two", "three"):
        box.text_frame.add_paragraph().text = line
    add_click_sequence(slide, [[(box.shape_id, "fade", i)] for i in range(3)])
    assert _validate(schema, prs) == []


def test_alt_text_and_the_decorative_flag_validate_on_every_figure_kind(schema):
    """`descr` on each non-visual block, and the decorative extension inside `cNvPr`'s
    `a:extLst`, which the schema types as a closed sequence ending in that list."""
    from deckwright.utils.a11y import describe

    prs, slide, text, blank, chart = _deck()
    shapes = {s.shape_id: s for s in slide.shapes}
    describe(shapes[chart], alt="A chart")
    describe(shapes[text], decorative=True)
    describe(shapes[blank], alt="A box")
    assert _validate(schema, prs) == []


def test_the_gate_catches_the_decorative_list_out_of_order(schema):
    """The negative control for the case above: a click action after `a:extLst` breaks the
    sequence, so an extension written in the wrong place would not pass."""
    from deckwright.utils.a11y import describe

    prs, slide, text, blank, chart = _deck()
    describe(next(s for s in slide.shapes if s.shape_id == text), decorative=True)
    sabotaged = _slide_xml(prs).replace(b"</a:extLst>", b'</a:extLst><a:hlinkClick r:id=""/>')

    assert not schema.validate(etree.fromstring(sabotaged))
    assert "hlinkClick" in schema.error_log[0].message


def test_a_click_action_beside_the_decorative_flag_validates_in_order(schema):
    """`goto:` writes `a:hlinkClick` into a `cNvPr` whose `a:extLst` alt text may already
    hold; the sequence puts the click first whichever was written first."""
    from deckwright.compile.goto import write_gotos
    from deckwright.utils.a11y import describe

    prs, slide, text, blank, chart = _deck()
    second = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = {s.shape_id: s for s in slide.shapes}
    describe(shapes[text], decorative=True)
    write_gotos(
        [(shapes[text], "next"), (shapes[blank], "back"), (shapes[chart], "back")],
        slides={"back": second},
    )
    assert _validate(schema, prs) == []
    tags = [el.tag.split("}")[1] for el in shapes[text]._element.nvSpPr.cNvPr]
    assert tags == ["hlinkClick", "extLst"]


def test_a_link_run_with_its_colour_extension_validates(schema):
    """`a:hlinkClick` sits after the run's fill and face in `a:rPr`, and carries the
    hyperlink-colour extension in its own `a:extLst`."""
    from deckwright.utils.shapes import para
    from pptx.dml.color import RGBColor

    prs, slide, text, blank, chart = _deck()
    frame = next(s for s in slide.shapes if s.shape_id == blank).text_frame
    para(frame, "Read [the guide](https://example.com/g) first", 18, RGBColor(0, 0, 0), first=True)
    assert _validate(schema, prs) == []


def test_a_cjk_run_marked_with_its_language_and_face_validates(schema):
    from deckwright.compile.eastasian import mark_east_asian

    prs, slide, text, blank, chart = _deck()
    frame = next(s for s in slide.shapes if s.shape_id == blank).text_frame
    run = frame.paragraphs[0].add_run()
    run.text = "売上は伸びました"
    run.font.name = "Helvetica"
    run.hyperlink.address = "https://example.com"
    mark_east_asian(slide, ea={"ja": "Hiragino Sans"}, lang=None, where="s")
    assert _validate(schema, prs) == []


@pytest.fixture(scope="module")
def math_schema():
    return ooxml.schema("shared-math.xsd")


_EVERY_NODE = (
    r"\left( \sum_{i=1}^{n} x_i^2 \right) + \sqrt[3]{\frac{a}{b}} - \lim_{x \to 0} y_j"
    r" + \int_0^1 z^2 + \prod w + \sqrt{q} + \text{if } p \leq \infty"
)


def _math(tex: str):
    """The `m:oMathPara` an equation writes, less the DrawingML run properties PowerPoint
    keeps where the ISO schema has Word's: those are DrawingML's to validate, not Office
    Math's."""
    from deckwright.components._omml import omml
    from deckwright.components._tex import parse

    root = etree.fromstring(omml(parse(tex, where="t"), size_pt=24, ink="1A1D21").encode())
    for props in list(root.iter(f"{{{_DML}}}rPr")):
        props.getparent().remove(props)
    return root.find(f"{{{_OMML}}}oMathPara")


def test_every_node_an_equation_writes_is_valid_office_math(math_schema):
    assert math_schema.validate(_math(_EVERY_NODE)), [e.message for e in math_schema.error_log]


def test_the_math_gate_catches_a_fraction_written_upside_down(math_schema):
    """The negative control: `CT_F` is `num` then `den`."""
    para = _math(_EVERY_NODE)
    fraction = para.find(f".//{{{_OMML}}}f")
    num, den = fraction.find(f"{{{_OMML}}}num"), fraction.find(f"{{{_OMML}}}den")
    fraction.remove(num)
    fraction.append(num)

    assert den is not None
    assert not math_schema.validate(para)
    assert "num" in math_schema.error_log[0].message


@pytest.mark.parametrize("which", ["Choice", "Fallback"])
def test_a_slide_carrying_an_equation_validates_in_each_branch(
    schema, math_schema, which, ctx_factory
):
    """The component wraps the whole shape, so each reader sees a complete, valid slide."""
    from deckwright.layouts.components import get_component

    ctx = ctx_factory({"equation": {"tex": r"x = \frac{-b \pm \sqrt{b^2 - 4ac}}{2a}"}})
    get_component("equation")(ctx)
    root = etree.fromstring(etree.tostring(ctx.slide._element))

    assert schema.validate(ooxml.branch(root, which)), [e.message for e in schema.error_log]
    if which == "Choice":
        (para,) = root.iter(f"{{{_OMML}}}oMathPara")
        for props in list(para.iter(f"{{{_DML}}}rPr")):
            props.getparent().remove(props)
        assert math_schema.validate(para), [e.message for e in math_schema.error_log]


@pytest.mark.parametrize("kind", ["fade", "wipeup", "wiperight"])
@pytest.mark.parametrize("by", ["category", "series", "all"])
def test_a_chart_build_then_a_click_sequence_validates(schema, by, kind):
    """The bars, then the claim: one timeline, the chart's `bldGraphic` and the claim's
    `bldP` in one build list."""
    prs, slide, text, blank, chart = _deck()
    add_chart_build(slide, chart, by, parts=3, kind=kind, then=[[(text, kind)], [(blank, "fade")]])
    assert _validate(schema, prs) == []


def test_a_chart_build_then_a_click_sequence_numbers_every_node_once(schema):
    """The schema does not hold `cTn` ids unique; the chart's numbering and the sequence's
    must not meet, or PowerPoint has two nodes under one id."""
    prs, slide, text, blank, chart = _deck()
    add_chart_build(slide, chart, "category", parts=3, then=[[(text, "fade")], [(blank, "fade")]])
    ids = [n.get("id") for n in etree.fromstring(_slide_xml(prs)).iter(f"{{{_PML}}}cTn")]
    assert len(ids) == len(set(ids)) == 4 + 3 * 4 + 2 * (2 + 3)


# --- the whole deck: every part of a real build ------------------------------------------

CATALOGUE = pytest.mark.xdist_group("catalogue")


@pytest.fixture(scope="module")
def catalogue(tmp_path_factory):
    """The capability catalogue against `base`, on a cache apart from the user's so its
    `document:` exercise starts the browser."""
    from deckwright.conform.demo import demo

    root = tmp_path_factory.mktemp("catalogue")
    with pytest.MonkeyPatch.context() as patch:
        patch.setenv("DECKWRIGHT_CACHE_DIR", str(root / "cache"))
        return demo("base", root / "deck")


def _members(deck) -> list[str]:
    with zipfile.ZipFile(deck) as archive:
        return archive.namelist()


def _read(deck, part: str) -> bytes:
    with zipfile.ZipFile(deck) as archive:
        return archive.read(part)


def _scope(root, inside: str | None):
    return {
        None: root,
        "Choice": root.find(f".//{{{_MC}}}Choice"),
        "Fallback": root.find(f".//{{{_MC}}}Fallback"),
    }[inside]


def _bogus(data: bytes, inside: str | None = None) -> bytes:
    """``data`` with an element no schema allows as the first child of its root, or of the
    first element in its ``mc:Choice`` or ``mc:Fallback``."""
    root = etree.fromstring(data)
    where = root if inside is None else _scope(root, inside)[0]
    where.insert(0, etree.Element("{urn:deckwright:bogus}bogus"))
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# An attribute each kind of part carries, and a value its type refuses; the presentation
# properties carry none, so they get a `p:showPr` boolean attribute given a non-boolean value.
_BAD_VALUES = (
    (f"{{{_DML}}}off", "x", "warp"),
    (f"{{{_DML}}}srgbClr", "val", "nothex"),
    (f"{{{_PML}}}sldSz", "cx", "warp"),
    (f"{{{_PML}}}transition", "spd", "warp"),
    (f"{{{_DML}}}sx", "n", "warp"),
    (f"{{{_DML}}}tblStyleLst", "def", "warp"),
)


def _bad_value(data: bytes, inside: str | None = None) -> bytes:
    """``data`` with one attribute set to a value its type refuses, inside the branch named."""
    root = etree.fromstring(data)
    scope = _scope(root, inside)
    if scope.tag == f"{{{_PML}}}presentationPr":
        scope.insert(0, etree.Element(f"{{{_PML}}}showPr", useTimings="warp"))
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    for tag, attribute, value in _BAD_VALUES:
        target = scope.find(f".//{tag}") if scope.tag != tag else scope
        if target is not None:
            target.set(attribute, value)
            return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    raise LookupError("no attribute to break")


_DAMAGE = {"an element no schema allows": _bogus, "a value no type allows": _bad_value}

# Office's own extension namespaces, which no vendored schema defines, so a lax wildcard
# or a stripped Choice passes them unjudged. One more is a change to look at.
_PASSED_OVER = frozenset(
    {
        "http://schemas.microsoft.com/office/drawing/2007/8/2/chart",
        "http://schemas.microsoft.com/office/drawing/2010/main",
        "http://schemas.microsoft.com/office/drawing/2012/chart",
        "http://schemas.microsoft.com/office/drawing/2017/decorative",
        "http://schemas.microsoft.com/office/drawing/2018/hyperlinkcolor",
        "http://schemas.microsoft.com/office/powerpoint/2010/main",
        "http://schemas.microsoft.com/office/powerpoint/2015/09/main",
    }
)


@CATALOGUE
def test_every_part_of_the_built_catalogue_validates(catalogue):
    """Every feature on one real deck: every part valid, every `.xml` and `.rels` member
    accounted for, and what the verdict passed over pinned."""
    from deckwright.conform.exercise import EXERCISE

    report = ooxml.validate_deck(catalogue)
    assert report.invalid == {}
    assert sum(p.startswith("ppt/slides/") for p in report.valid) == len(EXERCISE)
    assert any(p.startswith("ppt/charts/") for p in report.valid)
    assert any(p.startswith("ppt/notesSlides/") for p in report.valid)
    xml_members = {m for m in _members(catalogue) if m.lower().endswith((".xml", ".rels"))}
    assert xml_members <= set(report.valid) | set(report.unchecked)
    assert set(report.unchecked.values()) == ooxml.PACKAGE_ROOTS
    assert {ns for seen in report.passed_over.values() for ns in seen} == _PASSED_OVER
    assert report.unreached == {}


_KINDS = (
    "first slide",
    "last slide",
    "the equation slide, in its Choice",
    "the morph slide, in its Fallback",
    "a chart",
    "a chart carrying AlternateContent",
    "the layout",
    "the master",
    "the theme",
    "the notes slide",
    "the notes master",
    "the notes master's theme",
    "the presentation",
    "the presentation properties",
    "the view properties",
    "the table styles",
)


def _part(deck, kind: str) -> tuple[str, str | None]:
    """One part of that kind in the catalogue, and the branch to break inside, if any."""
    members = _members(deck)

    def numbered(prefix):
        hits = [m for m in members if re.fullmatch(prefix + r"[A-Za-z]+\d+\.xml", m)]
        return sorted(hits, key=lambda m: int(re.search(r"(\d+)\.xml$", m).group(1)))

    def holding(names, token, present=True):
        return next(n for n in names if (token in _read(deck, n)) is present)

    slides, charts = numbered("ppt/slides/"), numbered("ppt/charts/")
    return {
        "first slide": (slides[0], None),
        "last slide": (slides[-1], None),
        "the equation slide, in its Choice": (holding(slides, b"oMathPara"), "Choice"),
        "the morph slide, in its Fallback": (holding(slides, b"morph"), "Fallback"),
        "a chart": (holding(reversed(charts), b"AlternateContent", present=False), None),
        "a chart carrying AlternateContent": (holding(charts, b"AlternateContent"), None),
        "the layout": (numbered("ppt/slideLayouts/")[0], None),
        "the master": (numbered("ppt/slideMasters/")[0], None),
        "the theme": (numbered("ppt/theme/")[0], None),
        "the notes slide": (numbered("ppt/notesSlides/")[0], None),
        "the notes master": (numbered("ppt/notesMasters/")[0], None),
        "the notes master's theme": (numbered("ppt/theme/")[-1], None),
        "the presentation": ("ppt/presentation.xml", None),
        "the presentation properties": ("ppt/presProps.xml", None),
        "the view properties": ("ppt/viewProps.xml", None),
        "the table styles": ("ppt/tableStyles.xml", None),
    }[kind]


@CATALOGUE
@pytest.mark.parametrize("damage", sorted(_DAMAGE))
@pytest.mark.parametrize("kind", _KINDS)
def test_the_whole_deck_gate_names_the_one_part_broken(catalogue, tmp_path, kind, damage):
    """The negative control, once per kind of part the real build holds and per kind of
    fault: the part is named, that part alone, under the branch that carries the fault."""
    part, inside = _part(catalogue, kind)
    data = _DAMAGE[damage](_read(catalogue, part), inside)
    broken = ooxml.replace_part(catalogue, part, data, tmp_path / "broken.pptx")

    report = ooxml.validate_deck(broken)
    assert list(report.invalid) == [part]
    if damage == "an element no schema allows":
        assert "bogus" in report.invalid[part][0]
    if inside:
        assert all(e.startswith(f"[{inside}]") for e in report.invalid[part])


@CATALOGUE
@pytest.mark.parametrize("kind", _KINDS)
def test_the_whole_deck_gate_holds_each_part_to_its_content_type(catalogue, tmp_path, kind):
    """A valid part of another kind stored where `[Content_Types].xml` promises this one."""
    part, _ = _part(catalogue, kind)
    other, _ = _part(catalogue, "first slide" if "theme" in kind else "the theme")
    swapped = ooxml.replace_part(catalogue, part, _read(catalogue, other), tmp_path / "s.pptx")

    report = ooxml.validate_deck(swapped)
    assert list(report.invalid) == [part]
    assert "expects" in report.invalid[part][0]


@CATALOGUE
def test_the_whole_deck_gate_does_not_stop_at_the_first_fault(catalogue, tmp_path):
    slide, _ = _part(catalogue, "first slide")
    chart, _ = _part(catalogue, "a chart")
    once = ooxml.replace_part(
        catalogue, slide, _bogus(_read(catalogue, slide)), tmp_path / "1.pptx"
    )
    twice = ooxml.replace_part(once, chart, _bogus(_read(catalogue, chart)), tmp_path / "2.pptx")

    assert set(ooxml.validate_deck(twice).invalid) == {slide, chart}


def _edited(deck, part: str, edit, out):
    """``deck`` copied to ``out`` with ``edit`` applied to ``part``'s root."""
    root = etree.fromstring(_read(deck, part))
    edit(root)
    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return ooxml.replace_part(deck, part, data, out)


def _math_runs(root):
    return root.findall(f".//{{{_OMML}}}r")


def _second_math_root(root):
    first = root.find(f".//{{{_OMML}}}oMathPara")
    second = copy.deepcopy(first)
    second.insert(0, etree.Element(f"{{{_OMML}}}bogus"))
    first.addnext(second)


def _late_run_properties(root):
    run = next(r for r in _math_runs(root) if r.find(f"{{{_DML}}}rPr") is not None)
    run.append(run.find(f"{{{_DML}}}rPr"))


def _twice_run_properties(root):
    run = next(r for r in _math_runs(root) if r.find(f"{{{_DML}}}rPr") is not None)
    run.find(f"{{{_DML}}}rPr").addnext(copy.deepcopy(run.find(f"{{{_DML}}}rPr")))


def _small_third_run(root):
    root.findall(f".//{{{_OMML}}}r/{{{_DML}}}rPr")[2].set("sz", "50")


def _bare_omath(root):
    bare = etree.SubElement(root.find(f".//{{{_A14}}}m"), f"{{{_OMML}}}oMath")
    etree.SubElement(bare, f"{{{_OMML}}}bogus")


def _word_in_maths(root):
    root.find(f".//{{{_OMML}}}r").append(etree.Element(f"{{{ooxml.WML}}}br"))


_MATH_FAULTS = {
    "a WordprocessingML element in the maths": (_word_in_maths, "WordprocessingML"),
    "a fault in a second math root": (_second_math_root, "[math 2]"),
    "run properties after the run's text": (_late_run_properties, "rPr"),
    "two run properties in one run": (_twice_run_properties, "rPr"),
    "a size below the floor in the third run's properties": (_small_third_run, "sz"),
    "a bare oMath under the maths element": (_bare_omath, "bogus"),
}


@CATALOGUE
@pytest.mark.parametrize("fault", sorted(_MATH_FAULTS))
def test_the_whole_deck_gate_judges_every_piece_of_the_maths(catalogue, tmp_path, fault):
    """Every math root, run and run property, each judged apart from the Choice view that
    strips the `a14:m` holding them; the stripped namespace is still recorded."""
    part, _ = _part(catalogue, "the equation slide, in its Choice")
    edit, named = _MATH_FAULTS[fault]
    report = ooxml.validate_deck(_edited(catalogue, part, edit, tmp_path / "m.pptx"))

    assert list(report.invalid) == [part]
    assert any(e.startswith("[math ") and named in e for e in report.invalid[part]), report.invalid
    assert _A14 in report.passed_over[part]


@CATALOGUE
@pytest.mark.parametrize(
    "damage, named",
    [
        ("truncated", "XMLSyntaxError"),
        ("doctype", "DOCTYPE"),
        ("an internal entity", "DOCTYPE"),
        ("Latin-1", "not 1.0 in UTF-8"),
        ("XML 1.1", "not 1.0 in UTF-8"),
    ],
)
def test_a_part_that_cannot_be_judged_is_a_finding_under_its_name(
    catalogue, tmp_path, damage, named
):
    slide, _ = _part(catalogue, "first slide")
    data = _read(catalogue, slide)
    if damage == "truncated":
        data = data[: len(data) // 2]
    elif damage == "doctype":
        data = data.replace(b"?>", b"?><!DOCTYPE sld>", 1)
    elif damage == "an internal entity":
        data = data.replace(b"?>", b'?><!DOCTYPE sld [<!ENTITY e "x">]>', 1)
        data = data.replace(b' name="', b' name="&e;', 1)
    elif damage == "Latin-1":
        data = re.sub(rb"encoding=(['\"])UTF-8\1", b"encoding='ISO-8859-1'", data, count=1)
    else:
        data = re.sub(rb"version=(['\"])1\.0\1", b"version='1.1'", data, count=1)
    assert data != _read(catalogue, slide)
    report = ooxml.validate_deck(ooxml.replace_part(catalogue, slide, data, tmp_path / "d.pptx"))

    assert list(report.invalid) == [slide]
    assert named in report.invalid[slide][0]


@pytest.mark.parametrize("broken", ["Choice", "Fallback"])
def test_the_whole_part_check_reads_each_branch_on_its_own(schema, broken):
    """A fault confined to one branch is reported under that branch's name."""
    root = _equation_slide()
    sp = root.find(f".//{{{_MC}}}{broken}/{{{_PML}}}sp")
    sp.remove(sp.find(f"{{{_PML}}}spPr"))

    errors = ooxml.validate(root, schema)
    assert errors and all(e.startswith(f"[{broken}] ") for e in errors)
    assert "spPr" in errors[0]


def test_a_second_choice_is_read_as_a_view_of_its_own(schema):
    """A reader with the second Choice's namespace but not the first's sees the second."""
    root = _equation_slide()
    alternate = root.find(f".//{{{_MC}}}AlternateContent")
    second = etree.Element(f"{{{_MC}}}Choice", nsmap={"p14": _P14}, Requires="p14")
    etree.SubElement(second, f"{{{_PML}}}bogus")
    alternate.insert(1, second)

    errors = ooxml.validate(root, schema)
    assert errors and all(e.startswith("[Choice 2] ") for e in errors)
    assert "bogus" in errors[0]


def test_a_required_namespace_the_subtree_spells_its_own_way_is_still_stripped(schema):
    """Were `Requires` resolved after detaching, the `a14` it names would not resolve in the
    subtree, which spells that namespace `x`, and the maths element would stay to fail the view."""
    assert ooxml.validate(_equation_slide(prefix_on_root=True), schema) == []


def _wrapper_fault(root, fault: str) -> None:
    alternate = root.find(f".//{{{_MC}}}AlternateContent")
    choice = alternate.find(f"{{{_MC}}}Choice")
    fallback = alternate.find(f"{{{_MC}}}Fallback")
    if fault == "no Choice":
        alternate.remove(choice)
    elif fault == "two Fallbacks":
        alternate.append(copy.deepcopy(fallback))
    elif fault == "a Fallback before the Choice":
        alternate.insert(0, fallback)
    elif fault == "a stray child":
        alternate.append(etree.Element(f"{{{_PML}}}bogus"))
    elif fault == "an attribute on the wrapper":
        alternate.set("foo", "1")
    elif fault == "Requires on the Fallback":
        fallback.set("Requires", "a14")
    elif fault == "a Choice with no Requires":
        del choice.attrib["Requires"]
    elif fault == "an undeclared prefix":
        choice.set("Requires", "zz")
    elif fault == "an understood namespace":
        choice.set("Requires", "p")
    elif fault == "markup compatibility itself":
        choice.set("Requires", "mc")
    elif fault == "text in the wrapper":
        alternate.text = "junk"
    elif fault == "an attribute on the Choice":
        choice.set("foo", "1")
    elif fault == "a Fallback outside a wrapper":
        root.find(f".//{{{_PML}}}spTree").append(etree.Element(f"{{{_MC}}}Fallback"))
    elif fault == "an unknown markup-compatibility element":
        root.find(f".//{{{_PML}}}spTree").append(etree.Element(f"{{{_MC}}}Bogus"))
    elif fault == "MustUnderstand naming an undeclared prefix":
        root.set(f"{{{_MC}}}MustUnderstand", "zz")
    elif fault == "MustUnderstand naming a namespace nothing here has":
        choice.set(f"{{{_MC}}}MustUnderstand", "a14")
    elif fault == "ProcessContent naming what nothing ignores":
        choice.set(f"{{{_MC}}}ProcessContent", "a14")
    else:
        root.set(f"{{{_MC}}}Ignorable", "zz")


@pytest.mark.parametrize(
    "fault, named",
    [
        ("no Choice", "one or more Choice"),
        ("two Fallbacks", "one or more Choice"),
        ("a Fallback before the Choice", "one or more Choice"),
        ("a stray child", "one or more Choice"),
        ("an attribute on the wrapper", "AlternateContent carries attribute foo"),
        ("Requires on the Fallback", "Fallback carries attribute Requires"),
        ("a Choice with no Requires", "Choice with no Requires"),
        ("an undeclared prefix", "Requires names the undeclared prefix 'zz'"),
        ("an understood namespace", f"Requires names {_PML}"),
        ("markup compatibility itself", f"Requires names {_MC}"),
        ("an undeclared Ignorable prefix", "Ignorable names the undeclared prefix 'zz'"),
        ("text in the wrapper", "one or more Choice"),
        ("an attribute on the Choice", "Choice carries attribute foo"),
        ("a Fallback outside a wrapper", "Fallback outside an AlternateContent"),
        ("an unknown markup-compatibility element", "Bogus is no markup-compatibility element"),
        (
            "MustUnderstand naming an undeclared prefix",
            "MustUnderstand names the undeclared prefix",
        ),
        ("MustUnderstand naming a namespace nothing here has", f"MustUnderstand names {_A14}"),
        ("ProcessContent naming what nothing ignores", f"ProcessContent names {_A14}"),
    ],
)
def test_markup_compatibility_itself_is_judged(schema, fault, named):
    """What no branch view can see: the shape of a wrapper, its attributes, and what a
    Choice `Requires` or a part's `mc:Ignorable` names."""
    root = _equation_slide()
    _wrapper_fault(root, fault)

    assert any(e.startswith("[mc] ") and named in e for e in ooxml.validate(root, schema))


def _ignorable_slide(*, declared: bool):
    """A slide carrying a `p14` element and attribute where the ISO schema allows neither,
    under an `mc:Ignorable="p14"` on its root when ``declared``."""
    prs, slide, text, blank, chart = _deck()
    data = _slide_xml(prs).replace(
        b"<p:sld ", f'<p:sld xmlns:mc="{_MC}" xmlns:p14="{_P14}" '.encode(), 1
    )
    root = etree.fromstring(data)
    if declared:
        root.set(f"{{{_MC}}}Ignorable", "p14")
    tree = root.find(f".//{{{_PML}}}spTree")
    tree.set(f"{{{_P14}}}flag", "1")
    tree.insert(0, etree.Element(f"{{{_P14}}}extra"))
    return root


def test_what_an_ignorable_namespace_carries_is_skipped_as_a_reader_skips_it(schema):
    assert ooxml.validate(_ignorable_slide(declared=True), schema) == []
    assert ooxml.unjudged(_ignorable_slide(declared=True))[0] == {_P14}
    assert ooxml.validate(_ignorable_slide(declared=False), schema) != []
    root = _ignorable_slide(declared=True)
    root.find(f".//{{{_P14}}}extra").append(etree.Element(f"{{{_DML}}}tbl"))
    assert f"{{{_DML}}}tbl" in ooxml.unjudged(root)[1]


def test_an_ignorable_naming_an_iso_namespace_skips_nothing(schema):
    root = _ignorable_slide(declared=True)
    root.set(f"{{{_MC}}}Ignorable", "p14 a")
    root.find(f".//{{{_DML}}}off").set("x", "warp")
    assert any("off" in e for e in ooxml.validate(root, schema))


def test_an_ignorable_on_a_branch_is_honoured_in_that_branch(schema):
    root = _equation_slide()
    alternate = root.find(f".//{{{_MC}}}AlternateContent")
    fallback = alternate.find(f"{{{_MC}}}Fallback")
    replacement = etree.Element(f"{{{_MC}}}Fallback", nsmap={"s": "urn:skip"})
    replacement.set(f"{{{_MC}}}Ignorable", "s")
    replacement.extend(list(fallback))
    fallback.addprevious(replacement)
    alternate.remove(fallback)
    replacement.find(f"{{{_PML}}}sp").insert(1, etree.Element("{urn:skip}extra"))

    assert ooxml.validate(root, schema) == []
    del replacement.attrib[f"{{{_MC}}}Ignorable"]
    assert any("extra" in e for e in ooxml.validate(root, schema))


def _text_placed(where: str):
    if where == "after an element an Ignorable drops":
        root = _ignorable_slide(declared=True)
        root.find(f".//{{{_P14}}}extra").tail = "junk"
        return root
    root = _equation_slide()
    alternate = root.find(f".//{{{_MC}}}AlternateContent")
    if where == "after the wrapper":
        alternate.tail = "junk"
    elif where == "inside the Choice":
        alternate.find(f"{{{_MC}}}Choice").text = "junk"
    else:
        root.find(f".//{{{_A14}}}m").tail = "junk"
    return root


@pytest.mark.parametrize(
    "where",
    [
        "after the wrapper",
        "inside the Choice",
        "after the element the Choice strips",
        "after an element an Ignorable drops",
    ],
)
def test_text_a_reader_keeps_is_judged_wherever_markup_compatibility_moves_it(schema, where):
    assert any("Character content" in e for e in ooxml.validate(_text_placed(where), schema))


def _powerpoint_transition_slide(*, ignorable: bool):
    """A slide with a transition as PowerPoint writes one: `p14` content and a `p14`
    attribute in a Choice that `Requires` it, under a root `mc:Ignorable` when ``ignorable``."""
    prs, slide, text, blank, chart = _deck()
    data = _slide_xml(prs).replace(
        b"<p:sld ", f'<p:sld xmlns:mc="{_MC}" xmlns:p14="{_P14}" '.encode(), 1
    )
    root = etree.fromstring(data)
    if ignorable:
        root.set(f"{{{_MC}}}Ignorable", "p14")
    root.find(f"{{{_PML}}}clrMapOvr").addnext(
        etree.fromstring(
            f'<mc:AlternateContent xmlns:mc="{_MC}" xmlns:p="{_PML}" xmlns:p14="{_P14}">'
            '<mc:Choice Requires="p14"><p:transition spd="slow" p14:dur="2000"><p14:prism/>'
            '</p:transition></mc:Choice><mc:Fallback><p:transition spd="slow"><p:fade/>'
            "</p:transition></mc:Fallback></mc:AlternateContent>"
        )
    )
    return root


@pytest.mark.parametrize("ignorable", [False, True], ids=["Requires only", "and Ignorable"])
def test_a_transition_as_powerpoint_writes_it_validates(schema, ignorable):
    assert ooxml.validate(_powerpoint_transition_slide(ignorable=ignorable), schema) == []


def test_a_wrapper_with_fewer_choices_gives_every_later_view_its_first():
    root = etree.fromstring(
        f'<r xmlns:mc="{_MC}" xmlns:p14="{_P14}" xmlns:p15="urn:p15"><mc:AlternateContent>'
        '<mc:Choice Requires="p14"><first/></mc:Choice><mc:Choice Requires="p15"><second/>'
        "</mc:Choice></mc:AlternateContent></r>"
    )
    assert [c.tag for c in ooxml.branch(root, "Choice", 1)] == ["second"]
    assert [c.tag for c in ooxml.branch(root, "Choice", 2)] == ["first"]


def test_the_axis_id_copy_rewrites_only_what_python_pptx_signs(schema):
    """python-pptx writes `c:axId` and `c:crossAx` signed; the copy unsigns those and nothing
    else, leaves the part as it was, and any other bad value is still the schema's to name."""
    prs, slide, text, blank, chart = _deck()
    space = next(s for s in slide.shapes if s.shape_id == chart).chart._chartSpace
    axes = (f"{{{ooxml.CHART}}}axId", f"{{{ooxml.CHART}}}crossAx")
    assert any(int(e.get("val")) < 0 for e in space.iter(*axes))
    assert not schema.validate(space)
    assert "axId" in schema.error_log[0].message
    assert ooxml.validate(space, schema) == []

    unsigned = ooxml.unsigned_axis_ids(space)
    assert any(int(e.get("val")) < 0 for e in space.iter(*axes))
    for original, copied in zip(space.iter(*axes), unsigned.iter(*axes), strict=True):
        copied.set("val", original.get("val"))
    assert etree.tostring(unsigned) == etree.tostring(space)

    elsewhere = copy.deepcopy(space)
    elsewhere.find(f".//{{{ooxml.CHART}}}ser/{{{ooxml.CHART}}}idx").set("val", "-1")
    assert any("idx" in e for e in ooxml.validate(elsewhere, schema))
    with pytest.raises(AssertionError, match="idx"):
        ooxml.assert_chart_valid(elsewhere)
    for value in ("abc", None):
        on_axis = copy.deepcopy(space)
        axis = on_axis.find(f".//{{{ooxml.CHART}}}axId")
        axis.set("val", value) if value else axis.attrib.pop("val")
        assert any("axId" in e for e in ooxml.validate(on_axis, schema)), value


def test_a_bullet_size_in_thousandths_is_read_as_powerpoint_writes_it(schema):
    """PowerPoint writes `a:buSzPct` in thousandths of a percent where the schema wants
    `N%`; the copy rewrites a whole percent and leaves anything else to be named."""
    prs, slide, text, blank, chart = _deck()
    root = etree.fromstring(_slide_xml(prs))
    properties = etree.Element(f"{{{_DML}}}pPr")
    root.find(f".//{{{_DML}}}p").insert(0, properties)
    size = etree.SubElement(properties, f"{{{_DML}}}buSzPct", val="100000")

    rewritten = ooxml.percent_bullet_sizes(root)
    assert size.get("val") == "100000"
    rewritten.find(f".//{{{_DML}}}buSzPct").set("val", "100000")
    assert etree.tostring(rewritten) == etree.tostring(root)
    for value, valid in (
        ("100000", True),
        ("100%", True),
        ("500000", False),
        ("112500", False),
        ("\u0661\u0660\u0660\u0660\u0660\u0660", False),
    ):
        size.set("val", value)
        errors = ooxml.validate(root, schema)
        assert (errors == []) is valid, (value, errors)
        assert valid or "buSzPct" in errors[0]


def test_a_part_libxml2_will_not_canonicalise_still_has_a_canonical_form():
    """A relative namespace URI defeats C14N; the canonical form falls back to the
    serialisation, still blind to blank text between elements."""
    root = etree.fromstring(b'<r xmlns:x="relative"><x:b/></r>')
    with pytest.raises(etree.C14NError):
        etree.tostring(root, method="c14n")
    spaced = etree.fromstring(b'<r xmlns:x="relative">\n  <x:b/>\n</r>')
    other = etree.fromstring(b'<r xmlns:x="relative"><x:c v="1"/></r>')
    assert ooxml.canonical(root) == ooxml.canonical(spaced)
    assert ooxml.canonical(root) != ooxml.canonical(other)


def _retyped(
    data: bytes,
    part: str,
    content_type: str | None,
    *,
    duplicate: bool = False,
    upper: bool = False,
) -> bytes:
    """``[Content_Types].xml`` with ``part``'s override set to ``content_type`` (removed when
    ``None``), written twice when ``duplicate``, its PartName upper-cased when ``upper``."""
    root = etree.fromstring(data)
    overrides = root.iter(f"{{{ooxml.CONTENT_TYPES}}}Override")
    override = next(o for o in overrides if o.get("PartName") == "/" + part)
    if content_type is None:
        root.remove(override)
    else:
        override.set("ContentType", content_type)
    if upper:
        override.set("PartName", override.get("PartName").upper())
    if duplicate:
        root.append(copy.deepcopy(override))
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _extended(root, where: str) -> None:
    """Content the schemas pass over, added to a slide."""
    if where == "a namespace on an attribute alone":
        root.find(f".//{{{_A14}}}m").set(f"{{{_P14}}}flag", "1")
    elif where == "an element in no namespace inside an extLst entry":
        ext = etree.SubElement(root.find(f".//{{{_DML}}}extLst"), f"{{{_DML}}}ext", uri="{ABCD}")
        etree.SubElement(ext, "loose")
    elif where == "an ISO element directly inside an extLst entry":
        ext = etree.SubElement(root.find(f".//{{{_DML}}}extLst"), f"{{{_DML}}}ext", uri="{ABCD}")
        etree.SubElement(ext, f"{{{_DML}}}solidFill")
    elif where == "an ISO element under an undefined one inside an extLst entry":
        ext = etree.SubElement(root.find(f".//{{{_DML}}}extLst"), f"{{{_DML}}}ext", uri="{ABCD}")
        holder = etree.SubElement(ext, "{urn:x}holder")
        etree.SubElement(holder, f"{{{_PML}}}sldSz", cx="warp")
    elif where == "an ISO element in a wrapper inside an extLst entry":
        ext = etree.SubElement(root.find(f".//{{{_DML}}}extLst"), f"{{{_DML}}}ext", uri="{ABCD}")
        alternate = etree.SubElement(ext, f"{{{_MC}}}AlternateContent")
        choice = etree.SubElement(
            alternate, f"{{{_MC}}}Choice", nsmap={"p14": _P14}, Requires="p14"
        )
        etree.SubElement(choice, f"{{{_DML}}}solidFill")
    elif where == "an ISO element in a math run property's extension":
        properties = root.find(f".//{{{_OMML}}}r/{{{_DML}}}rPr")
        extensions = etree.SubElement(properties, f"{{{_DML}}}extLst")
        ext = etree.SubElement(extensions, f"{{{_DML}}}ext", uri="{ABCD}")
        etree.SubElement(ext, f"{{{_DML}}}solidFill")
    else:
        etree.SubElement(root.find(f".//{{{_A14}}}m"), f"{{{_DML}}}tbl")


_RECORDED = {
    "a namespace on an attribute alone": ("passed_over", _P14),
    "an element in no namespace inside an extLst entry": ("passed_over", ooxml.NO_NAMESPACE),
    "an ISO element directly inside an extLst entry": ("unreached", f"{{{_DML}}}solidFill"),
    "an ISO element under an undefined one inside an extLst entry": (
        "unreached",
        f"{{{_PML}}}sldSz",
    ),
    "an ISO element in a wrapper inside an extLst entry": ("unreached", f"{{{_DML}}}solidFill"),
    "an ISO element in a math run property's extension": ("unreached", f"{{{_DML}}}solidFill"),
    "a global ISO element under one a reader strips": ("unreached", f"{{{_DML}}}tbl"),
}


@CATALOGUE
@pytest.mark.parametrize("broken", [False, True], ids=["valid part", "invalid part"])
@pytest.mark.parametrize("where", sorted(_RECORDED))
def test_what_a_verdict_passes_over_is_recorded_whatever_the_verdict(
    catalogue, tmp_path, where, broken
):
    part, _ = _part(catalogue, "the equation slide, in its Choice")
    if "extLst" in where:
        slides = [s for s in _members(catalogue) if s.startswith("ppt/slides/")]
        part = next(s for s in slides if b"<a:extLst" in _read(catalogue, s))

    def edit(root):
        _extended(root, where)
        if broken:
            root.insert(0, etree.Element("{urn:deckwright:bogus}bogus"))

    report = ooxml.validate_deck(_edited(catalogue, part, edit, tmp_path / "x.pptx"))
    assert (part in report.invalid) is broken
    field, item = _RECORDED[where]
    assert item in getattr(report, field)[part]


@CATALOGUE
@pytest.mark.parametrize("how", ["an unlisted part", "an override with no content type"])
def test_a_part_with_no_content_type_is_a_fault(catalogue, tmp_path, how):
    slide, _ = _part(catalogue, "first slide")
    if how == "an unlisted part":
        part = "ppt/stray.xyz"
        deck = ooxml.replace_part(catalogue, part, b"<stray/>", tmp_path / "s.pptx")
    else:
        part = slide
        types = _retyped(_read(catalogue, ooxml.CONTENT_TYPES_PART), slide, "")
        deck = ooxml.replace_part(catalogue, ooxml.CONTENT_TYPES_PART, types, tmp_path / "s.pptx")

    assert ooxml.validate_deck(deck).invalid == {part: ("no content type in [Content_Types].xml",)}


_SLIDE_TYPE = f"{ooxml.OFFICE}presentationml.slide+xml"


@CATALOGUE
@pytest.mark.parametrize(
    "content_type, upper, named",
    [
        ("application/xml", False, "a type that does not own it"),
        ("image/png", False, "a type that does not own it"),
        (f"{ooxml.OFFICE}presentationml.unlisted+xml", False, "a type that does not own it"),
        (f"{ooxml.OFFICE}theme+xml", False, "expects"),
        (f"{ooxml.OFFICE}theme+xml", True, "expects"),
    ],
)
def test_a_slide_cannot_hide_behind_another_content_type(
    catalogue, tmp_path, content_type, upper, named
):
    """The override moved, the bytes untouched: no other type, generic, unlisted, not XML at
    all, or another Office type, shelters a root an Office type owns, however its PartName is cased."""
    slide, _ = _part(catalogue, "first slide")
    types = _retyped(_read(catalogue, ooxml.CONTENT_TYPES_PART), slide, content_type, upper=upper)
    retyped = ooxml.replace_part(catalogue, ooxml.CONTENT_TYPES_PART, types, tmp_path / "r.pptx")

    report = ooxml.validate_deck(retyped)
    assert list(report.invalid) == [slide]
    assert named in report.invalid[slide][0]


@CATALOGUE
@pytest.mark.parametrize(
    "damage, named",
    [
        ("truncated", "XMLSyntaxError"),
        ("duplicate", "appears twice"),
        ("absent", "absent from the package"),
    ],
)
def test_a_broken_content_types_part_is_a_fault_of_its_own(catalogue, tmp_path, damage, named):
    slide, _ = _part(catalogue, "first slide")
    data: bytes | None = _read(catalogue, ooxml.CONTENT_TYPES_PART)
    if damage == "truncated":
        data = data[: len(data) // 2]
    elif damage == "duplicate":
        data = _retyped(data, slide, _SLIDE_TYPE, duplicate=True)
    else:
        data = None
    out = tmp_path / "c.pptx"
    report = ooxml.validate_deck(ooxml.replace_part(catalogue, ooxml.CONTENT_TYPES_PART, data, out))

    assert ooxml.CONTENT_TYPES_PART in report.invalid
    assert named in report.invalid[ooxml.CONTENT_TYPES_PART][0]


@CATALOGUE
@pytest.mark.parametrize("twin", ["ppt/slides/Slide1.xml", "ppt/slides/slide1.xml"])
def test_two_members_one_name_apart_from_case_are_both_faults(catalogue, tmp_path, twin):
    """A member whose name differs from another only in case, or repeats it outright."""
    slide = "ppt/slides/slide1.xml"
    out = tmp_path / "t.pptx"
    with zipfile.ZipFile(catalogue) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            dst.writestr(info, src.read(info.filename))
        with pytest.warns(UserWarning) if twin == slide else contextlib.nullcontext():
            dst.writestr(twin, src.read(slide))

    report = ooxml.validate_deck(out)
    named = "appears twice in the archive" if twin == slide else "differs only in case"
    assert {slide, twin} <= set(report.invalid)
    assert all(named in report.invalid[name][0] for name in {slide, twin})


def _with_typed_part(deck, part: str, content_type: str, root: str, out):
    """``deck`` with ``part`` added, an empty element ``root``, under ``content_type``."""
    types = etree.fromstring(_read(deck, ooxml.CONTENT_TYPES_PART))
    override = f"{{{ooxml.CONTENT_TYPES}}}Override"
    etree.SubElement(types, override, PartName="/" + part, ContentType=content_type)
    data = etree.tostring(types, xml_declaration=True, encoding="UTF-8", standalone=True)
    staged = ooxml.replace_part(deck, ooxml.CONTENT_TYPES_PART, data, out.with_suffix(".t.pptx"))
    return ooxml.replace_part(staged, part, etree.tostring(etree.Element(root)), out)


_OFFICE_TYPE = "application/vnd.openxmlformats-officedocument."
_PML_TYPE = f"{_OFFICE_TYPE}presentationml."
_MACRO_TYPE = "application/vnd.ms-powerpoint."
_PRESENTATION = f"{{{_PML}}}presentation"
# Each content type a presentation package may declare, and the root its part must carry, as
# ISO/IEC 29500-1 and -2 give them.
_CONTENT_TYPE_ROOTS = {
    f"{_PML_TYPE}presentation.main+xml": _PRESENTATION,
    f"{_PML_TYPE}template.main+xml": _PRESENTATION,
    f"{_PML_TYPE}slideshow.main+xml": _PRESENTATION,
    f"{_MACRO_TYPE}presentation.macroenabled.main+xml": _PRESENTATION,
    f"{_MACRO_TYPE}template.macroenabled.main+xml": _PRESENTATION,
    f"{_MACRO_TYPE}slideshow.macroenabled.main+xml": _PRESENTATION,
    f"{_MACRO_TYPE}addin.macroenabled.main+xml": _PRESENTATION,
    f"{_PML_TYPE}slide+xml": f"{{{_PML}}}sld",
    f"{_PML_TYPE}slidelayout+xml": f"{{{_PML}}}sldLayout",
    f"{_PML_TYPE}slidemaster+xml": f"{{{_PML}}}sldMaster",
    f"{_PML_TYPE}notesslide+xml": f"{{{_PML}}}notes",
    f"{_PML_TYPE}notesmaster+xml": f"{{{_PML}}}notesMaster",
    f"{_PML_TYPE}handoutmaster+xml": f"{{{_PML}}}handoutMaster",
    f"{_PML_TYPE}presprops+xml": f"{{{_PML}}}presentationPr",
    f"{_PML_TYPE}viewprops+xml": f"{{{_PML}}}viewPr",
    f"{_PML_TYPE}tablestyles+xml": f"{{{_DML}}}tblStyleLst",
    f"{_PML_TYPE}tags+xml": f"{{{_PML}}}tagLst",
    f"{_PML_TYPE}commentauthors+xml": f"{{{_PML}}}cmAuthorLst",
    f"{_PML_TYPE}comments+xml": f"{{{_PML}}}cmLst",
    f"{_PML_TYPE}slideupdateinfo+xml": f"{{{_PML}}}sldSyncPr",
    f"{_OFFICE_TYPE}theme+xml": f"{{{_DML}}}theme",
    f"{_OFFICE_TYPE}themeoverride+xml": f"{{{_DML}}}themeOverride",
    f"{_OFFICE_TYPE}thememanager+xml": f"{{{_DML}}}themeManager",
    f"{_OFFICE_TYPE}drawingml.chart+xml": f"{{{ooxml.CHART}}}chartSpace",
    f"{_OFFICE_TYPE}drawingml.chartshapes+xml": f"{{{ooxml.CHART}}}userShapes",
    f"{_OFFICE_TYPE}drawingml.diagramdata+xml": f"{{{ooxml.DIAGRAM}}}dataModel",
    f"{_OFFICE_TYPE}drawingml.diagramlayout+xml": f"{{{ooxml.DIAGRAM}}}layoutDef",
    f"{_OFFICE_TYPE}drawingml.diagramstyle+xml": f"{{{ooxml.DIAGRAM}}}styleDef",
    f"{_OFFICE_TYPE}drawingml.diagramcolors+xml": f"{{{ooxml.DIAGRAM}}}colorsDef",
    "application/vnd.openxmlformats-package.relationships+xml": (
        "{http://schemas.openxmlformats.org/package/2006/relationships}Relationships"
    ),
    "application/vnd.openxmlformats-package.core-properties+xml": (
        "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}coreProperties"
    ),
    f"{_OFFICE_TYPE}extended-properties+xml": (
        "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}Properties"
    ),
    f"{_OFFICE_TYPE}custom-properties+xml": (
        "{http://schemas.openxmlformats.org/officeDocument/2006/custom-properties}Properties"
    ),
}


def test_the_content_type_table_is_the_standards():
    assert ooxml.EXPECTED_ROOT == _CONTENT_TYPE_ROOTS


@pytest.mark.parametrize("content_type, root", sorted(_CONTENT_TYPE_ROOTS.items()))
def test_each_content_type_holds_its_part_to_its_own_root(tmp_path, content_type, root):
    """Its own root passes the root check, and another row's root does not."""
    from tests.conftest import save_blank_deck

    stock = save_blank_deck(tmp_path / "stock.pptx")
    other = next(r for r in sorted(set(_CONTENT_TYPE_ROOTS.values())) if r != root)
    probe = "ppt/probe.xml"
    right = ooxml.validate_deck(
        _with_typed_part(stock, probe, content_type, root, tmp_path / "r.pptx")
    )
    wrong = ooxml.validate_deck(
        _with_typed_part(stock, probe, content_type, other, tmp_path / "w.pptx")
    )

    assert probe in {*right.valid, *right.unchecked, *right.invalid}
    if probe in right.invalid:
        assert not any("expects" in e or "own it" in e for e in right.invalid[probe])
    assert "expects" in wrong.invalid[probe][0]


@pytest.mark.parametrize(
    "case, part, content_type, root, verdict",
    [
        (
            "a slide named .bin under the slide type",
            "ppt/slides/extra.bin",
            f"{_PML_TYPE}slide+xml",
            f"{{{_PML}}}sld",
            "invalid",
        ),
        (
            "generic XML under another name",
            "ppt/data.dat",
            "application/xml",
            "{urn:x}data",
            "unchecked",
        ),
        ("a VML drawing", "ppt/drawings/v1.vml", f"{_OFFICE_TYPE}vmlDrawing", "xml", "unchecked"),
        (
            "a VML drawing holding a slide",
            "ppt/drawings/v1.vml",
            f"{_OFFICE_TYPE}vmlDrawing",
            f"{{{_PML}}}sld",
            "own it",
        ),
        (
            "a declared root under an unlisted Office type",
            "ppt/probe.xml",
            f"{_PML_TYPE}unlisted+xml",
            f"{{{_DML}}}tbl",
            "own it",
        ),
    ],
)
def test_the_package_judges_every_xml_part_however_it_is_named_or_typed(
    tmp_path, case, part, content_type, root, verdict
):
    from tests.conftest import save_blank_deck

    stock = save_blank_deck(tmp_path / "stock.pptx")
    report = ooxml.validate_deck(
        _with_typed_part(stock, part, content_type, root, tmp_path / "p.pptx")
    )
    if verdict == "unchecked":
        assert part in report.unchecked
    elif verdict == "invalid":
        assert part in report.invalid
    else:
        assert verdict in report.invalid[part][0]


def test_the_content_types_part_is_found_whatever_its_case(tmp_path):
    from tests.conftest import save_blank_deck

    stock = save_blank_deck(tmp_path / "stock.pptx")
    out = tmp_path / "lower.pptx"
    with zipfile.ZipFile(stock) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            name = (
                info.filename.lower()
                if info.filename == ooxml.CONTENT_TYPES_PART
                else info.filename
            )
            dst.writestr(name, src.read(info.filename))

    report = ooxml.validate_deck(out)
    assert report.invalid == {}
    assert "[content_types].xml" in report.unchecked


def test_a_global_iso_element_a_lax_wildcard_reaches_is_judged_and_not_recorded(schema):
    """Directly inside an `extLst` entry, a global element is the wildcard's to judge; the same
    element under an element a reader strips is recorded instead (`_RECORDED`)."""
    prs, slide, text, blank, chart = _deck()
    root = etree.fromstring(_slide_xml(prs))
    extensions = etree.SubElement(root.find(f".//{{{_PML}}}cNvPr"), f"{{{_DML}}}extLst")
    ext = etree.SubElement(extensions, f"{{{_DML}}}ext", uri="{ABCD}")
    etree.SubElement(ext, f"{{{_DML}}}tbl")

    assert any("tbl" in e for e in ooxml.validate(root, schema))
    assert f"{{{_DML}}}tbl" not in ooxml.unjudged(root)[1]
