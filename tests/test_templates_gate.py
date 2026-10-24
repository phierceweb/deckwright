"""Hand `tests/test_templates.py`'s own assertions a stock `Presentation()`, sound and broken,
and fail where they cannot tell the two apart. Needs no brand template, so it runs everywhere."""

from __future__ import annotations

import os
import re
import zipfile
from types import SimpleNamespace

import pytest
from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from deckwright.conform.derive import derive
from deckwright.theme.stock import is_stock_accent
from tests.ooxml import CONTENT_TYPES, CONTENT_TYPES_PART, DML, OFFICE, PML, replace_part


@pytest.fixture(scope="module")
def blank(tmp_path_factory):
    """Stock Office, 16:9, and nothing else — the same file `synthetic_template` makes."""
    path = tmp_path_factory.mktemp("blank") / "blank.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.save(str(path))
    return path


@pytest.fixture(scope="module")
def blank_bind(blank):
    return derive(blank).get("bind", {})


def test_a_blank_office_file_binds_no_brand_accent(blank_bind):
    """Stock accents say nothing about a brand, so `derive` rejects all six. Read `bind` here, never
    a loaded `Theme` — `palette.accents` answers with deckwright's own four defaults for this file."""
    bound = [role for role in blank_bind if role.startswith("accent-")]
    assert bound == [], (
        f"a blank Office file bound {bound} — either derive stopped rejecting stock "
        f"accents, or the stock table no longer covers this Office generation"
    )


def test_every_accent_in_a_blank_file_is_one_microsoft_ships(blank):
    """Why the assertion above holds, stated over the file rather than the code path,
    so a `derive` that silently stopped filtering cannot make both agree."""
    from deckwright.theme.clrscheme import parse_color_scheme, read_theme_xml

    scheme = parse_color_scheme(read_theme_xml(Presentation(str(blank)).slide_masters[0]))
    accents = {slot: scheme[slot] for slot in scheme if slot.startswith("accent")}
    assert accents, "the stock scheme declares no accents — this file changed shape"
    assert [s for s, v in sorted(accents.items()) if not is_stock_accent(v)] == []


def test_the_generated_sample_does_pass_that_same_gate(tmp_path):
    """The negative control: if a blank file and a designed one both fail, the gate is broken rather
    than discriminating."""
    from deckwright.conform.sample import write_sample

    bind = derive(write_sample(tmp_path / "sample.pptx")).get("bind", {})
    bound = [role for role in bind if role.startswith("accent-")]
    assert bound == [f"accent-{i}" for i in range(1, 7)]


def test_a_loaded_theme_still_reports_accents_for_a_template_that_bound_none(blank):
    """Asserts the `Theme.palette` fallback EXISTS — correct behaviour for rendering, and wrong as
    evidence about a template."""
    from deckwright.theme import load_theme

    theme = load_theme(None)
    assert theme.palette.accents, (
        "the default palette no longer carries accents — if that is deliberate, the "
        "corpus assertions may read palette.accents again"
    )


# --- the loop closed: the real assertions, handed the failure case ------------------


@pytest.fixture(scope="module")
def blank_built(blank, tmp_path_factory):
    """What `conform` would hand the corpus assertions for a blank template. Assembled rather than
    built: the assertions below read only the derived theme and the template beside it."""
    import yaml

    out = tmp_path_factory.mktemp("blank-conform")
    theme = derive(blank)
    theme["template"] = os.path.relpath(blank, out)
    theme_path = out / "blank.theme.yaml"
    theme_path.write_text(yaml.safe_dump(theme, sort_keys=False))

    class _Built:
        template = blank.name
        theme = theme_path

    return _Built()


def test_a_blank_office_file_offers_no_brand_accent_to_bind(blank_built):
    """What makes the accent assertions in `test_templates.py` non-vacuous: they read the
    template's own scheme, so a file carrying only stock colours has nothing for them to examine.
    Revert either to reading `Theme.palette` — which deckwright fills from its own defaults — and this
    reddens, because the blank file would suddenly appear to own brand accents."""
    import tests.test_templates as corpus

    assert corpus._brand_accents_in_scheme(blank_built) == {}
    assert corpus._bound_accents(blank_built) == {}


def test_those_same_assertions_accept_the_generated_sample(tmp_path):
    """The other half: a designed template must offer brand accents and pass both assertions, or
    the check above is merely blind rather than discriminating."""
    import yaml

    import tests.test_templates as corpus
    from deckwright.conform.sample import write_sample

    sample = write_sample(tmp_path / "sample.pptx")
    theme = derive(sample)
    theme["template"] = sample.name
    theme_path = tmp_path / "sample.theme.yaml"
    theme_path.write_text(yaml.safe_dump(theme, sort_keys=False))

    class _Built:
        template = sample.name
        theme = theme_path

    built = _Built()
    corpus.test_every_brand_accent_the_template_owns_is_bound(built)
    corpus.test_no_unedited_microsoft_accent_is_adopted_as_a_brand_colour(built)


# --- the schema assertions, handed a stock file ------------------------------------------


@pytest.fixture
def stock(tmp_path):
    """A one-slide stock Office file, standing as template and build alike."""
    from tests.conftest import save_blank_deck

    return save_blank_deck(tmp_path / "stock.pptx")


@pytest.fixture
def furnished(tmp_path):
    """A stock file whose one slide carries a chart and speaker notes."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("S", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), data
    )
    slide.notes_slide.notes_text_frame.text = "spoken"
    path = tmp_path / "furnished.pptx"
    prs.save(str(path))
    return path


def _damaged(
    deck,
    part: str,
    out,
    *,
    bogus: bool = False,
    bad_value: int = 0,
    shifted: bool = False,
    doctype: bool = False,
    reorder: bool = False,
    indent: bool = False,
    retype: str | None = None,
    truncate: bool = False,
    blank_beside_comment: bool = False,
):
    """``deck`` copied to ``out`` with ``part`` changed as the keywords say."""
    with zipfile.ZipFile(deck) as archive:
        root = etree.fromstring(archive.read(part))
        types = etree.fromstring(archive.read(CONTENT_TYPES_PART))
    if bogus:
        root.insert(0, etree.Element("{urn:deckwright:bogus}bogus"))
    if bad_value:
        root.findall(f".//{{{DML}}}off")[bad_value - 1].set("x", "warp")
    if blank_beside_comment:
        offset = root.find(f".//{{{DML}}}off")
        offset.text = "  "
        offset.append(etree.Comment("x"))
    if reorder:
        mapping = root.find(f".//{{{PML}}}clrMap")
        pairs = list(mapping.attrib.items())
        mapping.attrib.clear()
        for name, value in reversed(pairs):
            mapping.set(name, value)
    data = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True, pretty_print=indent
    )
    if shifted:
        data = data.replace(b"?>", b"?>\n<!-- shifted -->\n", 1)
    if doctype:
        data = data.replace(b"?>", b"?><!DOCTYPE sldMaster>", 1)
    if truncate:
        data = data[: len(data) // 2]
    result = replace_part(deck, part, data, out)
    if retype:
        override = next(
            o for o in types.iter(f"{{{CONTENT_TYPES}}}Override") if o.get("PartName") == "/" + part
        )
        override.set("ContentType", retype)
        typed = etree.tostring(types, xml_declaration=True, encoding="UTF-8", standalone=True)
        result = replace_part(result, CONTENT_TYPES_PART, typed, out.with_suffix(".typed.pptx"))
    return result


def _built(deck, template):
    return SimpleNamespace(deck=deck, template=str(template))


def test_the_schema_assertions_accept_a_stock_file(stock):
    import tests.test_templates as corpus

    corpus.test_every_part_the_build_wrote_is_valid_ooxml(_built(stock, stock))
    corpus.test_the_templates_own_parts_are_reported_apart(_built(stock, stock))
    corpus.test_what_the_build_leaves_unjudged_is_known(_built(stock, stock))


_MASTER = "ppt/slideMasters/slideMaster1.xml"
_SLIDE = "ppt/slides/slide1.xml"
_NOTES = "ppt/notesSlides/notesSlide1.xml"
_CHART = "ppt/charts/chart1.xml"
_THEME_TYPE = f"{OFFICE}theme+xml"


@pytest.mark.parametrize(
    "case, part, template_damage, build_damage",
    [
        ("a slide broken in the build only", _SLIDE, {}, {"bogus": True}),
        ("a notes slide broken the same way in both", _NOTES, {"bogus": True}, {"bogus": True}),
        ("a chart broken the same way in both", _CHART, {"bogus": True}, {"bogus": True}),
        ("a master truncated in both", _MASTER, {"truncate": True}, {"truncate": True}),
        (
            "a sound master given blank text beside a comment",
            _MASTER,
            {},
            {"blank_beside_comment": True},
        ),
        ("a master broken in the build only", _MASTER, {}, {"bogus": True}),
        ("a slide broken the same way in both", _SLIDE, {"bogus": True}, {"bogus": True}),
        (
            "a master with the template's fault and another",
            _MASTER,
            {"bad_value": 1},
            {"bad_value": 1, "bogus": True},
        ),
        ("a master with the same fault elsewhere", _MASTER, {"bad_value": 1}, {"bad_value": 2}),
        ("a sound master the build retyped", _MASTER, {}, {"retype": _THEME_TYPE}),
        (
            "a master with the template's fault, retyped",
            _MASTER,
            {"bogus": True},
            {"bogus": True, "retype": _THEME_TYPE},
        ),
        ("a sound master the build gave a DOCTYPE", _MASTER, {}, {"doctype": True}),
        (
            "a master with the template's fault, given a DOCTYPE",
            _MASTER,
            {"bogus": True},
            {"bogus": True, "doctype": True},
        ),
    ],
)
def test_a_fault_the_template_does_not_account_for_is_the_builds_own(
    furnished, tmp_path, case, part, template_damage, build_damage
):
    """A slide, notes slide or chart is the build's whatever the template's holds; a carried
    part is the build's unless the template's has its digest and its errors."""
    import tests.test_templates as corpus

    template = (
        _damaged(furnished, part, tmp_path / "template.pptx", **template_damage)
        if template_damage
        else furnished
    )
    build = _damaged(furnished, part, tmp_path / "build.pptx", **build_damage)
    with pytest.raises(AssertionError, match=re.escape(part)):
        corpus.test_every_part_the_build_wrote_is_valid_ooxml(_built(build, template))


@pytest.mark.parametrize(
    "template_extra, build_extra",
    [
        ({}, {}),
        ({}, {"shifted": True}),
        ({}, {"reorder": True}),
        ({"indent": True}, {}),
    ],
    ids=["same bytes", "a comment before the root", "attributes reordered", "indented template"],
)
def test_a_fault_the_template_already_has_is_reported_apart(
    stock, tmp_path, template_extra, build_extra
):
    """The template's faulty master carried over with the same canonical form, however its
    bytes differ: the build's assertion passes, and the other names the part as the template's."""
    import tests.test_templates as corpus

    template = _damaged(stock, _MASTER, tmp_path / "template.pptx", bogus=True, **template_extra)
    build = _damaged(stock, _MASTER, tmp_path / "build.pptx", bogus=True, **build_extra)
    corpus.test_every_part_the_build_wrote_is_valid_ooxml(_built(build, template))
    with pytest.raises(pytest.xfail.Exception, match=re.escape(_MASTER)):
        corpus.test_the_templates_own_parts_are_reported_apart(_built(build, template))


def _foreign_in_an_extension(stock) -> bytes:
    with zipfile.ZipFile(stock) as archive:
        root = etree.fromstring(archive.read(_SLIDE))
    ext_list = etree.SubElement(root.find(f".//{{{PML}}}cNvPr"), f"{{{DML}}}extLst")
    ext = etree.SubElement(ext_list, f"{{{DML}}}ext", uri="{ABCD}")
    etree.SubElement(ext, "{urn:stray}thing")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


@pytest.mark.parametrize(
    "case", ["a part with a root nobody declared", "a namespace that is not Microsoft's"]
)
def test_what_a_build_leaves_unjudged_must_be_of_a_known_kind(stock, tmp_path, case):
    import tests.test_templates as corpus

    if case == "a part with a root nobody declared":
        part, data = "ppt/stray.xml", b'<stray xmlns="urn:stray"/>'
    else:
        part, data = _SLIDE, _foreign_in_an_extension(stock)
    build = replace_part(stock, part, data, tmp_path / "build.pptx")
    with pytest.raises(AssertionError, match="stray"):
        corpus.test_what_the_build_leaves_unjudged_is_known(_built(build, stock))
